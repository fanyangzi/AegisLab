"""Browser workspace: validated snapshots, versioned mutations and resource checks.

This is a local/single-owner workspace, NOT an institutional identity system or
chemical-safety certification. The existing review API is left intact.
"""
from __future__ import annotations

import copy
import hashlib
import hmac
import json
import os
import re
import sqlite3
import secrets
import shutil
import zipfile
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from posixpath import normpath as posix_normpath
from typing import Any, Literal
from urllib.parse import urlparse
from uuid import uuid4

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Request
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from .parser import parse_sop

RULE_PACK = "resource-prerequisites/1.0"
ROOT = Path(__file__).resolve().parents[2]
SOURCE_REGISTRY_PATH = ROOT / "data" / "knowledge_sources.json"

# Original uploads are retained only as explicitly attached, bounded source
# artifacts.  The text extractor remains the authority for the workspace
# document; this volume lets a reviewer reopen the bytes that produced that
# text without exposing filesystem paths or accepting arbitrary file types.
ARTIFACT_MAX_BYTES = 8 * 1024 * 1024
ARTIFACT_SUFFIXES = frozenset({".txt", ".md", ".docx", ".pdf"})
ARCHIVE_MAX_BYTES = 32 * 1024 * 1024
ARCHIVE_MAX_FILES = 512
ARCHIVE_FORMAT = "aegislab-workspace-zip/v1"
_ARTIFACT_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,127}$")

def artifact_root_for_db(db_path: str) -> Path:
    configured = os.getenv("AEGIS_SPATIAL_ARTIFACT_DIR", "").strip()
    root = Path(configured).expanduser() if configured else Path(db_path).with_suffix(".artifacts")
    root.mkdir(parents=True, exist_ok=True)
    return root

def _artifact_filename(filename: str) -> str:
    """Return a display filename without allowing path components through."""
    clean = Path(str(filename or "source")).name.strip() or "source"
    return clean[:240]


def _safe_archive_member(name: str) -> str:
    """Validate a ZIP member name before it is read or written.

    ZIP member names are POSIX paths even on macOS/Windows.  Never pass a
    caller supplied name to ``extractall``; the restore path only accepts the
    two manifest files and ``artifacts/<id>/<basename>`` entries.
    """
    value = str(name or "")
    if not value or "\x00" in value or value.startswith(("/", "\\")) or "\\" in value:
        raise ValueError("存档包含不安全的文件路径。")
    normalized = posix_normpath(value)
    if normalized != value or normalized in {".", ".."} or any(part in {"", ".", ".."} for part in value.split("/")):
        raise ValueError("存档包含路径穿越条目。")
    return value


def _safe_artifact_storage_path(root: Path, storage_key: str) -> Path:
    """Resolve an internal storage key while preventing path escape."""
    candidate = (root / str(storage_key)).resolve()
    resolved_root = root.resolve()
    if candidate.parent != resolved_root or candidate == resolved_root:
        raise HTTPException(500, "来源文件存储记录无效。")
    return candidate

WORKSPACE_ROLES = frozenset({"instructor", "researcher", "safety_owner", "ehs_manager", "administrator"})

# The identity is resolved once by local_owner for each request and stored on
# Request.state.  Handlers then receive the same server-side role/principal
# without trusting a role/actor field from the JSON body.

def token_digest(token: str) -> str:
    """Hash a bearer token before it reaches durable storage or diagnostics."""
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def job_execution_mode() -> str:
    """Return the explicit queued-job execution mode for this process."""
    value = os.getenv("AEGIS_SPATIAL_JOB_EXECUTION", "process-local").strip().lower()
    return value if value in {"process-local", "worker"} else "process-local"


def worker_lease_seconds() -> float:
    """Return the durable Worker lease timeout.

    A running row is only eligible for recovery after this interval.  Keeping
    the value configurable makes a local restart drill deterministic while a
    normal deployment gets a long enough lease for model calls.  Invalid or
    non-positive values deliberately fall back to the safe default instead of
    allowing two workers to reclaim a live job immediately.
    """
    raw = os.getenv("AEGIS_SPATIAL_WORKER_LEASE_SECONDS", "900").strip()
    try:
        value = float(raw)
    except (TypeError, ValueError):
        return 900.0
    if value <= 0 or value > 7 * 24 * 60 * 60:
        return 900.0
    return value

# The browser scene uses a deliberately small set of semantic resource kinds.
# Keeping their approximate footprint here gives the server enough information
# to place a newly-created resource without making the scene layout part of the
# approval basis.  These are placement hints, not a claim about certified
# equipment dimensions.
EQUIPMENT_FOOTPRINTS: dict[str, tuple[float, float]] = {
    "hood": (3.0, 1.5), "bench": (2.4, 1.4), "storage": (1.8, 1.2),
    "analyzer": (2.2, 1.8), "sink": (1.8, 1.3), "waste": (1.4, 1.2),
    "clean_bench": (2.4, 1.5), "biosafety_cabinet": (2.4, 1.6), "balance": (1.6, 1.1),
    "centrifuge": (1.8, 1.5), "pcr": (1.6, 1.3), "gc": (2.4, 1.7), "hplc": (2.5, 1.7),
    "uv_vis": (1.8, 1.3), "mass_spec": (3.0, 2.0), "elemental_analyzer": (2.6, 1.8),
    "evaporator": (1.8, 1.4), "vacuum_pump": (1.4, 1.1), "nitrogen_blowdown": (1.8, 1.2),
    "water_purification": (1.8, 1.1), "refrigerator": (1.5, 1.2), "freezer": (1.7, 1.3),
    "drying_oven": (1.8, 1.4), "furnace": (1.6, 1.4), "gas_cabinet": (1.7, 1.3),
    "flammable_cabinet": (1.8, 1.2), "acid_base_cabinet": (1.8, 1.2), "solvent_waste": (1.4, 1.1),
    "solid_waste": (1.4, 1.1), "eyewash": (1.0, .9), "safety_shower": (1.2, 1.0),
    "gas_detector": (.8, .8), "environment_monitor": (.8, .8), "access_control": (1.0, .6),
    "chromatography": (2.5, 1.7), "ion_chromatography": (2.5, 1.7),
    "glovebox": (2.8, 1.6), "nmr": (3.2, 2.4), "ftir": (1.8, 1.4), "raman": (2.0, 1.5),
    "xrd": (3.0, 2.0), "xps": (3.4, 2.4), "sem": (3.0, 2.0), "tem": (3.6, 2.6), "afm": (2.2, 1.8), "epr": (2.8, 2.0),
    "icp_ms": (3.0, 2.0), "icp_oes": (2.8, 1.9), "aas": (2.4, 1.7), "xrf": (2.5, 1.7), "spr": (2.1, 1.6), "dls": (1.8, 1.4),
    "bet": (2.5, 1.8), "tga": (2.2, 1.7), "dsc": (2.0, 1.5), "reactor": (2.8, 2.0), "microwave_reactor": (2.0, 1.6),
    "photoreactor": (2.2, 1.6), "electrochemistry": (1.8, 1.4), "preparative_hplc": (2.8, 1.9), "solid_phase_extraction": (1.8, 1.4),
    "liquid_nitrogen": (1.6, 1.4), "cold_trap": (1.4, 1.2), "gas_manifold": (1.8, 1.2), "exhaust_treatment": (2.0, 1.4),
    "lyophilizer": (2.5, 1.8), "spray_dryer": (2.4, 1.8), "co2_incubator": (1.8, 1.5), "gel_documentation": (1.8, 1.4),
    "electrophoresis": (1.6, 1.2), "flow_cytometer": (2.5, 1.8), "fluorescence_microscope": (2.2, 1.7), "cleanroom": (3.0, 2.5),
}
PLACEMENT_GAP = 0.35

def now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

def uid(prefix: str) -> str:
    return f"{prefix}-{str(uuid4())[:8]}"

def dt(value: str) -> datetime:
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        raise ValueError("时间必须包含时区。")
    return parsed

def overlap(a: dict, b: dict) -> bool:
    return dt(a["start"]) < dt(b["end"]) and dt(b["start"]) < dt(a["end"])

def canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


# The browser workspace has one configured local owner, but business changes
# still need a durable hand-off point.  These entries are a local outbox, not
# an assertion that email, chat, or institutional notifications were sent.
# Keeping the mapping server-side means a browser cannot manufacture a
# notification by posting an arbitrary event type.
MUTATION_NOTIFICATION_TYPES: dict[str, tuple[str, str]] = {
    "approval_decision": ("approval.updated", "安全负责人"),
    "approve": ("approval.updated", "安全负责人"),
    "save_task": ("task.created", "任务负责人"),
    "assign_task": ("task.assigned", "任务负责人"),
    "update_task": ("task.updated", "任务负责人"),
    "save_document": ("evidence.submitted", "安全负责人"),
    "verify_document": ("evidence.verified", "安全负责人"),
    "transition_plan": ("plan.lifecycle", "计划负责人"),
    "save_branch": ("plan.branch_saved", "计划负责人"),
    "apply_branch": ("plan.branch_applied", "安全负责人"),
}


def mutation_notification(result: dict, action: str, payload: dict, actor: str, request_id: str) -> dict | None:
    """Build one auditable outbox record for a meaningful workspace change.

    The record deliberately contains the resulting revision and event summary,
    rather than a mutable copy of the whole workspace.  Consumers must fetch
    the current object before acting, and the outbox remains useful after a
    browser reconnect without pretending to deliver an external message.
    """
    descriptor = MUTATION_NOTIFICATION_TYPES.get(action)
    if not descriptor:
        return None
    event_type, audience_label = descriptor
    event = (result.get("events") or [{}])[-1]
    object_id = str(event.get("objectId") or payload.get("id") or payload.get("planId") or "workspace")
    plan_id = payload.get("planId")
    if not plan_id and action in {"save_task", "save_branch"}:
        nested = payload.get("task") or payload.get("branch") or {}
        plan_id = nested.get("planId")
    if not plan_id and action == "apply_branch":
        branch = next((item for item in result.get("branches", []) if item.get("id") == payload.get("id")), None)
        plan_id = branch.get("planId") if branch else None
    task = None
    task_id = object_id if action in {"save_task", "assign_task", "update_task"} else payload.get("id")
    if task_id and action in {"save_task", "assign_task", "update_task"}:
        task = next((item for item in result.get("tasks", []) if item.get("id") == task_id), None)
        plan_id = plan_id or (task or {}).get("planId")
    audience = audience_label
    if task and task.get("owner"):
        audience = f"任务负责人:{task['owner']}"
    title = {
        "approval.updated": "审批链需要关注",
        "task.created": "新增协作任务",
        "task.assigned": "协作任务已转交",
        "task.updated": "协作任务状态已更新",
        "evidence.submitted": "资料等待独立核验",
        "evidence.verified": "资料核验状态已更新",
        "plan.lifecycle": "计划生命周期已更新",
        "plan.branch_saved": "候选方案等待复核",
        "plan.branch_applied": "候选方案已应用并触发重检",
    }.get(event_type, "工作区状态已更新")
    return {
        "eventId": uid("notify"),
        "idempotencyKey": f"mutation:{request_id}",
        "eventType": event_type,
        "aggregateType": "task" if task else "plan" if plan_id else "workspace",
        "aggregateId": object_id,
        "planId": plan_id,
        "audience": audience,
        "title": title,
        "summary": str(event.get("summary") or "工作区状态已更新")[:500],
        "actor": actor,
        "workspaceRevision": int(result.get("revision", 0)),
        "sourceEventId": event.get("id"),
        "status": "pending",
        "createdAt": str(event.get("at") or now()),
        "acknowledgedAt": None,
        "acknowledgedBy": None,
    }

class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid")

class Lab(StrictModel):
    id: str = Field(min_length=1, max_length=128)
    name: str = Field(min_length=1, max_length=180)
    description: str = Field(max_length=2000)
    width: float = Field(ge=6, le=40)
    depth: float = Field(ge=6, le=40)

class Maintenance(StrictModel):
    id: str
    start: str
    end: str
    reason: str = Field(max_length=2000)

class EquipmentSourceCitation(StrictModel):
    """A reviewable pointer from an example asset to a public source record.

    The citation proves only that a public page is relevant to a category or
    control topic. It never asserts that the page is an institutional asset
    register or that the example resource is currently installed there.
    """
    sourceId: str = Field(min_length=1, max_length=120)
    relation: Literal["category-or-control-reference"]
    note: str = Field(min_length=1, max_length=500)


EquipmentFamily = Literal[
    "containment", "workstation", "separation", "spectroscopy", "mass_analysis",
    "structure_material", "thermal_process", "cold_chain", "gas_vacuum",
    "bio_molecular", "safety_response", "waste_environment", "utility",
]
EvidenceKind = Literal["sop", "inspection", "sds", "training", "calibration", "reservation", "permit", "sensor"]

class Equipment(StrictModel):
    id: str
    labId: str
    code: str = Field(min_length=1, max_length=60)
    name: str = Field(min_length=1, max_length=180)
    kind: Literal[
        "hood", "bench", "storage", "analyzer", "sink", "waste", "clean_bench", "biosafety_cabinet",
        "balance", "centrifuge", "pcr", "gc", "hplc", "uv_vis", "mass_spec", "elemental_analyzer",
        "evaporator", "vacuum_pump", "nitrogen_blowdown", "water_purification", "refrigerator", "freezer",
        "drying_oven", "furnace", "gas_cabinet", "flammable_cabinet", "acid_base_cabinet", "solvent_waste",
        "solid_waste", "eyewash", "safety_shower", "gas_detector", "environment_monitor", "access_control",
        "chromatography", "ion_chromatography", "glovebox", "nmr", "ftir", "raman", "xrd", "xps", "sem", "tem", "afm", "epr",
        "icp_ms", "icp_oes", "aas", "xrf", "spr", "dls", "bet", "tga", "dsc", "reactor", "microwave_reactor", "photoreactor",
        "electrochemistry", "preparative_hplc", "solid_phase_extraction", "liquid_nitrogen", "cold_trap", "gas_manifold", "exhaust_treatment",
        "lyophilizer", "spray_dryer", "co2_incubator", "gel_documentation", "electrophoresis", "flow_cytometer", "fluorescence_microscope", "cleanroom"
    ]
    x: float = Field(ge=-20, le=20)
    z: float = Field(ge=-20, le=20)
    rotation: float = Field(ge=-360, le=360)
    # Optional semantic room/zone label.  Existing workspaces do not need to
    # carry it; a branch may propose a zone change without inventing a lab
    # or pretending that a coordinate is a named operational area.
    zone: str | None = Field(default=None, max_length=180)
    capacity: int = Field(ge=1, le=100, strict=True)
    status: Literal["available", "unavailable", "unknown"]
    owner: str = Field(max_length=180)
    description: str = Field(max_length=5000)
    requiresEvidence: bool
    specVersion: int = Field(ge=1, strict=True)
    maintenance: list[Maintenance] = Field(max_length=100)
    # Operational profile fields are optional so old exported workspaces remain valid.
    family: EquipmentFamily | None = None
    subtype: str | None = Field(default=None, max_length=180)
    aliases: list[str] = Field(default_factory=list, max_length=30)
    typicalUse: str | None = Field(default=None, max_length=1000)
    hazardTags: list[str] = Field(default_factory=list, max_length=30)
    controlTags: list[str] = Field(default_factory=list, max_length=40)
    utilityRequirements: list[str] = Field(default_factory=list, max_length=30)
    requiredTraining: list[str] = Field(default_factory=list, max_length=30)
    calibrationOrInspection: list[str] = Field(default_factory=list, max_length=30)
    evidenceKinds: list[EvidenceKind] = Field(default_factory=list, max_length=20)
    sourceRefs: list[str] = Field(default_factory=list, max_length=20)
    sourceCitations: list[EquipmentSourceCitation] = Field(default_factory=list, max_length=20)

class Step(StrictModel):
    id: str
    text: str = Field(min_length=1, max_length=20000)
    line: int = Field(ge=1, strict=True)
    equipmentIds: list[str] = Field(max_length=100)
    confirmed: bool

class Reservation(StrictModel):
    id: str
    equipmentId: str
    start: str
    end: str

class BranchPatch(StrictModel):
    """A reviewable field-level change inside an isolated plan branch.

    Branch reservations remain the candidate state used by the checker.  The
    patch records make the user's intent explicit, so a resource switch, a
    time change, or a resource-condition change can be accepted together
    without treating the branch as an opaque replacement snapshot.  ``before``
    is checked against the branch base and ``after`` against the candidate
    state at save/apply time.
    """
    path: Literal[
        "reservation.equipmentId", "reservation.start", "reservation.end", "equipment.status",
        "plan.owner", "equipment.owner", "equipment.zone", "equipment.x", "equipment.z",
        "equipment.rotation", "equipment.controlTags", "equipment.requiredTraining",
        "document.equipmentId", "document.specVersion", "document.validFrom", "document.validUntil",
    ]
    targetId: str = Field(min_length=1, max_length=128)
    # Reservation values are strings, but operational fields can be numeric,
    # nullable strings, or bounded string lists.  Keeping the exact JSON value
    # in before/after avoids lossy string coercion while preserving a readable
    # field-level audit record.
    before: Any
    after: Any
    reason: str = Field(min_length=1, max_length=500)
    # Candidate edits are initially accepted for backwards compatibility with
    # saved branches.  A reviewer can explicitly reject one field without
    # applying the remaining branch edits; the decision remains part of the
    # immutable branch audit trail.
    disposition: Literal["accepted", "rejected"] = "accepted"
    reviewedBy: str | None = Field(default=None, max_length=180)
    reviewedAt: str | None = None
    reviewNote: str | None = Field(default=None, max_length=2000)

class BranchEquipmentState(StrictModel):
    """One equipment condition in an isolated candidate branch."""

    equipmentId: str = Field(min_length=1, max_length=128)
    status: Literal["available", "unavailable", "unknown"]


# Branch fields deliberately name a bounded set of business properties.  A
# generic JSON pointer would let a candidate mutate evidence, approval
# history, or unrelated workspace records without a reviewable contract.
BRANCH_PATCH_PATHS = frozenset({
    "reservation.equipmentId", "reservation.start", "reservation.end", "equipment.status",
    "plan.owner", "equipment.owner", "equipment.zone", "equipment.x", "equipment.z",
    "equipment.rotation", "equipment.controlTags", "equipment.requiredTraining",
    "document.equipmentId", "document.specVersion", "document.validFrom", "document.validUntil",
})

_BRANCH_PATCH_FIELDS: dict[str, tuple[str, str, str]] = {
    "reservation.equipmentId": ("reservation", "equipmentId", "text"),
    "reservation.start": ("reservation", "start", "text"),
    "reservation.end": ("reservation", "end", "text"),
    "equipment.status": ("equipment", "status", "status"),
    "plan.owner": ("plan", "owner", "text"),
    "equipment.owner": ("equipment", "owner", "text"),
    "equipment.zone": ("equipment", "zone", "nullable-text"),
    "equipment.x": ("equipment", "x", "coordinate"),
    "equipment.z": ("equipment", "z", "coordinate"),
    "equipment.rotation": ("equipment", "rotation", "rotation"),
    "equipment.controlTags": ("equipment", "controlTags", "tags"),
    "equipment.requiredTraining": ("equipment", "requiredTraining", "tags"),
    "document.equipmentId": ("document", "equipmentId", "nullable-text"),
    "document.specVersion": ("document", "specVersion", "positive-int"),
    "document.validFrom": ("document", "validFrom", "text"),
    "document.validUntil": ("document", "validUntil", "text"),
}


def _branch_patch_value(path: str, value: Any) -> Any:
    """Validate and return one field-level branch value.

    ``BranchPatch`` keeps values as JSON so list and numeric changes do not
    silently turn into strings.  This helper is the path-specific type gate
    used by save, review, candidate projection, and apply.
    """
    spec = _BRANCH_PATCH_FIELDS.get(path)
    if spec is None:
        raise ValueError("候选字段路径无效。")
    kind = spec[2]
    if kind == "text":
        if not isinstance(value, str) or not value.strip() or len(value) > 1000:
            raise ValueError(f"{path} 必须是非空文本。")
        return value
    if kind == "nullable-text":
        if value is not None and (not isinstance(value, str) or len(value) > 1000):
            raise ValueError(f"{path} 必须是文本或 null。")
        return value
    if kind == "status":
        if value not in {"available", "unavailable", "unknown"}:
            raise ValueError("设备状态候选值无效。")
        return value
    if kind in {"coordinate", "rotation"}:
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ValueError(f"{path} 必须是数字。")
        numeric = float(value)
        if kind == "coordinate" and not -20 <= numeric <= 20:
            raise ValueError(f"{path} 超出实验室坐标范围。")
        if kind == "rotation" and not -360 <= numeric <= 360:
            raise ValueError(f"{path} 超出朝向范围。")
        return value
    if kind == "positive-int":
        if isinstance(value, bool) or not isinstance(value, int) or value < 1:
            raise ValueError(f"{path} 必须是正整数。")
        return value
    if kind == "tags":
        if not isinstance(value, list) or len(value) > 40 or any(not isinstance(item, str) or not item.strip() or len(item) > 180 for item in value):
            raise ValueError(f"{path} 必须是最多 40 项的文本列表。")
        if len(value) != len(set(value)):
            raise ValueError(f"{path} 不能包含重复项。")
        return value
    raise ValueError("候选字段类型未配置。")


def _branch_patch_target(workspace: dict, plan_id: str, path: str, target_id: str) -> tuple[dict, str] | None:
    """Resolve a branch patch to one mutable object and field.

    The resolver is intentionally scoped to the branch's plan for plan and
    reservation fields.  Equipment and documents remain workspace objects but
    are still required to exist before a patch can be saved.
    """
    spec = _BRANCH_PATCH_FIELDS.get(path)
    if spec is None:
        return None
    collection, field, _ = spec
    if collection == "plan":
        plan = next((item for item in workspace.get("plans", []) if item["id"] == plan_id), None)
        return (plan, field) if plan is not None and target_id == plan_id else None
    if collection == "reservation":
        plan = next((item for item in workspace.get("plans", []) if item["id"] == plan_id), None)
        if plan is None:
            return None
        reservation = next((item for item in plan.get("reservations", []) if item["id"] == target_id), None)
        return (reservation, field) if reservation is not None else None
    if collection == "equipment":
        equipment = next((item for item in workspace.get("equipment", []) if item["id"] == target_id), None)
        return (equipment, field) if equipment is not None else None
    if collection == "document":
        document = next((item for item in workspace.get("documents", []) if item["id"] == target_id), None)
        return (document, field) if document is not None else None
    return None

class Plan(StrictModel):
    id: str
    code: str = Field(min_length=1, max_length=80)
    labId: str
    name: str = Field(min_length=1, max_length=180)
    owner: str = Field(min_length=1, max_length=180)
    revision: int = Field(ge=1, strict=True)
    source: str = Field(min_length=1, max_length=100000)
    sourceName: str = Field(max_length=240)
    steps: list[Step] = Field(max_length=500)
    reservations: list[Reservation] = Field(max_length=500)
    archived: bool
    # Explicit lifecycle is kept separate from the derived safety status.  Old
    # workspaces omit this field (``None`` is interpreted as draft, or archived
    # when archived=true); a transition operation writes it only after the
    # server rechecks the current plan and its evidence basis.
    lifecycle: Literal["draft", "submitted", "revising", "archived", "cancelled"] | None = None

class DocumentLocation(StrictModel):
    """A stable pointer into the extracted source, without claiming page OCR.

    ``page`` is populated for PDF sources, ``paragraph`` for DOCX body/table
    units, and ``lineStart``/``lineEnd`` for plain text.  The coordinates are
    copied from the extractor and let a client return to the exact extracted
    segment.  Missing coordinates mean that
    the extractor could not provide that coordinate; they are never inferred.
    """
    kind: Literal["page", "paragraph", "line", "range"]
    page: int | None = Field(default=None, ge=1, strict=True)
    paragraph: int | None = Field(default=None, ge=1, strict=True)
    lineStart: int | None = Field(default=None, ge=1, strict=True)
    lineEnd: int | None = Field(default=None, ge=1, strict=True)
    label: str | None = Field(default=None, max_length=240)

class Document(StrictModel):
    id: str
    name: str = Field(min_length=1, max_length=240)
    kind: Literal["inspection", "sop", "reference"]
    equipmentId: str | None
    specVersion: int = Field(ge=1, strict=True)
    status: Literal["pending", "verified", "rejected"]
    validFrom: str
    validUntil: str
    content: str = Field(min_length=1, max_length=200000)
    source: str = Field(min_length=1, max_length=1000)
    reviewedBy: str | None
    reviewedAt: str | None
    # Provenance fields are optional for backwards-compatible imports.  They
    # describe where a document came from and when its text was captured; they
    # do not assert that the source is authoritative or that the evidence is
    # applicable to every plan.  Applicability is evaluated separately below.
    sourceType: Literal["upload", "manual", "external", "system"] | None = None
    sourceLocator: str | None = Field(default=None, max_length=1000)
    capturedAt: str | None = None
    contentLength: int | None = Field(default=None, ge=0, strict=True)
    provenanceNote: str | None = Field(default=None, max_length=1000)
    sourceFileName: str | None = Field(default=None, max_length=240)
    sourceFileSize: int | None = Field(default=None, ge=0, strict=True)
    sourceMediaType: str | None = Field(default=None, max_length=180)
    sourceLocations: list[DocumentLocation] = Field(default_factory=list, max_length=2000)
    # Optional link to the controlled original-byte volume.  Older workspace
    # exports omit this field; text-only records remain valid and explicit.
    sourceArtifactId: str | None = Field(default=None, max_length=180)
    sourceArtifactVersion: int | None = Field(default=None, ge=1, strict=True)


def _source_span(
    *,
    document_id: str | None = None,
    source_artifact_id: str | None = None,
    source_artifact_version: int | None = None,
    locator: str | None = None,
    line: int | None = None,
    page: int | None = None,
    quoted_text: str | None = None,
    confirmed_by: str | None = None,
    confirmed_at: str | None = None,
) -> dict[str, Any]:
    """Build a sparse, JSON-safe source span without filling unknown facts.

    A span is intentionally sparse: a source artifact version is only useful
    when its artifact id is present, and confirmation metadata is only copied
    when both reviewer fields were recorded.  Callers can therefore attach
    ``sourceSpan: {}`` to every graph object while never inventing a location.
    """
    span: dict[str, Any] = {}
    if document_id:
        span["documentId"] = str(document_id)
    if source_artifact_id:
        span["sourceArtifactId"] = str(source_artifact_id)
        if source_artifact_version is not None:
            span["sourceArtifactVersion"] = int(source_artifact_version)
    if locator:
        span["locator"] = str(locator)
    if line is not None and int(line) > 0:
        span["line"] = int(line)
    if page is not None and int(page) > 0:
        span["page"] = int(page)
    if quoted_text is not None:
        span["quotedText"] = str(quoted_text)
    if confirmed_by and confirmed_at:
        span["confirmedBy"] = str(confirmed_by)
        span["confirmedAt"] = str(confirmed_at)
    return span


def _plan_source_span(plan: dict[str, Any], line: int | None) -> dict[str, Any]:
    """Return an exact line anchor into the saved plan revision, if valid."""
    if line is None:
        return {}
    try:
        line_number = int(line)
    except (TypeError, ValueError):
        return {}
    lines = str(plan.get("source") or "").splitlines()
    if line_number < 1 or line_number > len(lines):
        return {}
    # The plan revision is part of the locator so a stale Copilot result does
    # not silently point at a different saved source revision.
    locator = f"plan:{plan.get('id')}@revision:{plan.get('revision')}#line:{line_number}"
    return _source_span(locator=locator, line=line_number, quoted_text=lines[line_number - 1])


def _document_source_span(document: dict[str, Any], *, line: int | None = None, page: int | None = None) -> dict[str, Any]:
    """Resolve a document's recorded extraction coordinate without guessing."""
    locations = [item for item in (document.get("sourceLocations") or []) if isinstance(item, dict)]
    selected: dict[str, Any] | None = None
    if line is not None:
        try:
            line_number = int(line)
        except (TypeError, ValueError):
            return {}
        selected = next((item for item in locations if item.get("lineStart") is not None and item.get("lineEnd") is not None and int(item["lineStart"]) <= line_number <= int(item["lineEnd"])), None)
        if selected is None:
            return {}
    elif page is not None:
        try:
            page_number = int(page)
        except (TypeError, ValueError):
            return {}
        selected = next((item for item in locations if item.get("page") == page_number), None)
        if selected is None:
            return {}
    elif locations:
        selected = locations[0]

    line_number = line
    page_number = page
    quoted_text: str | None = None
    if selected:
        page_number = page_number if page_number is not None else selected.get("page")
        if line_number is None and selected.get("lineStart") == selected.get("lineEnd"):
            line_number = selected.get("lineStart")
        if line_number is not None:
            lines = str(document.get("content") or "").splitlines()
            if 1 <= int(line_number) <= len(lines):
                quoted_text = lines[int(line_number) - 1]
    return _source_span(
        document_id=document.get("id"),
        source_artifact_id=document.get("sourceArtifactId"),
        source_artifact_version=document.get("sourceArtifactVersion"),
        locator=document.get("sourceLocator"),
        line=line_number,
        page=page_number,
        quoted_text=quoted_text,
        confirmed_by=document.get("reviewedBy") if document.get("status") == "verified" else None,
        confirmed_at=document.get("reviewedAt") if document.get("status") == "verified" else None,
    )

class Task(StrictModel):
    id: str
    title: str = Field(min_length=1, max_length=240)
    planId: str | None
    equipmentId: str | None
    owner: str = Field(min_length=1, max_length=180)
    dueAt: str
    # ``returned`` and ``cancelled`` are persisted business states.  ``reopened``
    # is kept explicit so a completed/returned item can be reopened without
    # erasing the preceding decision from the audit trail.
    status: Literal["open", "submitted", "done", "returned", "cancelled", "reopened"]
    note: str = Field(max_length=5000)
    sourceId: str | None = Field(default=None, max_length=180)
    sourceLine: int | None = Field(default=None, ge=1)
    documentIds: list[str] = Field(default_factory=list, max_length=100)
    actionKind: str | None = Field(default=None, max_length=60)
    actionTarget: dict[str, Any] | None = None

class Decision(StrictModel):
    id: str
    planId: str
    revision: int
    basis: str
    actor: str
    comment: str = Field(min_length=1, max_length=2000)
    at: str
    # Decisions are immutable evidence of a human disposition.  Legacy
    # records omit these fields and remain interpreted as a normal approval.
    disposition: Literal["approved", "returned", "exception"] = "approved"
    actorRole: str = Field(default="safety_owner", max_length=80)
    stage: str = Field(default="safety_owner", max_length=80)
    checkSnapshot: list[dict[str, Any]] = Field(default_factory=list, max_length=500)
    exception: dict[str, Any] | None = None


def approval_stages() -> list[str]:
    """Return the server-configured approval stages without trusting the UI.

    The default remains the existing single safety-owner gate so old local
    workspaces keep their behavior.  Institutions can opt into a second EHS
    stage through an environment setting without changing the data contract.
    """
    raw = os.getenv("AEGIS_APPROVAL_STAGES", "safety_owner")
    stages = [item.strip() for item in raw.split(",") if item.strip()]
    return stages or ["safety_owner"]


def decision_check_snapshot(checks: list[dict]) -> list[dict]:
    return [
        {"id": str(check.get("id", "")), "state": str(check.get("state", "unknown")), "title": str(check.get("title", ""))[:240]}
        for check in checks
    ]


def approval_state(w: dict, p: dict, checks: list[dict]) -> dict:
    """Derive current approval state from immutable decisions and current basis."""
    decisions = [d for d in w.get("decisions", []) if d.get("planId") == p.get("id")]
    current_basis = basis(w, p)
    # Only the latest decision determines whether the current chain needs a
    # recheck. Older decisions remain immutable history; a new decision at
    # the current basis starts a fresh chain.
    stale = bool(decisions and decisions[-1].get("basis") != current_basis)
    latest = decisions[-1] if decisions else None
    stages = approval_stages()
    approved_stages = {item.get("stage", item.get("actorRole", "safety_owner")) for item in decisions if item.get("disposition", "approved") == "approved" and item.get("basis") == current_basis}
    pending = next((stage for stage in stages if stage not in approved_stages), None)
    disposition = latest.get("disposition", "approved") if latest else None
    return {
        "stages": stages,
        "latest": latest,
        "stale": stale,
        "pendingStage": pending,
        "disposition": disposition,
        "checkSnapshot": decision_check_snapshot(checks),
        "latestCheckSnapshot": latest.get("checkSnapshot", []) if latest else [],
    }

class Event(StrictModel):
    id: str
    at: str
    actor: str
    action: str
    objectId: str
    summary: str

class Branch(StrictModel):
    id: str
    name: str = Field(min_length=1, max_length=180)
    planId: str
    baseRevision: int
    baseBasis: str
    reservations: list[Reservation] = Field(max_length=500)
    createdAt: str
    patches: list[BranchPatch] = Field(default_factory=list, max_length=200)
    # Optional for old saved branches.  Only the selected condition is
    # projected into the candidate workspace; the formal equipment ledger is
    # unchanged until an explicit apply_branch operation succeeds.
    equipmentStates: list[BranchEquipmentState] = Field(default_factory=list, max_length=200)

class Workspace(StrictModel):
    schemaVersion: Literal[1]
    revision: int = Field(ge=0, strict=True)
    name: str = Field(min_length=1, max_length=180)
    provenance: Literal["user", "example"]
    laboratories: list[Lab] = Field(max_length=100)
    equipment: list[Equipment] = Field(max_length=2000)
    plans: list[Plan] = Field(max_length=2000)
    documents: list[Document] = Field(max_length=5000)
    tasks: list[Task] = Field(max_length=5000)
    decisions: list[Decision] = Field(max_length=5000)
    events: list[Event] = Field(max_length=5000)
    branches: list[Branch] = Field(max_length=2000)


def empty_workspace() -> dict:
    return dict(schemaVersion=1, revision=0, name="我的实验工作区", provenance="user", laboratories=[], equipment=[], plans=[], documents=[], tasks=[], decisions=[], events=[], branches=[])


CURRENT_WORKSPACE_SCHEMA = 1


def migrate_workspace_import(payload: dict) -> tuple[dict, dict[str, Any]]:
    """Normalize an exported workspace before validation and restore.

    Browser exports are long-lived files: a workspace written before optional
    provenance, lifecycle, task, or branch fields existed should remain
    importable after a server upgrade.  Migration is deliberately narrow and
    loss-averse.  It only unwraps a known export envelope, renames explicitly
    documented snake_case aliases, fills fields whose absence has an
    unambiguous empty/default meaning, and bumps the schema marker.  Unknown
    fields are left in place so ``Workspace.model_validate`` rejects them
    instead of silently discarding evidence.

    A future schema is rejected rather than guessed.  The report is returned
    to the API so a reviewer can see whether a restore was a no-op or used a
    legacy migration; it contains no document contents or credentials.
    """
    if not isinstance(payload, dict):
        raise ValueError("工作区存档必须是 JSON 对象。")
    raw = copy.deepcopy(payload)
    changes: list[str] = []
    envelope_unwrapped = False

    # Accept the exact object emitted by ``GET /export`` as well as the
    # nested workspace object used by the browser restore picker.  An actual
    # workspace has ``laboratories``; that guard prevents accidentally
    # unwrapping a malformed workspace which happens to contain a ``workspace``
    # metadata field.
    nested = raw.get("workspace")
    if "laboratories" not in raw and isinstance(nested, dict):
        raw = copy.deepcopy(nested)
        envelope_unwrapped = True
        changes.append("解包服务端导出信封")

    if "schemaVersion" in raw:
        schema = raw["schemaVersion"]
    elif "schema_version" in raw:
        schema = raw.pop("schema_version")
        changes.append("将 schema_version 迁移为 schemaVersion")
    else:
        # Very early browser drafts did not carry a schema marker.  Treating
        # those as v0 is safe because only additive defaults are applied.
        schema = 0
        changes.append("缺少 schemaVersion，按 v0 迁移")
    if isinstance(schema, bool) or not isinstance(schema, int) or schema < 0:
        raise ValueError("工作区存档版本必须是非负整数。")
    if schema > CURRENT_WORKSPACE_SCHEMA:
        raise ValueError(
            f"工作区存档版本 {schema} 高于服务端支持的 {CURRENT_WORKSPACE_SCHEMA}，请先升级服务端。"
        )

    # These fields have a semantically safe default.  Missing lists mean no
    # records were present in the old export; missing revision/provenance are
    # metadata defaults and do not claim any laboratory fact.
    defaults: dict[str, Any] = {
        "revision": 0,
        "name": "导入的实验工作区",
        "provenance": "user",
        "laboratories": [],
        "equipment": [],
        "plans": [],
        "documents": [],
        "tasks": [],
        "decisions": [],
        "events": [],
        "branches": [],
    }
    for key, value in defaults.items():
        if key not in raw:
            raw[key] = copy.deepcopy(value)
            changes.append(f"补齐缺省字段 {key}")

    # Document provenance was introduced with camelCase names.  Migrate only
    # explicit aliases; do not infer source authority or invent timestamps.
    document_aliases = {
        "equipment_id": "equipmentId", "spec_version": "specVersion",
        "valid_from": "validFrom", "valid_until": "validUntil",
        "reviewed_by": "reviewedBy", "reviewed_at": "reviewedAt",
        "source_type": "sourceType", "source_locator": "sourceLocator",
        "captured_at": "capturedAt", "content_length": "contentLength",
        "provenance_note": "provenanceNote", "source_file_name": "sourceFileName",
        "source_file_size": "sourceFileSize", "source_media_type": "sourceMediaType",
        "source_locations": "sourceLocations", "source_artifact_id": "sourceArtifactId",
        "source_artifact_version": "sourceArtifactVersion",
    }
    for document in raw.get("documents", []):
        if not isinstance(document, dict):
            continue
        for old, new in document_aliases.items():
            if new not in document and old in document:
                document[new] = document.pop(old)
                changes.append(f"资料字段 {old} → {new}")

    # Legacy plans only had ``archived``.  The field remains optional in the
    # current contract, so there is no need to invent a lifecycle value.  A
    # missing archived flag is safe to interpret as an active plan.
    for plan in raw.get("plans", []):
        if isinstance(plan, dict) and "archived" not in plan:
            plan["archived"] = False
            changes.append("计划缺少 archived，按未归档迁移")

    raw["schemaVersion"] = CURRENT_WORKSPACE_SCHEMA
    if schema != CURRENT_WORKSPACE_SCHEMA:
        changes.append(f"工作区 schema v{schema} → v{CURRENT_WORKSPACE_SCHEMA}")
    report = {
        "fromVersion": schema,
        "toVersion": CURRENT_WORKSPACE_SCHEMA,
        "migrated": bool(changes),
        "changes": changes,
        "envelopeUnwrapped": envelope_unwrapped,
    }
    return raw, report

def read_source_registry() -> list[dict]:
    """Read the public source directory without treating it as workspace fact."""
    try:
        raw = json.loads(SOURCE_REGISTRY_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise HTTPException(503, "公开来源目录暂时不可用。") from error
    if not isinstance(raw, list):
        raise HTTPException(503, "公开来源目录格式无效。")
    return raw


def validate_workspace(w: dict) -> dict:
    # Validate without normalizing numeric JSON types: stable TS/Python fingerprints.
    Workspace.model_validate(w)
    for key in ("laboratories", "equipment", "plans", "documents", "tasks", "decisions", "events", "branches"):
        ids = [x["id"] for x in w[key]]
        if len(ids) != len(set(ids)) or any(not isinstance(x, str) or not x or len(x) > 128 for x in ids):
            raise ValueError("对象 ID 必须唯一且有效。")
    labs = {x["id"]: x for x in w["laboratories"]}
    equipment = {x["id"]: x for x in w["equipment"]}
    plans = {x["id"]: x for x in w["plans"]}
    def interval(a: str, b: str):
        if dt(a) >= dt(b):
            raise ValueError("结束时间必须晚于开始时间。")
    codes: set[tuple] = set()
    for e in equipment.values():
        if e["labId"] not in labs:
            raise ValueError("资源所属实验室不存在。")
        l = labs[e["labId"]]
        if abs(e["x"]) > l["width"] / 2 - .5 or abs(e["z"]) > l["depth"] / 2 - .5:
            raise ValueError("设备位置超出实验室边界。")
        code = (e["labId"], e["code"])
        if code in codes:
            raise ValueError("同一实验室内设备编号不能重复。")
        codes.add(code)
        for m in e["maintenance"]:
            interval(m["start"], m["end"])
    for p in plans.values():
        if p["labId"] not in labs or not p["name"].strip() or not p["owner"].strip() or not p["source"].strip():
            raise ValueError("计划基础资料不完整。")
        lines = p["source"].splitlines()
        for s in p["steps"]:
            if s["line"] > len(lines) or lines[s["line"] - 1].strip() != s["text"] or any(equipment.get(i, {}).get("labId") != p["labId"] for i in s["equipmentIds"]):
                raise ValueError("步骤必须准确引用原文，不得编造行号。")
        for i, r in enumerate(p["reservations"]):
            interval(r["start"], r["end"])
            if equipment.get(r["equipmentId"], {}).get("labId") != p["labId"]:
                raise ValueError("预约设备必须属于当前实验室。")
            if any(other["equipmentId"] == r["equipmentId"] and overlap(other, r) for other in p["reservations"][i + 1:]):
                raise ValueError("同一计划不能重复占用同一资源的重叠时段。")
        if p.get("archived") and p.get("lifecycle", "archived") not in {"archived"}:
            raise ValueError("已归档计划必须保持 archived 生命周期。")
    for d in w["documents"]:
        if d["equipmentId"] is not None and d["equipmentId"] not in equipment:
            raise ValueError("资料关联的设备不存在。")
        interval(d["validFrom"], d["validUntil"])
        if not d["content"].strip() or not d["source"].strip():
            raise ValueError("资料内容和来源为必填项。")
        if d.get("capturedAt") is not None:
            # A provenance timestamp is a claim about when the captured text
            # entered the workspace; reject ambiguous local timestamps.
            dt(d["capturedAt"])
        if d.get("contentLength") is not None and d["contentLength"] != len(d["content"]):
            raise ValueError("资料内容长度与 provenance 记录不一致。")
        if d.get("sourceFileName") is not None and not str(d["sourceFileName"]).strip():
            raise ValueError("资料原始文件名不能是空白。")
        if d.get("sourceFileSize") is not None and d.get("sourceType") != "upload":
            raise ValueError("原始文件大小只能用于上传资料。")
        for location in d.get("sourceLocations", []):
            coordinates = [location.get("page"), location.get("paragraph"), location.get("lineStart"), location.get("lineEnd")]
            if not any(value is not None for value in coordinates):
                raise ValueError("资料定位至少要包含页、段落或行坐标。")
            if location.get("lineStart") is not None and location.get("lineEnd") is not None and location["lineStart"] > location["lineEnd"]:
                raise ValueError("资料行定位范围无效。")
    for t in w["tasks"]:
        dt(t["dueAt"])
        if (t["planId"] is not None and t["planId"] not in plans) or (t["equipmentId"] is not None and t["equipmentId"] not in equipment):
            raise ValueError("任务关联的对象不存在。")
    for b in w["branches"]:
        if b["planId"] not in plans:
            raise ValueError("方案关联的计划不存在。")
    if len(json.dumps(w, ensure_ascii=False)) > 5_000_000:
        raise ValueError("工作区超过当前容量，请先归档。")
    return w


def equipment_footprint(kind: str) -> tuple[float, float]:
    """Return a conservative placement footprint for a resource kind.

    The footprint is only used for collision-free default placement.  It is
    intentionally kept outside ``Equipment`` so changing a visual placement
    cannot alter a resource's capability or review basis.
    """
    return EQUIPMENT_FOOTPRINTS.get(kind, (1.5, 1.2))


def equipment_overlaps(a: dict, b: dict, *, x: float | None = None, z: float | None = None) -> bool:
    """Whether two resources would occupy the same scene slot.

    Rotation is ignored for the conservative default layout.  The scene can
    still rotate a resource after it has been created; this check only prevents
    the common failure where every new resource is saved at ``(0, 0)``.
    """
    # Resources in different laboratories have independent coordinate spaces.
    # Comparing their scene coordinates would make a new object in one room
    # appear to collide with an object in another room.
    if a.get("labId") != b.get("labId"):
        return False
    ax = a["x"] if x is None else x
    az = a["z"] if z is None else z
    aw, ad = equipment_footprint(a["kind"])
    bw, bd = equipment_footprint(b["kind"])
    return (
        abs(ax - b["x"]) < (aw + bw) / 2 + PLACEMENT_GAP
        and abs(az - b["z"]) < (ad + bd) / 2 + PLACEMENT_GAP
    )


def find_equipment_position(w: dict, equipment: dict) -> tuple[float, float]:
    """Find a deterministic free position inside the owning laboratory.

    The center is preferred for the first resource.  Subsequent candidates are
    sampled on a stable grid from the back-left toward the front-right so a
    retry produces the same workspace.  An explicit position is never changed
    by this helper; callers opt into it only for newly-created resources that
    requested automatic placement.
    """
    lab = next((item for item in w["laboratories"] if item["id"] == equipment["labId"]), None)
    if lab is None:
        raise ValueError("资源所属实验室不存在。")
    width, depth = equipment_footprint(equipment["kind"])
    half_w, half_d = width / 2, depth / 2
    x_min, x_max = -lab["width"] / 2 + half_w, lab["width"] / 2 - half_w
    z_min, z_max = -lab["depth"] / 2 + half_d, lab["depth"] / 2 - half_d
    if x_min > x_max or z_min > z_max:
        raise ValueError("实验室空间不足，无法自动放置该资源。")

    existing = [item for item in w["equipment"] if item["labId"] == equipment["labId"] and item["id"] != equipment["id"]]
    # Start at the visual center, then use a fixed lattice.  Rounding avoids
    # accumulating floating-point noise in exported workspace JSON.
    candidates: list[tuple[float, float]] = [(0.0, 0.0)]
    step_x = max(1.0, min(width + PLACEMENT_GAP, 2.4))
    step_z = max(1.0, min(depth + PLACEMENT_GAP, 2.0))
    x = x_min
    while x <= x_max + 1e-6:
        z = z_min
        while z <= z_max + 1e-6:
            candidates.append((round(x, 2), round(z, 2)))
            z += step_z
        x += step_x
    # Prefer candidates closest to the center after the explicit center entry.
    candidates[1:] = sorted(candidates[1:], key=lambda point: (point[0] ** 2 + point[1] ** 2, point[1], point[0]))
    for candidate_x, candidate_z in candidates:
        if not (x_min - 1e-6 <= candidate_x <= x_max + 1e-6 and z_min - 1e-6 <= candidate_z <= z_max + 1e-6):
            continue
        if not any(equipment_overlaps(equipment, other, x=candidate_x, z=candidate_z) for other in existing):
            return candidate_x, candidate_z
    raise ValueError("实验室没有可用的自动放置位置，请扩大空间或调整已有资源。")


def basis(w: dict, p: dict) -> str:
    """Fingerprint only facts that can change this plan's checks.

    Display edits, camera placement, unrelated future maintenance and distant
    reservations must not revoke a recorded decision.  The fingerprint still
    changes for a source/step/booking edit, a relevant equipment state or
    specification change, and evidence whose applicability window intersects
    the plan's actual reservations.
    """
    reservations = sorted((dict(r) for r in p.get("reservations", [])), key=lambda r: (r.get("equipmentId", ""), r.get("start", ""), r.get("end", ""), r.get("id", "")))
    ids = {r["equipmentId"] for r in reservations}
    def relevant(interval: dict, windows: list[dict]) -> bool:
        return any(overlap(interval, window) for window in windows)
    relevant_equipment = []
    for e in w["equipment"]:
        if e["id"] not in ids:
            continue
        e_reservations = [r for r in reservations if r["equipmentId"] == e["id"]]
        maintenance = sorted((m for m in e.get("maintenance", []) if relevant(m, e_reservations)), key=lambda m: (m.get("start", ""), m.get("end", ""), m.get("id", "")))
        relevant_equipment.append({
            "id": e["id"], "labId": e["labId"], "kind": e["kind"],
            "status": e["status"], "capacity": e["capacity"],
            "requiresEvidence": e["requiresEvidence"], "specVersion": e["specVersion"],
            "maintenance": maintenance,
        })
    relevant_equipment.sort(key=lambda e: e["id"])
    relevant_documents = []
    for d in w["documents"]:
        if d.get("equipmentId") not in ids:
            continue
        doc_window = {"start": d["validFrom"], "end": d["validUntil"]}
        if not relevant(doc_window, reservations) and d.get("kind") == "inspection":
            continue
        document_facts = {k: d.get(k) for k in ("id", "equipmentId", "kind", "specVersion", "status", "validFrom", "validUntil", "content", "source", "reviewedBy", "reviewedAt")}
        # Keep legacy fingerprints byte-for-byte stable when an imported
        # document has no provenance metadata.  New records include only the
        # fields that were actually captured, so absent metadata is not a
        # fabricated claim and the browser/Python parity contract survives.
        document_facts.update({k: d[k] for k in ("sourceType", "sourceLocator", "capturedAt", "contentLength", "provenanceNote", "sourceFileName", "sourceFileSize", "sourceMediaType", "sourceLocations") if d.get(k) is not None})
        relevant_documents.append(document_facts)
    relevant_documents.sort(key=lambda d: (d.get("equipmentId") or "", d.get("kind", ""), d.get("id", "")))
    other = []
    for q in w["plans"]:
        if q["id"] == p["id"] or q["archived"] or q.get("lifecycle") == "cancelled":
            continue
        for r in q.get("reservations", []):
            if r["equipmentId"] in ids and relevant(r, reservations):
                other.append(dict(planId=q["id"], **r))
    other.sort(key=lambda r: (r.get("equipmentId", ""), r.get("start", ""), r.get("end", ""), r.get("planId", ""), r.get("id", "")))
    plan_facts = {
        "id": p["id"], "labId": p["labId"], "revision": p.get("revision"), "sourceName": p.get("sourceName"), "source": p["source"],
        "steps": [{k: s.get(k) for k in ("id", "text", "line", "equipmentIds", "confirmed")} for s in sorted(p.get("steps", []), key=lambda s: (s.get("line", 0), s.get("id", "")))],
        "reservations": reservations,
    }
    return canonical(dict(rulePack=RULE_PACK, plan=plan_facts, equipment=relevant_equipment,
        documents=relevant_documents, otherReservations=other))


def _effective_branch_snapshots(w: dict, branch: dict) -> tuple[list[dict], list[dict]]:
    """Project only accepted field patches onto the current branch base.

    This is used after a reviewer rejects an individual suggestion.  The
    legacy branch shape still carries a complete snapshot; once a patch has an
    explicit disposition, the accepted patches become the authoritative
    candidate and rejected fields fall back to the base value.
    """
    base_plan = next((item for item in w["plans"] if item["id"] == branch["planId"]), None)
    if base_plan is None:
        raise ValueError("方案关联的计划不存在。")
    reservations = copy.deepcopy(base_plan.get("reservations", []))
    equipment_by_id = {item["id"]: copy.deepcopy(item) for item in w.get("equipment", [])}
    reservation_by_id = {item["id"]: item for item in reservations}
    for patch in branch.get("patches", []):
        if patch.get("disposition", "accepted") != "accepted":
            continue
        path = patch["path"]
        if path == "equipment.status":
            equipment = equipment_by_id.get(patch["targetId"])
            if equipment is not None:
                equipment["status"] = patch["after"]
            continue
        row = reservation_by_id.get(patch["targetId"])
        if row is not None:
            row[path.removeprefix("reservation.")] = patch["after"]
    states = [
        {"equipmentId": patch["targetId"], "status": patch["after"]}
        for patch in branch.get("patches", [])
        if patch["path"] == "equipment.status" and patch.get("disposition", "accepted") == "accepted"
    ]
    return reservations, states


def branch_candidate(w: dict, branch: dict) -> dict:
    """Build and validate a branch projection without mutating ``w``.

    ``reservations`` is the candidate snapshot retained for compatibility
    with existing branches.  New branches additionally carry field-level
    ``patches``.  Each patch is checked against the current base value and
    against the candidate snapshot, which prevents a stale or misleading diff
    from being stored as if it described the selected branch.
    """
    Branch.model_validate(branch)
    candidate = copy.deepcopy(w)
    plan = next((item for item in candidate["plans"] if item["id"] == branch["planId"]), None)
    base_plan = next((item for item in w["plans"] if item["id"] == branch["planId"]), None)
    if plan is None or base_plan is None:
        raise ValueError("方案关联的计划不存在。")
    patch_keys = [(item.get("path"), item.get("targetId")) for item in branch.get("patches", [])]
    if len(patch_keys) != len(set(patch_keys)):
        raise ValueError("候选字段不能重复指定同一路径和目标。")
    has_review_decisions = any("disposition" in patch for patch in branch.get("patches", []))
    if has_review_decisions:
        effective_reservations, effective_states = _effective_branch_snapshots(w, branch)
        if branch.get("reservations") != effective_reservations:
            raise ValueError("候选字段与人工取舍后的方案快照不一致，请重新生成方案。")
        if branch.get("equipmentStates", []) != effective_states:
            raise ValueError("候选设备状态与人工取舍后的审计字段不一致。")
        plan["reservations"] = effective_reservations
    else:
        plan["reservations"] = copy.deepcopy(branch["reservations"])
    equipment_by_id = {item["id"]: item for item in candidate["equipment"]}
    base_equipment_by_id = {item["id"]: item for item in w["equipment"]}
    equipment_states = branch.get("equipmentStates", [])
    if len({item["equipmentId"] for item in equipment_states}) != len(equipment_states):
        raise ValueError("候选设备状态不能重复指定同一资源。")
    for state in equipment_states:
        equipment = equipment_by_id.get(state["equipmentId"])
        if equipment is None:
            raise ValueError(f"候选字段目标不存在：{state['equipmentId']}。")
        equipment["status"] = state["status"]
    state_by_equipment = {item["equipmentId"]: item["status"] for item in equipment_states}
    reservation_by_id = {item["id"]: item for item in plan["reservations"]}
    base_reservation_by_id = {item["id"]: item for item in base_plan["reservations"]}
    for patch in branch.get("patches", []):
        path = patch["path"]
        target_id = patch["targetId"]
        if path not in BRANCH_PATCH_PATHS:
            raise ValueError("候选字段路径无效。")
        before = _branch_patch_value(path, patch.get("before"))
        after = _branch_patch_value(path, patch.get("after"))
        if before == after:
            raise ValueError("候选字段必须产生实际变更。")
        if patch.get("disposition", "accepted") not in {"accepted", "rejected"}:
            raise ValueError("候选字段取舍必须是 accepted 或 rejected。")
        if path == "equipment.status":
            base_equipment = base_equipment_by_id.get(target_id)
            candidate_equipment = equipment_by_id.get(target_id)
            if base_equipment is None or candidate_equipment is None:
                raise ValueError(f"候选字段目标不存在：{target_id}。")
            if base_equipment["status"] != before:
                raise ValueError("候选字段基线已变化，请重新创建方案。")
            expected_after = after if patch.get("disposition", "accepted") == "accepted" else before
            if candidate_equipment["status"] != expected_after:
                raise ValueError("候选字段与方案快照不一致，请重新生成方案。")
            if state_by_equipment.get(target_id) != expected_after:
                # Rejected fields intentionally disappear from the effective
                # equipmentStates snapshot; accepted fields must stay aligned.
                if patch.get("disposition", "accepted") == "accepted" or target_id in state_by_equipment:
                    raise ValueError("设备状态候选必须与审计字段保持一致。")
            continue
        # Reservation snapshots are supplied explicitly by the branch and
        # therefore must match the patch.  Other business objects are
        # projected directly from their field-level patch; no formal object
        # is mutated until apply_branch succeeds.
        if path.startswith("reservation."):
            field = path.removeprefix("reservation.")
            base_row = base_reservation_by_id.get(target_id)
            candidate_row = reservation_by_id.get(target_id)
            if base_row is None or candidate_row is None:
                raise ValueError(f"候选字段目标不存在：{target_id}。")
            if base_row[field] != before:
                raise ValueError("候选字段基线已变化，请重新创建方案。")
            expected_after = after if patch.get("disposition", "accepted") == "accepted" else before
            if candidate_row[field] != expected_after:
                raise ValueError("候选字段与方案快照不一致，请重新生成方案。")
            continue
        base_target = _branch_patch_target(w, branch["planId"], path, target_id)
        candidate_target = _branch_patch_target(candidate, branch["planId"], path, target_id)
        if base_target is None or candidate_target is None:
            raise ValueError(f"候选字段目标不存在：{target_id}。")
        base_object, field = base_target
        candidate_object, candidate_field = candidate_target
        if field != candidate_field or base_object.get(field) != before:
            raise ValueError("候选字段基线已变化，请重新创建方案。")
        expected_after = after if patch.get("disposition", "accepted") == "accepted" else before
        candidate_object[field] = copy.deepcopy(expected_after)
    audited_equipment_after = {
        patch["targetId"]: patch["after"] for patch in branch.get("patches", [])
        if patch.get("path") == "equipment.status" and patch.get("disposition", "accepted") == "accepted"
    }
    for state in equipment_states:
        base_equipment = base_equipment_by_id[state["equipmentId"]]
        if base_equipment["status"] != state["status"] and audited_equipment_after.get(state["equipmentId"]) != state["status"]:
            raise ValueError("设备状态候选必须与审计字段保持一致。")
    validate_workspace(candidate)
    return candidate


def _apply_accepted_branch_fields(workspace: dict, branch: dict) -> None:
    """Apply only accepted non-snapshot fields to the formal workspace.

    Reservations and equipment status retain their existing explicit apply
    logic because both have compatibility snapshots.  Every other field is
    copied through the same path/type resolver used during candidate checks;
    rejected suggestions remain audit history and never reach production
    state.
    """
    for patch in branch.get("patches", []):
        if patch.get("disposition", "accepted") != "accepted":
            continue
        path = patch["path"]
        if path.startswith("reservation.") or path == "equipment.status":
            continue
        target = _branch_patch_target(workspace, branch["planId"], path, patch["targetId"])
        if target is None:
            raise ValueError(f"候选字段目标不存在：{patch['targetId']}。")
        obj, field = target
        obj[field] = copy.deepcopy(_branch_patch_value(path, patch.get("after")))


def check_workspace(w: dict) -> list[dict]:
    results: list[dict] = []
    for p in [p for p in w["plans"] if not p["archived"] and p.get("lifecycle") != "cancelled"]:
        def add(code, state, title, detail, equipment_id=None, related=None, documents=None):
            line = next((s["line"] for s in p["steps"] if equipment_id in s["equipmentIds"]), None)
            results.append(dict(id=f"{p['id']}:{code}:{equipment_id or 'plan'}", planId=p["id"], equipmentId=equipment_id, code=code, state=state, title=title, detail=detail, relatedPlanIds=related or [], documentIds=documents or [], line=line))
        confirmed = bool(p["steps"]) and all(s["confirmed"] for s in p["steps"])
        add("SOURCE", "pass" if confirmed else "unknown", "原文与步骤已确认" if confirmed else "步骤提取需要人工确认", f"{len(p['steps'])} 个步骤绑定到原文行号。" if confirmed else "请核对原文、步骤及设备匹配；机器提取不代表事实已确认。")
        if not p["reservations"]:
            add("SCHEDULE", "unknown", "尚未设置资源使用时段", "设备、开始和结束时间是资源核验的必要输入。")
        ids = list(dict.fromkeys(r["equipmentId"] for r in p["reservations"]))
        for identifier in dict.fromkeys(i for step in p["steps"] for i in step["equipmentIds"]):
            if identifier not in ids:
                add("BINDING", "unknown", "原文设备与预约尚未对齐", "步骤引用的设备没有对应预约。更换资源后须修订原文并重新确认，不能只替换画面中的连线。", identifier)
        for identifier in ids:
            e = next((e for e in w["equipment"] if e["id"] == identifier), None)
            if not e:
                add("RESOURCE", "unknown", "资源不存在", "重新绑定有效资源。", identifier)
                continue
            reservations = [r for r in p["reservations"] if r["equipmentId"] == identifier]
            maintenance = e["status"] == "unavailable" or any(overlap(r, m) for r in reservations for m in e["maintenance"])
            state = "blocked" if maintenance else "unknown" if e["status"] == "unknown" else "pass"
            title = f"{e['code']} 在预约范围内不可用" if maintenance else f"{e['code']} 的台账状态待确认" if state == "unknown" else f"{e['code']} 的预约不冲突于维护"
            add("AVAILABILITY", state, title, "当前停用状态或维护窗口与使用计划交叠；需更换资源或调整时段。" if maintenance else "仅依据已登记的台账与维护窗口；不代表现场传感器状态。", identifier)
            assigned = [dict(planId=q["id"], **r) for q in w["plans"] if not q["archived"] and q.get("lifecycle") != "cancelled" for r in q["reservations"] if r["equipmentId"] == identifier]
            conflicts: list[str] = []
            for r in reservations:
                active = [a for a in assigned if overlap(a, r)]
                points = {dt(r["start"])} | {max(dt(a["start"]), dt(r["start"])) for a in active}
                for point in sorted(points):
                    at = [a for a in active if dt(a["start"]) <= point < dt(a["end"])]
                    if len(at) > e["capacity"]:
                        for a in at:
                            if a["planId"] != p["id"] and a["planId"] not in conflicts:
                                conflicts.append(a["planId"])
            add("CAPACITY", "blocked" if conflicts else "pass", f"{e['code']} 超出并发使用容量" if conflicts else f"{e['code']} 的资源容量已满足", f"登记容量为 {e['capacity']}；与 {'、'.join(conflicts)} 同时占用。" if conflicts else f"当前使用安排未超出登记容量 {e['capacity']}。相邻时段按半开区间计算。", identifier, conflicts)
            if e["requiresEvidence"]:
                docs = [d for d in w["documents"] if d["equipmentId"] == identifier and d["kind"] == "inspection"]
                valid = [d for d in docs if d["status"] == "verified" and d["specVersion"] == e["specVersion"]]
                # Two independently verified records for the same object and
                # specification can overlap the requested interval while
                # describing different observations.  Treat that as an
                # unresolved evidence conflict instead of arbitrarily picking
                # the first document.  A document is still listed with its
                # normal provenance so a reviewer can reconcile the pair.
                conflict_by_reservation: dict[str, list[str]] = {}
                for reservation in reservations:
                    matches = [
                        d for d in valid
                        if dt(d["validFrom"]) <= dt(reservation["start"])
                        and dt(d["validUntil"]) >= dt(reservation["end"])
                    ]
                    signatures = {canonical({"content": d.get("content", "")}) for d in matches}
                    if len(signatures) > 1:
                        conflict_by_reservation[reservation["id"]] = [d["id"] for d in matches]
                covered = all(any(dt(d["validFrom"]) <= dt(r["start"]) and dt(d["validUntil"]) >= dt(r["end"]) for d in valid) for r in reservations)
                conflicted_reservation_ids = set(conflict_by_reservation)
                if conflicted_reservation_ids:
                    covered = False
                applicability = []
                for d in docs:
                    covered_reservation_ids = [
                        r["id"] for r in reservations
                        if dt(d["validFrom"]) <= dt(r["start"]) and dt(d["validUntil"]) >= dt(r["end"])
                    ]
                    missing_reservation_ids = [r["id"] for r in reservations if r["id"] not in covered_reservation_ids]
                    if d["status"] != "verified":
                        state, reason = "candidate", "not-verified"
                    elif d["specVersion"] != e["specVersion"]:
                        state, reason = "mismatched", "spec-version-mismatch"
                    elif not covered_reservation_ids:
                        state, reason = "expired", "outside-reservation-window"
                    elif any(reservation_id in conflict_by_reservation and d["id"] in conflict_by_reservation[reservation_id] for reservation_id in covered_reservation_ids):
                        state, reason = "conflict", "overlapping-verified-conflict"
                    else:
                        state = "supports"
                        reason = "covers-all-reservations" if not missing_reservation_ids else "partial-reservation-coverage"
                    applicability_row = {
                        "documentId": d["id"], "state": state, "reason": reason,
                        "coverage": "complete" if not missing_reservation_ids and covered_reservation_ids else "partial" if covered_reservation_ids else "none",
                        "coverageBasis": "validity-window",
                        "effectiveForAllReservations": state == "supports" and not missing_reservation_ids,
                        "coveredReservationIds": covered_reservation_ids,
                        "missingReservationIds": missing_reservation_ids,
                        "equipmentId": identifier,
                        "specVersion": d["specVersion"], "expectedSpecVersion": e["specVersion"],
                        "validFrom": d["validFrom"], "validUntil": d["validUntil"],
                        "source": d["source"], "sourceType": d.get("sourceType"),
                        "sourceLocator": d.get("sourceLocator"), "capturedAt": d.get("capturedAt"),
                        "contentLength": d.get("contentLength"), "sourceFileName": d.get("sourceFileName"),
                        "sourceFileSize": d.get("sourceFileSize"), "sourceMediaType": d.get("sourceMediaType"),
                        "sourceLocations": d.get("sourceLocations", []),
                        "provenanceNote": d.get("provenanceNote"),
                        "reviewedBy": d.get("reviewedBy"), "reviewedAt": d.get("reviewedAt"),
                    }
                    conflicting_document_ids = sorted({
                        document_id
                        for reservation_id in covered_reservation_ids
                        for document_id in conflict_by_reservation.get(reservation_id, [])
                        if document_id != d["id"]
                    })
                    if conflicting_document_ids:
                        applicability_row["conflictingDocumentIds"] = conflicting_document_ids
                    applicability.append(applicability_row)
                evidence_detail = "对象、规格版本、核验状态与全部使用时段均匹配。" if covered else "需要已核验且覆盖整个使用时段的记录；上传、过期、其他设备的记录都不能代替。"
                if conflicted_reservation_ids:
                    conflict_documents = sorted({document_id for ids_for_reservation in conflict_by_reservation.values() for document_id in ids_for_reservation})
                    evidence_detail = f"同一预约时段存在 {len(conflict_documents)} 份内容不一致但均已核验的记录（{'、'.join(conflict_documents)}）；需负责人裁定后才能作为适用证据。"
                add("EVIDENCE", "pass" if covered else "unknown", f"{e['code']} 的检查证据适用" if covered else f"{e['code']} 缺少适用的检查证据", evidence_detail, identifier, documents=[d["id"] for d in docs])
                results[-1]["evidenceApplicability"] = applicability
                results[-1]["evidenceScope"] = {
                    "equipmentId": identifier,
                    "specVersion": e["specVersion"],
                    "reservationIds": [r["id"] for r in reservations],
                    "requiredCoverage": "all-reservations",
                }
                if conflicted_reservation_ids:
                    results[-1]["evidenceScope"]["conflictReservationIds"] = sorted(conflicted_reservation_ids)
    return results


def plan_status(w: dict, p: dict, checks: list[dict]) -> str:
    if p["archived"] or p.get("lifecycle") == "archived":
        return "archived"
    if p.get("lifecycle") == "cancelled":
        return "cancelled"
    approval = approval_state(w, p, [c for c in checks if c["planId"] == p["id"]])
    decision = approval["latest"]
    if approval["stale"]:
        return "recheck"
    if decision and approval["disposition"] == "returned":
        return "returned"
    if decision and approval["disposition"] == "exception" and not approval["stale"]:
        return "approved_with_exception"
    if not p["steps"] or any(not s["confirmed"] for s in p["steps"]):
        return "draft"
    relevant = [c for c in checks if c["planId"] == p["id"]]
    if any(c["state"] == "blocked" for c in relevant):
        return "blocked"
    if any(c["state"] == "unknown" for c in relevant):
        return "unknown"
    if decision and approval["pendingStage"] is None:
        return "approved"
    return "pending_approval" if decision else "ready"


def envelope(w: dict) -> dict:
    checks = check_workspace(w)
    return dict(workspace=w, checks=checks, statuses={p["id"]: plan_status(w, p, checks) for p in w["plans"]}, mode="server", rulePack=RULE_PACK, scope="资源时段、台账可用性、检查证据适用性与原文确认；不构成全域化学安全评估或开工许可。")


def expiry_recheck_report(w: dict, as_of: str, *, worker_name: str = "not-running") -> dict:
    """Return a reproducible expiry/overdue report for the current snapshot.

    This is deliberately a pure report builder.  It does not mark evidence
    verified, alter a plan, or create work items.  The durable endpoint stores
    the report in ``spatial_jobs`` so a later browser request can inspect the
    same result instead of silently treating a fresh calculation as a worker
    run.  Plan impact is derived from the existing applicability checks, so an
    expired record only affects a plan when its object, version and reservation
    window make that record relevant.
    """
    as_of_dt = dt(as_of)
    expired_evidence: list[dict] = []
    expired_ids: set[str] = set()
    for document in w.get("documents", []):
        until = dt(document["validUntil"])
        if document.get("status") == "verified" and until <= as_of_dt:
            expired_ids.add(document["id"])
            expired_evidence.append({
                "documentId": document["id"],
                "name": document["name"],
                "equipmentId": document.get("equipmentId"),
                "specVersion": document.get("specVersion"),
                "validUntil": document["validUntil"],
                "source": document.get("source"),
                "status": document.get("status"),
                "reason": "valid-until-reached",
            })

    active_task_states = {"open", "submitted", "reopened"}
    overdue_tasks = [
        {
            "taskId": task["id"],
            "title": task["title"],
            "planId": task.get("planId"),
            "equipmentId": task.get("equipmentId"),
            "owner": task.get("owner"),
            "dueAt": task["dueAt"],
            "status": task["status"],
            "reason": "due-at-reached",
        }
        for task in w.get("tasks", [])
        if task.get("status") in active_task_states and dt(task["dueAt"]) <= as_of_dt
    ]

    checks = check_workspace(w)
    impacted_plans: list[dict] = []
    for plan in w.get("plans", []):
        if plan.get("archived") or plan.get("lifecycle") == "cancelled":
            continue
        plan_checks = [check for check in checks if check["planId"] == plan["id"]]
        expired_check_ids: list[str] = []
        expired_plan_document_ids: set[str] = set()
        for check in plan_checks:
            if check.get("code") != "EVIDENCE":
                continue
            for applicability in check.get("evidenceApplicability", []):
                document_id = applicability.get("documentId")
                if document_id in expired_ids and applicability.get("state") in {"expired", "partial", "mismatched", "candidate"}:
                    expired_check_ids.append(check["id"])
                    expired_plan_document_ids.add(document_id)
                    break
        decisions = [decision for decision in w.get("decisions", []) if decision.get("planId") == plan["id"]]
        decision_invalidated = bool(decisions and decisions[-1].get("basis") != basis(w, plan))
        if not expired_check_ids and not decision_invalidated:
            continue
        reasons: list[str] = []
        if expired_check_ids:
            reasons.append("evidence-valid-until-reached")
        if decision_invalidated:
            reasons.append("recorded-decision-basis-changed")
        impacted_plans.append({
            "planId": plan["id"],
            "code": plan["code"],
            "revision": plan["revision"],
            "status": plan_status(w, plan, plan_checks),
            "reasonCodes": reasons,
            "expiredDocumentIds": sorted(expired_plan_document_ids),
            "checkIds": sorted(set(expired_check_ids)),
            "decisionInvalidated": decision_invalidated,
        })

    impacted_plans.sort(key=lambda item: item["planId"])
    expired_evidence.sort(key=lambda item: (item["validUntil"], item["documentId"]))
    overdue_tasks.sort(key=lambda item: (item["dueAt"], item["taskId"]))
    return {
        "runType": "expiry_recheck",
        "trigger": "manual",
        "asOf": as_of,
        "worker": worker_name,
        "scope": "已核验资料有效期、未完成任务截止时间、受影响的当前计划",
        "expiredEvidence": expired_evidence,
        "overdueTasks": overdue_tasks,
        "impactedPlans": impacted_plans,
        "counts": {
            "expiredEvidence": len(expired_evidence),
            "overdueTasks": len(overdue_tasks),
            "impactedPlans": len(impacted_plans),
        },
        "nextStep": "由负责人查看影响对象并手动补证、修订或重新复核；本次重检不会自动批准、取消或创建任务。",
    }


COPILOT_ACTION_KINDS = {
    "source_line", "equipment", "evidence", "rule", "relationship", "task", "scenario",
}


def _copilot_packet(w: dict, plan: dict, parsed: dict, evaluation: Any) -> dict:
    """Build the reviewable, structured part of the AI copilot response.

    The language model is only responsible for the prose summary.  This packet
    is deterministic and keeps every suggestion tied to a parsed step, a rule
    result, or a resource check so the UI can show *why* the assistant reached
    a conclusion.  It deliberately never returns approval or executable
    incident instructions.
    """
    plan_checks = [c for c in check_workspace(w) if c["planId"] == plan["id"]]
    failed_rules = [r for r in evaluation.results if not r.passed]

    def action(kind: str, target: dict[str, Any]) -> dict[str, Any]:
        """Attach a review-only navigation target to an AI item.

        These targets are deliberately declarative.  The browser uses them to
        focus an existing source, resource, evidence record, rule or graph;
        they never approve a plan or mutate the workspace by themselves.
        """
        if kind not in COPILOT_ACTION_KINDS:
            raise ValueError(f"未知的 Copilot 动作类型：{kind}")
        # Keep the target's discriminator aligned with the typed action.  The
        # browser treats this as a navigation hint only; it never mutates a
        # plan or bypasses a revision/approval check.
        target = dict(target)
        target.setdefault("type", kind)
        # Resolve a source span from the authoritative current plan/document
        # only.  Copilot prose never supplies coordinates; absent coordinates
        # remain absent so a model cannot fabricate a citation.
        span = target.get("sourceSpan")
        if not isinstance(span, dict) or not span:
            document_id = target.get("documentId")
            if document_id:
                document = next((item for item in w.get("documents", []) if item.get("id") == document_id), None)
                span = _document_source_span(document) if document else {}
            else:
                span = _plan_source_span(plan, target.get("line"))
        if span:
            target["sourceSpan"] = span
        return {"actionKind": kind, "actionTarget": target}

    def task_draft(item: dict[str, Any], title: str, note: str, owner: str, kind: str, target: dict[str, Any]) -> dict[str, Any]:
        """Carry the exact source item into a review-only task draft.

        A task form may be opened from several lists (evidence, conflicts, or
        a manual review fallback).  Capturing the source id and typed target
        prevents the form from silently falling back to the plan's first
        unresolved check.
        """
        draft_target = dict(target)
        draft_target.setdefault("type", kind)
        if not draft_target.get("sourceSpan"):
            document_id = draft_target.get("documentId")
            if document_id:
                document = next((item for item in w.get("documents", []) if item.get("id") == document_id), None)
                span = _document_source_span(document) if document else {}
            else:
                span = _plan_source_span(plan, draft_target.get("line"))
            if span:
                draft_target["sourceSpan"] = span
        return {
            "itemId": item.get("id"),
            "sourceId": item.get("sourceId") or item.get("id"),
            "title": title,
            "note": note,
            "owner": owner,
            "line": item.get("line"),
            "equipmentId": item.get("equipmentId"),
            "documentIds": list(item.get("documentIds") or []),
            "actionKind": kind,
            "actionTarget": draft_target,
        }
    steps = [step.model_dump(mode="json") if hasattr(step, "model_dump") else step for step in parsed.get("steps", [])]
    substances = [substance.model_dump(mode="json") if hasattr(substance, "model_dump") else substance for substance in parsed.get("substances", [])]
    plan_steps_by_line = {int(item.get("line")): item for item in plan.get("steps", []) if item.get("line") is not None}
    for step in steps:
        # Source line is the only legal join key.  Never borrow a plan step by
        # array position after a model has reordered or omitted entries.
        source_step = plan_steps_by_line.get(step.get("source_line") or step.get("line"))
        if source_step:
            step.setdefault("line", source_step.get("line"))
            step.setdefault("planEquipmentIds", source_step.get("equipmentIds", []))
    hazards: list[dict] = []
    relations: list[dict] = []
    for index, step in enumerate(steps, 1):
        hazard = step.get("hazard")
        if hazard or step.get("requires_ppe") or step.get("waste_stream") or step.get("ventilation"):
            source = f"SOP 原文 L{step.get('line')}" if step.get("line") else "SOP 原文（待定位）"
            hazards.append({
                "id": f"step-{step.get('order', index)}",
                "step": step.get("order", index),
                "line": step.get("line"),
                "title": hazard or "需要控制措施的步骤",
                "detail": step.get("instruction", ""),
                "controls": step.get("controls", []),
                "ppe": step.get("requires_ppe", []),
                "evidence": step.get("evidence_refs", []),
                "severity": "high" if hazard else "medium",
                "source": source,
                "origin": "source-anchored",
                **(action("source_line", {"planId": plan["id"], "line": step.get("line")}) if step.get("line") else action("rule", {"planId": plan["id"], "ruleId": "step-controls"})),
            })
        for substance in substances:
            substance_name = str(substance.get("name", "")).strip()
            if substance_name and substance_name in str(step.get("instruction", "")):
                relation_id = f"step-{step.get('order', index)}:substance-{substance_name}:uses"
            relations.append({"id": relation_id, "from": f"step-{step.get('order', index)}", "to": f"substance-{substance_name}", "kind": "uses", "label": "使用化学品", "source": f"SOP L{step.get('line')}" if step.get('line') else "SOP 原文", **action("relationship", {"planId": plan["id"], "line": step.get("line"), "relationId": relation_id, "from": f"step-{step.get('order', index)}", "to": f"substance-{substance_name}", "relation": "uses"})})
        for equipment in step.get("equipment", []):
            relation_id = f"step-{step.get('order', index)}:equipment-{equipment}:requires"
            relations.append({"id": relation_id, "from": f"step-{step.get('order', index)}", "to": f"equipment-{equipment}", "kind": "requires", "label": "需要设备", "source": f"SOP L{step.get('line')}" if step.get('line') else "SOP 原文", **action("relationship", {"planId": plan["id"], "line": step.get("line"), "relationId": relation_id, "from": f"step-{step.get('order', index)}", "to": f"equipment-{equipment}", "relation": "requires"})})
        for control in step.get("controls", []):
            relation_id = f"step-{step.get('order', index)}:control-{control}:controlled-by"
            relations.append({"id": relation_id, "from": f"step-{step.get('order', index)}", "to": f"control-{control}", "kind": "controlled-by", "label": "控制措施", "source": f"SOP L{step.get('line')}" if step.get('line') else "SOP 原文", **action("relationship", {"planId": plan["id"], "line": step.get("line"), "relationId": relation_id, "from": f"step-{step.get('order', index)}", "to": f"control-{control}", "relation": "controlled-by"})})
    # If a plan uses generic wording rather than a known substance, retain a
    # signal from the deterministic rule engine instead of pretending the
    # parser identified a chemical that was not present in the source.
    for result in failed_rules:
        if result.rule_id in {"ventilation-required", "emergency-controls", "incompatibility-separation", "step-controls"}:
            hazards.append({
                "id": f"rule-{result.rule_id}", "step": None,
                "title": result.message, "detail": result.recommended_control or result.message,
                "controls": [], "ppe": [], "evidence": list(result.evidence_refs),
                "severity": result.severity.value, "source": result.source or f"规则包 {result.rule_version}", "origin": "deterministic-rule",
                **action("rule", {"type": "rule", "planId": plan["id"], "ruleId": result.rule_id}),
            })
    # Keep the graph compact and stable when a long SOP contains repeated
    # mentions of the same control.
    unique_hazards: list[dict] = []
    seen_hazard: set[str] = set()
    for item in hazards:
        key = f"{item['title']}|{item['detail']}"
        if key not in seen_hazard:
            seen_hazard.add(key)
            unique_hazards.append(item)

    conflicts = []
    evidence_gaps = []
    for check in plan_checks:
        if check["state"] == "blocked":
            conflicts.append({
                "id": check["id"], "code": check["code"], "title": check["title"],
                "detail": check["detail"], "equipmentId": check.get("equipmentId"),
                "relatedPlanIds": check.get("relatedPlanIds", []),
                "line": check.get("line"), "severity": "critical",
                **(action("equipment", {"planId": plan["id"], "equipmentId": check.get("equipmentId")}) if check.get("equipmentId") else action("source_line", {"planId": plan["id"], "line": check.get("line")}) if check.get("line") else action("rule", {"planId": plan["id"], "ruleId": check.get("code")})),
            })
        elif check["state"] == "unknown":
            evidence_gaps.append({
                "id": check["id"], "code": check["code"], "title": check["title"],
                "detail": check["detail"], "equipmentId": check.get("equipmentId"),
                "line": check.get("line"), "documentIds": check.get("documentIds", []),
                "priority": "high" if check["code"] in {"EVIDENCE", "SCHEDULE", "BINDING"} else "medium",
                **(action("evidence", {"planId": plan["id"], "documentId": check["documentIds"][0], "equipmentId": check.get("equipmentId")}) if check.get("documentIds") else action("source_line", {"planId": plan["id"], "line": check.get("line")}) if check.get("line") else action("rule", {"planId": plan["id"], "ruleId": check.get("code")})),
            })

    next_actions = []
    for gap in evidence_gaps[:6]:
        if gap["code"] == "EVIDENCE":
            title = "补齐对应设备的有效检查证据"
            detail = "需同时满足对象、规格版本、核验状态和完整使用时段。"
        elif gap["code"] == "SOURCE":
            title = "确认 SOP 步骤与原文行号"
            detail = "确认后才会把机器提取结果纳入资源核验。"
        elif gap["code"] == "SCHEDULE":
            title = "补充资源使用时段"
            detail = "没有时段无法判断维护窗口与共享容量。"
        elif gap["code"] == "BINDING":
            title = "把原文设备绑定到正式预约"
            detail = "不能只替换画面中的关系线；需要修订计划。"
        else:
            title, detail = gap["title"], gap["detail"]
        target = gap.get("actionTarget") or {"type": "rule", "planId": plan["id"], "ruleId": gap.get("code")}
        kind = gap.get("actionKind") or "rule"
        owner = "实验负责人"
        next_actions.append({"id": f"action-{gap['id']}", "title": title, "detail": detail, "sourceId": gap["id"], "owner": owner, "status": "suggested", "line": gap.get("line"), "equipmentId": gap.get("equipmentId"), "documentIds": list(gap.get("documentIds") or []), "taskDraft": task_draft(gap, title, detail, owner, kind, target), **action(kind, target)})
    for conflict in conflicts[:4]:
        target = {"planId": plan["id"], "scenarioId": "resource-alternative"}
        owner = "计划负责人"
        next_actions.append({"id": f"action-{conflict['id']}", "title": "打开方案推演，比较可行资源或时段", "detail": "候选方案先在隔离副本中核验，正式计划不会被自动覆盖。", "sourceId": conflict["id"], "owner": owner, "status": "suggested", "line": conflict.get("line"), "equipmentId": conflict.get("equipmentId"), "taskDraft": task_draft(conflict, "复核资源或时段冲突", conflict["detail"], owner, "scenario", target), **action("scenario", target)})
    if not next_actions:
        target = {"planId": plan["id"], "ruleId": "HUMAN_REVIEW"}
        fallback = {"id": "action-review", "title": "由安全负责人复核当前版本", "detail": "AI 解释和规则结果不能替代机构审批或现场确认。", "sourceId": plan["id"], "owner": "安全负责人", "status": "required"}
        next_actions.append({**fallback, "taskDraft": task_draft(fallback, fallback["title"], fallback["detail"], fallback["owner"], "task", {"planId": plan["id"], "taskId": "draft-human-review"}), **action("task", {"planId": plan["id"], "taskId": "draft-human-review"})})

    # Produce two deterministic candidates for the scenario comparison.  They
    # are suggestions only; applying one still requires the existing branch
    # workflow and creates a new plan revision.
    baseline_counts = {"blocked": sum(c["state"] == "blocked" for c in plan_checks), "unknown": sum(c["state"] == "unknown" for c in plan_checks), "pass": sum(c["state"] == "pass" for c in plan_checks)}
    scenarios = [{"id": "baseline", "name": "当前正式安排", "kind": "baseline", "status": "current", "summary": "以当前计划和已登记资源为基线。", "counts": baseline_counts, "changes": [], "reasoning": ["保留正式计划作为比较基准", "不会因为 AI 结果自动修改"], "actionability": "reference", **action("scenario", {"planId": plan["id"], "scenarioId": "baseline"})}]
    # Alternate resource candidate: use the first available same-kind resource
    # for a conflicted/unknown reservation, if one exists.
    candidate_reservations = [dict(r) for r in plan["reservations"]]
    changes = []
    for reservation in candidate_reservations:
        current = next((e for e in w["equipment"] if e["id"] == reservation["equipmentId"]), None)
        if current is None:
            continue
        current_bad = any(c.get("equipmentId") == current["id"] and c["state"] in {"blocked", "unknown"} for c in plan_checks)
        if not current_bad:
            continue
        replacement = next((e for e in w["equipment"] if e["labId"] == plan["labId"] and e["id"] != current["id"] and e["kind"] == current["kind"] and e["status"] == "available" and not any(overlap(reservation, m) for m in e.get("maintenance", []))), None)
        if replacement:
            reservation["equipmentId"] = replacement["id"]
            changes.append({"from": current["id"], "to": replacement["id"], "reason": "同类别、当前台账可用且不落在维护窗口"})
    if changes:
        candidate = copy.deepcopy(w)
        candidate_plan = next(q for q in candidate["plans"] if q["id"] == plan["id"])
        candidate_plan["reservations"] = candidate_reservations
        candidate_checks = [c for c in check_workspace(candidate) if c["planId"] == plan["id"]]
        scenarios.append({"id": "resource-alternative", "name": "AI 资源替代候选", "kind": "alternative", "status": "suggested", "summary": "用同类别且当前可用资源替代受约束对象，结果仍需人工确认。", "counts": {"blocked": sum(c["state"] == "blocked" for c in candidate_checks), "unknown": sum(c["state"] == "unknown" for c in candidate_checks), "pass": sum(c["state"] == "pass" for c in candidate_checks)}, "changes": changes, "reasoning": ["同类别资源保持设备能力语义", "避开已登记维护窗口和停用状态", "替代结果将在隔离副本中重新核验"], "actionability": "branch", **action("scenario", {"planId": plan["id"], "scenarioId": "resource-alternative"})})
    scenarios.append({"id": "evidence-first", "name": "先补证再放行", "kind": "workflow", "status": "suggested", "summary": "不改变排程，优先处理证据缺口和原文确认。", "counts": baseline_counts, "changes": [{"action": "complete-evidence", "count": len(evidence_gaps)}], "reasoning": ["不改变正式资源预约", f"先闭环 {len(evidence_gaps)} 个证据或确认缺口", "由负责人完成复核后再形成新修订"], "actionability": "task", **action("scenario", {"planId": plan["id"], "scenarioId": "evidence-first"})})

    return {
        "entities": {"steps": len(steps), "hazards": unique_hazards, "substances": substances, "ppe": parsed.get("ppe", []), "equipment": parsed.get("equipment", []), "wasteStreams": parsed.get("waste_streams", []), "ventilation": parsed.get("ventilation"), "modelCandidates": parsed.get("model_candidates", [])},
        "relations": relations,
        "riskSignals": unique_hazards,
        "conflicts": conflicts,
        "evidenceGaps": evidence_gaps,
        "nextActions": next_actions,
        "scenarios": scenarios,
        "ruleSummary": {"total": len(plan_checks), **baseline_counts, "failedTextRules": [r.rule_id for r in failed_rules]},
    }


def _merge_model_parse(parsed: dict, model_fields: Any) -> tuple[dict, bool]:
    """Merge a model candidate without losing deterministic source anchors.

    The model is useful for prose SOPs where the deterministic vocabulary is
    intentionally small.  A model step is accepted only when it carries an
    explicit document revision and source line/text hash that matches the
    source ledger.  Array order, title similarity and missing spans are never
    allowed to inherit a trusted citation.
    """
    if not isinstance(model_fields, dict):
        return parsed, False
    merged = dict(parsed)
    # Candidate output is kept separate from the confirmed deterministic
    # parse.  A model may explain an anchored line, but it cannot silently
    # promote a guessed control or chemical into a fact used by rules.
    merged["model_candidates"] = []
    used = False

    raw_steps = model_fields.get("steps")
    if isinstance(raw_steps, list):
        candidates = []
        from .models import LabStep, Substance
        source_steps = list(parsed.get("steps", []))
        for index, item in enumerate(raw_steps):
            if not isinstance(item, dict):
                continue
            instruction = str(item.get("instruction", "")).strip()
            if not instruction:
                continue
            revision = item.get("documentRevisionId") or item.get("document_revision_id")
            source_line = item.get("sourceLine") or item.get("source_line")
            source_hash = item.get("sourceTextHash") or item.get("source_text_hash")
            try:
                source_line = int(source_line)
            except (TypeError, ValueError):
                source_line = None
            matched = next((candidate for candidate in source_steps if candidate.source_line == source_line), None)
            if not revision or revision != parsed.get("document_revision_id") or matched is None:
                # Unanchored model text remains visible only through the
                # deterministic fallback; it cannot become a cited fact.
                continue
            if source_hash and source_hash != matched.source_text_hash:
                continue
            if matched.instruction.strip() != instruction.strip():
                continue
            try:
                step = LabStep(
                    order=int(item.get("order") or index + 1),
                    instruction=instruction[:500],
                    hazard=str(item.get("hazard") or "").strip() or None,
                    controls=[str(v).strip() for v in item.get("controls", []) if str(v).strip()][:20],
                    requires_ppe=[str(v).strip() for v in item.get("requires_ppe", []) if str(v).strip()][:20],
                    equipment=[str(v).strip() for v in item.get("equipment", []) if str(v).strip()][:20],
                    ventilation=str(item.get("ventilation") or "").strip() or None,
                    waste_stream=str(item.get("waste_stream") or "").strip() or None,
                    evidence_refs=list(matched.evidence_refs),
                    document_revision_id=matched.document_revision_id,
                    source_line=matched.source_line,
                    source_text_hash=matched.source_text_hash,
                    control_assertions=list(matched.control_assertions),
                )
            except (TypeError, ValueError):
                continue
            candidates.append(step)
            merged["model_candidates"].append({
                "kind": "step",
                "documentRevisionId": matched.document_revision_id,
                "sourceLine": matched.source_line,
                "sourceTextHash": matched.source_text_hash,
                "instruction": instruction,
                "hazard": step.hazard,
                "controls": list(step.controls),
                "requiresPpe": list(step.requires_ppe),
                "equipment": list(step.equipment),
                "ventilation": step.ventilation,
                "wasteStream": step.waste_stream,
                "status": "candidate",
            })
        if candidates:
            # Deterministic fields remain authoritative.  This survives model
            # reordering, merging and omission without borrowing another
            # line's evidence or treating a candidate as confirmation.
            used = True

    raw_substances = model_fields.get("substances")
    if isinstance(raw_substances, list):
        candidates = []
        for item in raw_substances:
            if not isinstance(item, dict) or not str(item.get("name", "")).strip():
                continue
            if str(item.get("name", "")).strip() not in parsed.get("source", "") and str(item.get("name", "")).strip() not in "\n".join(getattr(s, "instruction", "") for s in parsed.get("steps", [])):
                continue
            try:
                name = str(item["name"]).strip()[:180]
                candidates.append({
                    "name": name,
                    "casNumber": str(item.get("cas_number") or "").strip() or None,
                    "hazardClasses": [str(v).strip() for v in item.get("hazard_classes", []) if str(v).strip()][:10],
                    "status": "candidate",
                })
            except (TypeError, ValueError):
                continue
        if candidates:
            # A model may suggest a chemical and a hazard class, but a name
            # substring is not a source anchor.  Keep these suggestions in a
            # visible candidate lane; only the deterministic parser (or a
            # later independently verified SDS) may change authoritative
            # substances used by rules.  This prevents prompt-injected JSON
            # from turning an untrusted model guess into a risk fact.
            merged["model_candidates"].append({
                "kind": "substances",
                "values": candidates,
                "status": "candidate",
                "authority": "none-until-source-anchored-and-reviewed",
            })
            used = True

    for key in ("ppe", "equipment", "waste_streams"):
        values = model_fields.get(key)
        if isinstance(values, list):
            normalized = sorted({str(value).strip() for value in values if str(value).strip()})
            if normalized:
                source_present = [value for value in normalized if any(value in str(getattr(step, "instruction", "")) for step in parsed.get("steps", []))]
                # A substring is not a source assertion: it can appear in a
                # negated sentence or in a prompt-injection aside. Keep the
                # model output visible, but never let it change the
                # authoritative rule inputs without a later human/SDS review.
                merged["model_candidates"].append({"kind": key, "values": normalized, "sourcePresent": source_present, "status": "candidate", "authority": "none-until-reviewed"})
                used = True
    ventilation = model_fields.get("ventilation")
    if isinstance(ventilation, str) and ventilation.strip():
        candidate = ventilation.strip()
        source_present = bool(merged.get("ventilation")) and any(candidate in str(getattr(step, "instruction", "")) for step in parsed.get("steps", []))
        merged["model_candidates"].append({"kind": "ventilation", "value": candidate, "sourcePresent": source_present, "status": "candidate", "authority": "none-until-reviewed"})
        used = True
    return merged, used


def apply_operation(w: dict, action: str, x: dict, actor: str) -> dict:
    next_w = copy.deepcopy(w)
    object_id = x.get("id", x.get("planId", "workspace"))
    summary = ""
    def get(collection, identifier):
        value = next((a for a in next_w[collection] if a["id"] == identifier), None)
        if value is None:
            raise ValueError("关联对象不存在。")
        return value
    def upsert(collection, value):
        index = next((i for i, a in enumerate(next_w[collection]) if a["id"] == value["id"]), None)
        if index is None:
            next_w[collection].append(value)
        else:
            next_w[collection][index] = value
    if action == "initialize":
        if any(w[key] for key in ("laboratories", "equipment", "plans", "documents", "tasks", "decisions", "branches")):
            raise ValueError("已有数据的工作区不能被初始化覆盖。")
        next_w = copy.deepcopy(x["workspace"])
        next_w["decisions"], next_w["events"], next_w["branches"] = [], [], []
        summary = "导入明确标记的样例工作区" if next_w["provenance"] == "example" else "建立空白工作区"
    elif action == "save_lab":
        # Validate the submitted object before touching the workspace.  The
        # final workspace validation below remains authoritative, while this
        # gives callers a useful error when a malformed lab is submitted by a
        # browser form.
        Lab.model_validate(x["lab"])
        upsert("laboratories", x["lab"])
        object_id, summary = x["lab"]["id"], "保存实验室资料"
    elif action == "save_equipment":
        e = copy.deepcopy(x["equipment"])
        previous = next((a for a in w["equipment"] if a["id"] == e["id"]), None)
        if not any(lab["id"] == e["labId"] for lab in next_w["laboratories"]):
            raise ValueError("资源所属实验室不存在。")
        # The form starts a new resource at (0, 0).  Treat that untouched
        # default as a request for server-side placement so repeated saves do
        # not pile models on top of one another.  Callers that intentionally
        # want (0, 0) may pass ``autoPlace: false`` at the operation level.
        at_origin = abs(float(e["x"])) < 1e-9 and abs(float(e["z"])) < 1e-9
        auto_place = at_origin and x.get("autoPlace", True) is not False
        if previous is None and auto_place:
            e["x"], e["z"] = find_equipment_position(next_w, e)
        e["specVersion"] = previous["specVersion"] + int(any(e[k] != previous[k] for k in ("kind", "capacity", "requiresEvidence"))) if previous else 1
        upsert("equipment", e)
        object_id, summary = e["id"], f"更新资源 {e['code']}，核验依赖同步重算"
    elif action == "save_plan":
        p = copy.deepcopy(x["plan"])
        previous = next((a for a in w["plans"] if a["id"] == p["id"]), None)
        if previous and previous["revision"] != x.get("expectedPlanRevision"):
            raise ValueError("计划已变更，请重新载入后保存。")
        if previous and (previous["archived"] or previous.get("lifecycle") == "archived") and not p.get("archived"):
            raise ValueError("已归档计划只能引用或复制，不能直接复活为当前计划。")
        p["revision"] = previous["revision"] + 1 if previous else 1
        if previous and p["source"] != previous["source"]:
            for s in p["steps"]:
                s["confirmed"] = False
            p["lifecycle"] = "revising"
        elif previous:
            p["lifecycle"] = previous.get("lifecycle", "draft")
        else:
            p["lifecycle"] = p.get("lifecycle", "draft")
        if p.get("lifecycle") == "submitted" and (not p["steps"] or any(not s["confirmed"] for s in p["steps"])):
            raise ValueError("计划提交前必须完成全部步骤的人工确认。")
        if p.get("archived"):
            # Archive is a business transition, not a display toggle.  Check
            # the still-active candidate before setting archived=True, because
            # check_workspace intentionally excludes archived plans.
            if not p["steps"] or any(not s["confirmed"] for s in p["steps"]):
                raise ValueError("计划归档前必须完成全部步骤的人工确认。")
            candidate = copy.deepcopy(next_w)
            candidate_plan = next((q for q in candidate["plans"] if q["id"] == p["id"]), None)
            if candidate_plan is None:
                candidate["plans"].append(copy.deepcopy(p))
                candidate_plan = candidate["plans"][-1]
            else:
                candidate_plan.update(copy.deepcopy(p))
            candidate_plan["archived"] = False
            candidate_plan["lifecycle"] = "submitted"
            archive_checks = [c for c in check_workspace(candidate) if c["planId"] == p["id"]]
            if not archive_checks or any(c["state"] != "pass" for c in archive_checks):
                raise ValueError("计划仍有阻断或未知项，不能归档。")
            decisions = [d for d in w["decisions"] if d["planId"] == p["id"]]
            # Saving an archive record increments the plan revision for the
            # audit log.  Compare the approval against the pre-save revision;
            # lifecycle metadata itself must not invalidate an otherwise
            # unchanged decision.
            basis_plan = copy.deepcopy(candidate_plan)
            if previous:
                basis_plan["revision"] = previous["revision"]
            if decisions and decisions[-1]["basis"] != basis(candidate, basis_plan):
                raise ValueError("当前计划基线与最近复核不一致，请先重新复核后归档。")
            p["lifecycle"] = "archived"
        elif p.get("lifecycle") == "archived":
            raise ValueError("归档生命周期必须同时设置 archived=true。")
        upsert("plans", p)
        object_id, summary = p["id"], f"保存计划 {p['code']} · v{p['revision']}"
    elif action == "confirm_steps":
        p = get("plans", x["planId"])
        if not p["steps"]:
            raise ValueError("没有可确认的步骤。")
        if p["archived"] or p.get("lifecycle") in {"archived", "cancelled"}:
            raise ValueError("已归档或已取消计划不能重新确认步骤。")
        for s in p["steps"]:
            s["confirmed"] = True
        p["revision"] += 1
        summary = "人工核对原文与全部步骤"
    elif action == "save_document":
        d = copy.deepcopy(x["document"])
        previous = next((item for item in w["documents"] if item["id"] == d["id"]), None)
        # Preserve capture provenance across edits while recording a bounded,
        # inspectable text length for the exact content that was submitted.
        # Metadata records where the text came from; it never asserts source
        # authority. Reviewers still verify source, object, version and time.
        if previous:
            for field in ("sourceType", "sourceLocator", "capturedAt", "provenanceNote", "sourceFileName", "sourceFileSize", "sourceMediaType", "sourceLocations", "sourceArtifactId", "sourceArtifactVersion"):
                if d.get(field) is None and previous.get(field) is not None:
                    d[field] = previous[field]
        d.setdefault("sourceType", "manual")
        d.setdefault("capturedAt", now())
        d["contentLength"] = len(d.get("content", ""))
        d.setdefault("sourceLocations", [])
        d.update(status="pending", reviewedBy=None, reviewedAt=None)
        upsert("documents", d)
        object_id, summary = d["id"], "提交资料，等待独立核验"
    elif action == "verify_document":
        d = get("documents", x["id"])
        if x.get("status") not in ("verified", "rejected"):
            raise ValueError("资料核验状态无效。")
        d.update(status=x["status"], reviewedBy=actor, reviewedAt=now())
        summary = "记录资料核验结果" if d["status"] == "verified" else "退回资料"
    elif action == "save_task":
        task = copy.deepcopy(x["task"])
        # An AI action carries a stable source id and action kind.  Reopening
        # the same suggestion must not create a second indistinguishable open
        # task; an explicitly completed task may be reopened only by changing
        # its source/action or by editing the existing record.
        if task.get("sourceId") and task.get("actionKind"):
            duplicate = next((existing for existing in next_w["tasks"] if existing["id"] != task.get("id")
                and existing.get("planId") == task.get("planId")
                and existing.get("sourceId") == task.get("sourceId")
                and existing.get("sourceLine") == task.get("sourceLine")
                and existing.get("actionKind") == task.get("actionKind")
                and existing.get("status") in {"open", "submitted", "returned", "reopened"}), None)
            if duplicate:
                raise ValueError("该 AI 建议已有未完成任务，请打开原任务继续处理。")
        upsert("tasks", task)
        object_id, summary = x["task"]["id"], "创建协作任务"
    elif action == "assign_task":
        # Assignment is an explicit collaboration mutation, separate from a
        # task status update.  It leaves the evidence/approval basis untouched
        # and records the hand-off reason in the audit event below.  The local
        # role gate in ``mutate_workspace`` keeps this operation from being
        # presented as a multi-user permission system.
        t = get("tasks", x["id"])
        if t["status"] in {"done", "cancelled"}:
            raise ValueError("已完成或已取消任务不能直接转交；请先重新打开任务。")
        new_owner = str(x.get("newOwner", "")).strip()
        note = str(x.get("note", "")).strip()
        expected_owner = x.get("expectedOwner")
        if expected_owner is not None and str(expected_owner) != t["owner"]:
            raise ValueError("任务负责人已变化，请刷新后重新转交。")
        if not new_owner or len(new_owner) > 180:
            raise ValueError("转交负责人不能为空且不能超过 180 个字符。")
        if new_owner == t["owner"]:
            raise ValueError("转交负责人必须与当前负责人不同。")
        if not note or len(note) > 2000:
            raise ValueError("转交任务必须填写不超过 2000 个字符的交接说明。")
        previous_owner = t["owner"]
        t["owner"] = new_owner
        object_id = t["id"]
        summary = f"转交协作任务：{previous_owner} → {new_owner}；{note}"
    elif action == "update_task":
        t = get("tasks", x["id"])
        requested = x.get("status")
        allowed = {"open", "submitted", "done", "returned", "cancelled", "reopened"}
        if requested not in allowed:
            raise ValueError("任务状态无效。")
        current = t["status"]
        transitions = {
            "open": {"open", "submitted", "done", "cancelled"},
            "submitted": {"submitted", "done", "returned", "cancelled"},
            "returned": {"returned", "reopened", "cancelled"},
            "reopened": {"reopened", "submitted", "done", "cancelled"},
            "done": {"done", "reopened"},
            "cancelled": {"cancelled", "reopened"},
        }
        if requested not in transitions.get(current, set()):
            raise ValueError(f"任务不能从“{current}”流转到“{requested}”。")
        note = str(x.get("note", t["note"])).strip()
        if requested in {"returned", "cancelled", "reopened"} and not note:
            raise ValueError("退回、取消或重开任务必须填写处理说明。")
        handoff_requested = any(key in x for key in ("newOwner", "expectedOwner", "handoffNote"))
        previous_owner = t["owner"]
        new_owner = ""
        handoff_note = ""
        if handoff_requested:
            # Status changes and ownership changes are one business operation.
            # This avoids leaving a task half-transferred when the second
            # request in a browser form fails or loses its revision race.
            if current in {"done", "cancelled"}:
                raise ValueError("已完成或已取消任务不能在同一操作中直接转交；请先重新打开任务。")
            expected_owner = x.get("expectedOwner")
            if expected_owner is None or str(expected_owner) != t["owner"]:
                raise ValueError("任务负责人已变化，请刷新后重新转交。")
            new_owner = str(x.get("newOwner", "")).strip()
            handoff_note = str(x.get("handoffNote", "")).strip()
            if not new_owner or len(new_owner) > 180:
                raise ValueError("转交负责人不能为空且不能超过 180 个字符。")
            if new_owner == t["owner"]:
                raise ValueError("转交负责人必须与当前负责人不同。")
            if not handoff_note or len(handoff_note) > 2000:
                raise ValueError("转交任务必须填写不超过 2000 个字符的交接说明。")
        t.update(status=requested, note=note)
        if handoff_requested:
            t["owner"] = new_owner
        labels = {"returned": "退回补充", "cancelled": "取消任务", "reopened": "重新打开"}
        summary = labels.get(requested, "更新任务进度") + "；不自动改变证据或审批"
        if handoff_requested:
            summary = f"{summary}；同时转交：{previous_owner} → {new_owner}；{handoff_note}"
    elif action == "save_branch":
        b = copy.deepcopy(x["branch"])
        p = get("plans", b["planId"])
        if p["revision"] != b["baseRevision"] or basis(w, p) != b["baseBasis"]:
            raise ValueError("方案基线已变化，请以当前计划重新创建。")
        branch_candidate(next_w, b)
        upsert("branches", b)
        object_id, summary = b["id"], "保存独立候选方案，正式计划未改变"
    elif action == "review_branch_patch":
        branch_id = str(x.get("branchId", ""))
        b = get("branches", branch_id)
        p = get("plans", b["planId"])
        if p["revision"] != b["baseRevision"] or basis(w, p) != b["baseBasis"]:
            raise ValueError("方案基线已变化，请重新创建后再审查候选字段。")
        path = x.get("path")
        target_id = str(x.get("targetId", ""))
        disposition = x.get("disposition")
        if path not in BRANCH_PATCH_PATHS:
            raise ValueError("候选字段路径无效。")
        if disposition not in {"accepted", "rejected"}:
            raise ValueError("候选字段取舍必须是 accepted 或 rejected。")
        note = str(x.get("reviewNote", "")).strip()
        if not note:
            raise ValueError("候选字段取舍必须填写复核说明。")
        patch = next((item for item in b.get("patches", []) if item.get("path") == path and item.get("targetId") == target_id), None)
        if patch is None:
            raise ValueError("候选字段不存在。")
        # Validate the unmodified branch before recording the decision, then
        # rebuild its candidate snapshot so rejected fields visibly fall back
        # to the formal plan while the suggestion itself remains in history.
        branch_candidate(next_w, b)
        patch["disposition"] = disposition
        patch["reviewedBy"] = actor
        patch["reviewedAt"] = now()
        patch["reviewNote"] = note[:2000]
        reservations, states = _effective_branch_snapshots(next_w, b)
        b["reservations"] = reservations
        b["equipmentStates"] = states
        branch_candidate(next_w, b)
        object_id = branch_id
        summary = f"复核候选字段：{path} · {target_id} · {'采纳' if disposition == 'accepted' else '拒绝'}；{note[:200]}"
    elif action == "apply_branch":
        b = get("branches", x["id"])
        p = get("plans", b["planId"])
        if p["revision"] != b["baseRevision"] or basis(w, p) != b["baseBasis"]:
            raise ValueError("方案已过期：计划、资源或证据发生变化，请重新比较。")
        candidate = branch_candidate(next_w, b)
        candidate_plan = next(q for q in candidate["plans"] if q["id"] == p["id"])
        p["reservations"] = copy.deepcopy(candidate_plan["reservations"])
        # Apply only the explicitly audited equipment conditions from the
        # isolated candidate.  Other workspace resources retain their formal
        # state; the branch cannot replace the whole equipment ledger.
        candidate_equipment = {item["id"]: item for item in candidate["equipment"]}
        audited_equipment_ids = {
            patch["targetId"] for patch in b.get("patches", [])
            if patch.get("path") == "equipment.status" and patch.get("disposition", "accepted") == "accepted"
        }
        for equipment in next_w["equipment"]:
            if equipment["id"] not in audited_equipment_ids:
                continue
            projected = candidate_equipment.get(equipment["id"])
            if projected is not None and projected["status"] != equipment["status"]:
                equipment["status"] = projected["status"]
        _apply_accepted_branch_fields(next_w, b)
        p["revision"] += 1
        patch_count = len(b.get("patches", []))
        detail = f"，记录 {patch_count} 个字段变更" if patch_count else ""
        object_id, summary = p["id"], f"应用方案「{b['name']}」，产生新修订并触发重新复核{detail}"
    elif action == "approve":
        p = get("plans", x["planId"])
        if p["revision"] != x.get("expectedPlanRevision") or p["archived"] or p.get("lifecycle") in {"archived", "cancelled"}:
            raise ValueError("计划版本已改变或已归档。")
        if not x.get("comment", "").strip():
            raise ValueError("必须填写人工复核意见。")
        checks = [c for c in check_workspace(next_w) if c["planId"] == p["id"]]
        if not checks or any(c["state"] != "pass" for c in checks):
            raise ValueError("仍有阻断或未知项，不能记录复核通过。")
        role = str(x.get("_serverRole", "safety_owner"))
        stage = approval_stages()[0]
        next_w["decisions"].append(dict(
            id=uid("decision"), planId=p["id"], revision=p["revision"],
            basis=basis(next_w, p), actor=actor, comment=x["comment"], at=now(),
            disposition="approved", actorRole=role, stage=stage,
            checkSnapshot=decision_check_snapshot(checks), exception=None,
        ))
        p["lifecycle"] = "submitted"
        summary = "记录当前版本的资源前置条件复核；不构成化学安全认证"
    elif action == "approval_decision":
        p = get("plans", x["planId"])
        if p["revision"] != x.get("expectedPlanRevision") or p["archived"] or p.get("lifecycle") in {"archived", "cancelled"}:
            raise ValueError("计划版本已改变或已归档。")
        comment = str(x.get("comment", "")).strip()
        if not comment:
            raise ValueError("必须填写人工复核意见。")
        disposition = x.get("disposition")
        if disposition not in {"approved", "returned", "exception"}:
            raise ValueError("审批决定必须是 approved、returned 或 exception。")
        role = str(x.get("_serverRole", "safety_owner"))
        if role not in {"safety_owner", "ehs_manager", "administrator"}:
            raise ValueError("当前工作区身份没有记录正式审批的权限。")
        checks = [c for c in check_workspace(next_w) if c["planId"] == p["id"]]
        if not checks:
            raise ValueError("计划没有可供复核的检查快照。")
        current = approval_state(next_w, p, checks)
        stages = approval_stages()
        pending_stage = current["pendingStage"] or stages[0]
        if len(stages) > 1 and role != pending_stage and role != "administrator":
            raise ValueError(f"当前审批链等待“{pending_stage}”阶段，不能由当前身份越级处理。")
        stage = pending_stage
        exception_payload = None
        if disposition == "approved":
            if any(c["state"] != "pass" for c in checks):
                raise ValueError("仍有阻断或未知项，不能记录复核通过。")
            p["lifecycle"] = "submitted"
            summary = f"记录审批阶段 {stage} 的复核通过；不构成化学安全认证"
        elif disposition == "returned":
            p["lifecycle"] = "revising"
            summary = f"审批阶段 {stage} 退回计划，等待补充后重新复核"
        else:
            if role not in {"ehs_manager", "administrator"}:
                raise ValueError("例外授权只能由 EHS 负责人或管理员记录。")
            raw_exception = x.get("exception")
            if not isinstance(raw_exception, dict):
                raise ValueError("例外授权必须说明检查项、原因、缓解措施和失效时间。")
            check_ids = raw_exception.get("checkIds")
            reason = str(raw_exception.get("reason", "")).strip()
            mitigation = str(raw_exception.get("mitigation", "")).strip()
            expires_at = str(raw_exception.get("expiresAt", "")).strip()
            if not isinstance(check_ids, list) or not check_ids or any(not isinstance(item, str) for item in check_ids):
                raise ValueError("例外授权至少要指定一个检查项。")
            by_id = {c["id"]: c for c in checks}
            selected = [by_id.get(item) for item in check_ids]
            if any(item is None for item in selected):
                raise ValueError("例外授权引用了当前计划不存在的检查项。")
            if any(item["state"] == "pass" for item in selected):
                raise ValueError("例外授权只能针对当前阻断或待确认项。")
            if not reason or not mitigation:
                raise ValueError("例外授权必须填写原因和缓解措施。")
            try:
                expiry = dt(expires_at)
            except (TypeError, ValueError) as error:
                raise ValueError("例外失效时间必须是带时区的 ISO 时间。") from error
            if expiry <= datetime.now(timezone.utc):
                raise ValueError("例外失效时间必须晚于当前时间。")
            exception_payload = {"checkIds": list(dict.fromkeys(check_ids)), "reason": reason[:2000], "mitigation": mitigation[:2000], "expiresAt": expires_at}
            p["lifecycle"] = "submitted"
            summary = f"审批阶段 {stage} 记录限时例外；原始检查仍保留为阻断或待确认"
        next_w["decisions"].append(dict(
            id=uid("decision"), planId=p["id"], revision=p["revision"],
            basis=basis(next_w, p), actor=actor, comment=comment, at=now(),
            disposition=disposition, actorRole=role, stage=stage,
            checkSnapshot=decision_check_snapshot(checks), exception=exception_payload,
        ))
        object_id = p["id"]
    elif action == "transition_plan":
        p = get("plans", x["planId"])
        requested = x.get("lifecycle", x.get("status"))
        allowed = {"draft", "submitted", "revising", "archived", "cancelled"}
        if requested not in allowed:
            raise ValueError("计划生命周期无效。")
        current = "archived" if p["archived"] else p.get("lifecycle", "draft")
        transitions = {
            "draft": {"draft", "submitted", "cancelled"},
            "submitted": {"submitted", "revising", "cancelled", "archived"},
            "revising": {"revising", "submitted", "cancelled"},
            "cancelled": {"cancelled"},
            "archived": {"archived"},
        }
        if requested not in transitions[current]:
            raise ValueError(f"计划不能从“{current}”流转到“{requested}”。")
        if requested == "submitted":
            if not p["steps"] or any(not s["confirmed"] for s in p["steps"]):
                raise ValueError("计划提交前必须完成全部步骤的人工确认。")
            p["lifecycle"] = "submitted"
            summary = "提交当前计划版本，等待资源与证据复核"
        elif requested == "revising":
            p["lifecycle"] = "revising"
            summary = "打开计划修订状态；正式版本未被候选覆盖"
        elif requested == "cancelled":
            p["lifecycle"] = "cancelled"
            p["archived"] = False
            summary = "取消计划；保留历史记录且不再参与当前资源检查"
        elif requested == "archived":
            if not p["steps"] or any(not s["confirmed"] for s in p["steps"]):
                raise ValueError("计划归档前必须完成全部步骤的人工确认。")
            candidate = copy.deepcopy(next_w)
            candidate_plan = next(q for q in candidate["plans"] if q["id"] == p["id"])
            candidate_plan["archived"] = False
            candidate_plan["lifecycle"] = "submitted"
            archive_checks = [c for c in check_workspace(candidate) if c["planId"] == p["id"]]
            if not archive_checks or any(c["state"] != "pass" for c in archive_checks):
                raise ValueError("计划仍有阻断或未知项，不能归档。")
            decisions = [d for d in w["decisions"] if d["planId"] == p["id"]]
            if decisions and decisions[-1]["basis"] != basis(candidate, candidate_plan):
                raise ValueError("当前计划基线与最近复核不一致，请先重新复核后归档。")
            p["lifecycle"] = "archived"
            p["archived"] = True
            summary = "归档当前计划版本；历史依据仍可追溯，不能直接复活"
        else:
            p["lifecycle"] = requested
            summary = "更新计划生命周期"
    else:
        raise ValueError("不支持的操作。")
    next_w["revision"] = w["revision"] + 1
    next_w["events"].append(dict(id=uid("event"), at=now(), actor=actor, action=action, objectId=object_id, summary=summary))
    return validate_workspace(next_w)


class WorkspaceStore:
    def __init__(self, path: str, *, recover_incomplete: bool | None = None):
        self.path = path
        # The browser deployment keeps process-local callbacks as the default
        # so existing callers remain deterministic.  A separate Worker can
        # opt into durable queued/running rows and take ownership explicitly.
        if recover_incomplete is None:
            recover_incomplete = os.getenv("AEGIS_SPATIAL_JOB_EXECUTION", "process-local").strip().lower() != "worker"
        self.recover_incomplete = bool(recover_incomplete)
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        with self.connection() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS spatial_workspace(id INTEGER PRIMARY KEY CHECK(id=1), revision INTEGER NOT NULL, payload TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS spatial_operations(id TEXT PRIMARY KEY, request_hash TEXT NOT NULL, result_revision INTEGER NOT NULL);
                CREATE TABLE IF NOT EXISTS spatial_analysis_runs(
                    run_id TEXT PRIMARY KEY, plan_id TEXT NOT NULL, plan_revision INTEGER NOT NULL,
                    generated_at TEXT NOT NULL, degraded INTEGER NOT NULL, payload TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_spatial_analysis_plan ON spatial_analysis_runs(plan_id, generated_at DESC);
                CREATE TABLE IF NOT EXISTS spatial_evaluation_runs(
                    evaluation_id TEXT PRIMARY KEY, plan_id TEXT NOT NULL, plan_revision INTEGER NOT NULL,
                    evaluated_at TEXT NOT NULL, basis TEXT NOT NULL, status TEXT NOT NULL, payload TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_spatial_evaluation_plan ON spatial_evaluation_runs(plan_id, evaluated_at DESC);
                CREATE TABLE IF NOT EXISTS spatial_jobs(
                    job_id TEXT PRIMARY KEY, kind TEXT NOT NULL, plan_id TEXT,
                    idempotency_key TEXT NOT NULL UNIQUE, status TEXT NOT NULL,
                    created_at TEXT NOT NULL, started_at TEXT, finished_at TEXT,
                    result_payload TEXT, error TEXT, request_payload TEXT
                );
                CREATE INDEX IF NOT EXISTS idx_spatial_jobs_plan ON spatial_jobs(plan_id, created_at DESC);
                CREATE TABLE IF NOT EXISTS spatial_workspace_snapshots(
                    snapshot_id TEXT PRIMARY KEY, revision INTEGER NOT NULL,
                    saved_at TEXT NOT NULL, actor TEXT NOT NULL, reason TEXT NOT NULL,
                    payload TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_spatial_workspace_snapshots_saved
                    ON spatial_workspace_snapshots(saved_at DESC);
                CREATE TABLE IF NOT EXISTS spatial_outbox(
                    event_id TEXT PRIMARY KEY,
                    idempotency_key TEXT NOT NULL UNIQUE,
                    event_type TEXT NOT NULL,
                    aggregate_type TEXT NOT NULL,
                    aggregate_id TEXT NOT NULL,
                    plan_id TEXT,
                    audience TEXT NOT NULL,
                    title TEXT NOT NULL,
                    summary TEXT NOT NULL,
                    actor TEXT NOT NULL,
                    workspace_revision INTEGER NOT NULL,
                    source_event_id TEXT,
                    status TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    acknowledged_at TEXT,
                    acknowledged_by TEXT
                );
                CREATE INDEX IF NOT EXISTS idx_spatial_outbox_status
                    ON spatial_outbox(status, created_at DESC);
                CREATE INDEX IF NOT EXISTS idx_spatial_outbox_plan
                    ON spatial_outbox(plan_id, created_at DESC);
                CREATE TABLE IF NOT EXISTS spatial_memberships(
                    principal TEXT PRIMARY KEY,
                    role TEXT NOT NULL,
                    token_hash TEXT NOT NULL UNIQUE,
                    enabled INTEGER NOT NULL DEFAULT 1 CHECK(enabled IN (0,1)),
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    disabled_at TEXT,
                    last_seen_at TEXT
                );
                CREATE INDEX IF NOT EXISTS idx_spatial_memberships_token
                    ON spatial_memberships(token_hash);
                CREATE TABLE IF NOT EXISTS spatial_membership_operations(
                    request_id TEXT PRIMARY KEY,
                    request_hash TEXT NOT NULL,
                    principal TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS spatial_source_artifacts(
                    artifact_id TEXT PRIMARY KEY,
                    document_id TEXT NOT NULL,
                    filename TEXT NOT NULL,
                    media_type TEXT NOT NULL,
                    size INTEGER NOT NULL,
                    storage_key TEXT NOT NULL UNIQUE,
                    uploaded_by TEXT NOT NULL,
                    uploaded_at TEXT NOT NULL,
                    version INTEGER NOT NULL DEFAULT 1,
                    status TEXT NOT NULL DEFAULT 'stored'
                );
                CREATE INDEX IF NOT EXISTS idx_spatial_source_artifacts_document
                    ON spatial_source_artifacts(document_id, uploaded_at DESC);
            """)
            # Upgrade local databases created before queued browser jobs kept
            # their replay inputs.  The nullable column is intentionally
            # additive so existing job IDs and terminal history remain intact.
            columns = {row[1] for row in conn.execute("PRAGMA table_info(spatial_jobs)").fetchall()}
            if "request_payload" not in columns:
                conn.execute("ALTER TABLE spatial_jobs ADD COLUMN request_payload TEXT")
            # There is no worker process in the browser-only deployment.  A
            # process that opens the database after a crash/restart therefore
            # must make an unfinished synchronous job explicit instead of
            # leaving a misleading `running` state forever.
            if self.recover_incomplete:
                conn.execute(
                    """UPDATE spatial_jobs
                       SET status='failed', finished_at=COALESCE(finished_at, ?),
                           error=COALESCE(error, ?)
                       WHERE status IN ('queued','running')""",
                    (now(), "服务重启时进程内任务未能继续；本地部署没有持久 Worker，任务已标记失败，可按幂等边界重新发起。"),
                )
            conn.execute("INSERT OR IGNORE INTO spatial_workspace VALUES(1,0,?)", (json.dumps(empty_workspace(), ensure_ascii=False),))
    @contextmanager
    def connection(self):
        conn = sqlite3.connect(self.path, timeout=10)
        try:
            with conn:
                yield conn
        finally:
            conn.close()
    def read(self) -> dict:
        with self.connection() as conn:
            return json.loads(conn.execute("SELECT payload FROM spatial_workspace WHERE id=1").fetchone()[0])

    @staticmethod
    def _artifact_row(row: sqlite3.Row | tuple | None) -> dict | None:
        if row is None:
            return None
        keys = ("artifactId", "documentId", "filename", "mediaType", "size", "storageKey", "uploadedBy", "uploadedAt", "version", "status")
        item = dict(zip(keys, row))
        item.pop("storageKey", None)
        return item

    def save_source_artifact(self, *, document_id: str, filename: str, media_type: str, data: bytes, actor: str) -> dict:
        """Atomically retain one bounded source file for a workspace document."""
        document_id = str(document_id or "").strip()
        safe_name = _artifact_filename(filename)
        suffix = Path(safe_name).suffix.lower()
        if not document_id:
            raise HTTPException(422, "来源文件必须关联一条资料记录。")
        if suffix not in ARTIFACT_SUFFIXES:
            raise HTTPException(422, "仅支持 TXT、Markdown、DOCX 和 PDF 来源文件。")
        if len(data) > ARTIFACT_MAX_BYTES:
            raise HTTPException(413, "文件不能超过 8 MB。")
        if not data:
            raise HTTPException(422, "来源文件不能为空。")
        artifact_id = uid("artifact")
        storage_key = f"{artifact_id}{suffix}"
        root = artifact_root_for_db(self.path)
        final_path = root / storage_key
        temporary = root / f".{storage_key}.upload"
        try:
            temporary.write_bytes(data)
            os.replace(temporary, final_path)
            with self.connection() as conn:
                conn.execute("BEGIN IMMEDIATE")
                workspace_row = conn.execute("SELECT payload FROM spatial_workspace WHERE id=1").fetchone()
                workspace = json.loads(workspace_row[0]) if workspace_row else empty_workspace()
                # DocumentForm uploads before its first save so the extracted
                # text and original bytes can travel together in one submit.
                # The document mutation later records this ID. Only the
                # documented client-side draft prefix can be provisional;
                # arbitrary missing document IDs are rejected.
                if not any(item.get("id") == document_id for item in workspace.get("documents", [])) and not document_id.startswith("document-"):
                    raise HTTPException(404, "关联资料不存在，未保留来源文件。")
                previous = conn.execute(
                    "SELECT version FROM spatial_source_artifacts WHERE document_id=? ORDER BY version DESC LIMIT 1",
                    (document_id,),
                ).fetchone()
                version = int(previous[0]) + 1 if previous else 1
                conn.execute(
                    "INSERT INTO spatial_source_artifacts(artifact_id,document_id,filename,media_type,size,storage_key,uploaded_by,uploaded_at,version,status) VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (artifact_id, document_id, safe_name, str(media_type or "application/octet-stream")[:180], len(data), storage_key, str(actor or "本机负责人")[:180], now(), version, "stored"),
                )
                row = conn.execute("SELECT artifact_id,document_id,filename,media_type,size,storage_key,uploaded_by,uploaded_at,version,status FROM spatial_source_artifacts WHERE artifact_id=?", (artifact_id,)).fetchone()
            return self._artifact_row(row) or {}
        except Exception:
            final_path.unlink(missing_ok=True)
            temporary.unlink(missing_ok=True)
            raise

    def source_artifact(self, artifact_id: str, *, document_id: str | None = None) -> dict | None:
        with self.connection() as conn:
            query = "SELECT artifact_id,document_id,filename,media_type,size,storage_key,uploaded_by,uploaded_at,version,status FROM spatial_source_artifacts WHERE artifact_id=?"
            params: tuple[Any, ...] = (artifact_id,)
            if document_id:
                query += " AND document_id=?"
                params += (document_id,)
            row = conn.execute(query, params).fetchone()
        return self._artifact_row(row)

    def source_artifacts(self, *, document_ids: set[str] | None = None) -> list[dict]:
        """List retained source metadata for the current workspace scope."""
        query = "SELECT artifact_id,document_id,filename,media_type,size,storage_key,uploaded_by,uploaded_at,version,status FROM spatial_source_artifacts"
        params: tuple[Any, ...] = ()
        if document_ids is not None:
            if not document_ids:
                return []
            placeholders = ",".join("?" for _ in document_ids)
            query += f" WHERE document_id IN ({placeholders})"
            params = tuple(sorted(document_ids))
        query += " ORDER BY document_id ASC, version ASC, artifact_id ASC"
        with self.connection() as conn:
            rows = conn.execute(query, params).fetchall()
        return [self._artifact_row(row) or {} for row in rows]

    def _source_artifact_storage(self, artifact_id: str, *, document_id: str | None = None) -> tuple[dict, Path] | None:
        """Return metadata and a path using the private storage key."""
        with self.connection() as conn:
            query = "SELECT artifact_id,document_id,filename,media_type,size,storage_key,uploaded_by,uploaded_at,version,status FROM spatial_source_artifacts WHERE artifact_id=?"
            params: tuple[Any, ...] = (artifact_id,)
            if document_id:
                query += " AND document_id=?"
                params += (document_id,)
            row = conn.execute(query, params).fetchone()
        metadata = self._artifact_row(row)
        if not metadata or not row:
            return None
        root = artifact_root_for_db(self.path)
        path = _safe_artifact_storage_path(root, str(row[5]))
        if not path.is_file():
            raise HTTPException(410, "来源文件记录存在，但原始文件已不可用。")
        try:
            if path.stat().st_size != int(metadata["size"]):
                raise HTTPException(409, "来源文件大小与记录不一致，无法导出。")
        except OSError as error:
            raise HTTPException(410, "来源文件记录存在，但原始文件已不可用。") from error
        return metadata, path

    def source_artifact_path(self, artifact_id: str, *, document_id: str | None = None) -> tuple[dict, Path] | None:
        return self._source_artifact_storage(artifact_id, document_id=document_id)

    def restore_with_artifacts(
        self,
        *,
        expected_revision: int,
        request_id: str,
        candidate: dict,
        artifacts: list[dict],
        actor: str,
        migration_report: dict | None = None,
    ) -> dict:
        """Restore a workspace and its retained source bytes atomically.

        Archive files are staged under unique storage keys before the SQLite
        transaction commits.  Existing files remain untouched until the new
        workspace and metadata are durable; a failed restore removes only the
        staged files.  This keeps the JSON workspace and artifact ledger
        aligned even if extraction or validation fails midway.
        """
        candidate, migration = migrate_workspace_import(candidate)
        if migration_report is not None:
            migration = migration_report
        validate_workspace(candidate)
        # The endpoint performs the detailed archive checks.  Repeat the
        # cross-object invariant here so direct callers cannot create an
        # unreferenced or missing source artifact row.
        documents = {str(item["id"]): item for item in candidate.get("documents", [])}
        by_id: dict[str, dict] = {}
        for item in artifacts:
            metadata = dict(item.get("metadata") or {})
            artifact_id = str(metadata.get("artifactId") or "")
            document_id = str(metadata.get("documentId") or "")
            data = item.get("data")
            if not _ARTIFACT_ID_RE.fullmatch(artifact_id) or document_id not in documents or not isinstance(data, (bytes, bytearray)):
                raise ValueError("来源存档的 artifact 元数据无效。")
            if artifact_id in by_id:
                raise ValueError("来源存档包含重复的 artifact。")
            filename = _artifact_filename(str(metadata.get("filename") or ""))
            suffix = Path(filename).suffix.lower()
            if suffix not in ARTIFACT_SUFFIXES or filename != str(metadata.get("filename")):
                raise ValueError("来源存档包含不支持或不安全的文件名。")
            raw_size = metadata.get("size")
            if isinstance(raw_size, bool) or not isinstance(raw_size, int) or raw_size < 0 or raw_size > ARTIFACT_MAX_BYTES or raw_size != len(data):
                raise ValueError("来源存档的文件大小与元数据不一致。")
            raw_version = metadata.get("version", 1)
            if isinstance(raw_version, bool) or not isinstance(raw_version, int) or raw_version < 1:
                raise ValueError("来源存档的版本无效。")
            media_type = metadata.get("mediaType", "application/octet-stream")
            if not isinstance(media_type, str) or not media_type.strip() or len(media_type) > 180:
                raise ValueError("来源存档的媒体类型无效。")
            status = metadata.get("status", "stored")
            if status != "stored":
                raise ValueError("来源存档的状态无效。")
            if documents[document_id].get("sourceArtifactId") != artifact_id:
                raise ValueError("资料与来源文件 artifact 关联不一致。")
            source_version = documents[document_id].get("sourceArtifactVersion")
            if source_version is not None and source_version != raw_version:
                raise ValueError("资料与来源文件版本关联不一致。")
            by_id[artifact_id] = {"metadata": metadata, "data": bytes(data), "filename": filename, "suffix": suffix}
        referenced = {str(item.get("sourceArtifactId")) for item in documents.values() if item.get("sourceArtifactId")}
        if referenced != set(by_id):
            raise ValueError("工作区中的来源文件引用与 ZIP artifact 清单不一致。")

        fingerprint = hashlib.sha256(canonical({
            "type": "restore_workspace_zip", "workspace": candidate,
            "artifacts": [{"metadata": item["metadata"], "size": len(item["data"])} for item in by_id.values()],
        }).encode()).hexdigest()
        root = artifact_root_for_db(self.path)
        staged: list[tuple[Path, Path]] = []
        snapshot_files: list[Path] = []
        snapshot_dir: Path | None = None
        old_paths: list[Path] = []
        committed = False
        try:
            with self.connection() as conn:
                conn.execute("BEGIN IMMEDIATE")
                row = conn.execute("SELECT revision,payload FROM spatial_workspace WHERE id=1").fetchone()
                previous = conn.execute("SELECT request_hash FROM spatial_operations WHERE id=?", (request_id,)).fetchone()
                if previous:
                    if previous[0] != fingerprint:
                        raise HTTPException(409, "同一恢复操作标识不能用于不同存档。")
                    return json.loads(row[1])
                if row[0] != expected_revision:
                    raise HTTPException(409, "工作区已被其他请求更新，请刷新后重试；未覆盖已有数据。")
                before = json.loads(row[1])
                old_rows = conn.execute("SELECT artifact_id,document_id,filename,media_type,size,storage_key,uploaded_by,uploaded_at,version,status FROM spatial_source_artifacts").fetchall()
                old_artifacts: list[dict] = []
                old_total = 0
                for old_row in old_rows:
                    old_metadata = self._artifact_row(old_row) or {}
                    old_path = _safe_artifact_storage_path(root, str(old_row[5]))
                    if not old_path.is_file():
                        raise HTTPException(409, "当前工作区存在缺失的原始来源文件，无法建立恢复前快照。")
                    old_size = int(old_row[4])
                    if old_path.stat().st_size != old_size or old_size > ARTIFACT_MAX_BYTES:
                        raise HTTPException(409, "当前工作区原始来源文件与记录不一致，无法建立恢复前快照。")
                    old_total += old_size
                    if old_total > ARCHIVE_MAX_BYTES:
                        raise HTTPException(413, "恢复前来源快照不能超过 32 MB。")
                    old_paths.append(old_path)
                    old_artifacts.append({"metadata": old_metadata, "path": old_path})

                snapshot_id = uid("snapshot")
                snapshot_dir = root / ".snapshots" / snapshot_id
                snapshot_dir.mkdir(parents=True, exist_ok=False)
                snapshot_manifest: list[dict] = []
                for old_item in old_artifacts:
                    metadata = dict(old_item["metadata"])
                    artifact_id = str(metadata["artifactId"])
                    suffix = Path(str(metadata["filename"])).suffix.lower()
                    snapshot_storage_key = f"artifact-{artifact_id}{suffix}"
                    snapshot_path = _safe_artifact_storage_path(snapshot_dir, snapshot_storage_key)
                    shutil.copyfile(old_item["path"], snapshot_path)
                    snapshot_files.append(snapshot_path)
                    metadata["snapshotStorageKey"] = snapshot_storage_key
                    snapshot_manifest.append(metadata)

                # Stage and publish new files before changing the database. A
                # generated storage key is intentionally private; it never
                # trusts archive filenames for filesystem locations.
                restore_token = uuid4().hex
                for artifact_id, item in by_id.items():
                    storage_key = f"restore-{restore_token}-{artifact_id}{item['suffix']}"
                    temporary = root / f".{storage_key}.upload"
                    final = _safe_artifact_storage_path(root, storage_key)
                    temporary.write_bytes(item["data"])
                    os.replace(temporary, final)
                    staged.append((final, temporary))
                    item["storageKey"] = storage_key

                saved_at = now()
                conn.execute(
                    "INSERT INTO spatial_workspace_snapshots(snapshot_id,revision,saved_at,actor,reason,payload) VALUES(?,?,?,?,?,?)",
                    (snapshot_id, int(row[0]), saved_at, actor, "restore_workspace_zip_before_replace", json.dumps({"workspace": before, "artifacts": snapshot_manifest}, ensure_ascii=False)),
                )
                restored = copy.deepcopy(candidate)
                restored["revision"] = int(row[0]) + 1
                restored.setdefault("events", [])
                restored["events"] = [*restored["events"], dict(
                    id=uid("event"), at=saved_at, actor=actor, action="restore_workspace_zip",
                    objectId="workspace", summary=(
                        f"从 ZIP 存档恢复工作区与 {len(by_id)} 份原始来源文件；恢复前快照 {snapshot_id}"
                        + (f"；存档迁移 v{migration['fromVersion']}→v{migration['toVersion']}" if migration["migrated"] else "")
                    ),
                )]
                validate_workspace(restored)
                conn.execute("DELETE FROM spatial_source_artifacts")
                for artifact_id, item in by_id.items():
                    metadata = item["metadata"]
                    conn.execute(
                        "INSERT INTO spatial_source_artifacts(artifact_id,document_id,filename,media_type,size,storage_key,uploaded_by,uploaded_at,version,status) VALUES(?,?,?,?,?,?,?,?,?,?)",
                        (artifact_id, metadata["documentId"], item["filename"], str(metadata.get("mediaType") or "application/octet-stream")[:180], len(item["data"]), item["storageKey"], str(metadata.get("uploadedBy") or actor)[:180], str(metadata.get("uploadedAt") or saved_at), int(metadata.get("version") or 1), str(metadata.get("status") or "stored")[:40]),
                    )
                conn.execute("UPDATE spatial_workspace SET revision=?,payload=? WHERE id=1", (restored["revision"], json.dumps(restored, ensure_ascii=False)))
                conn.execute("INSERT INTO spatial_operations VALUES(?,?,?)", (request_id, fingerprint, restored["revision"]))
            committed = True
            # DB commit succeeded. Old files are now unreachable and can be
            # removed; retaining a failed-restore orphan would leak source
            # bytes indefinitely.
            new_paths = {item[0] for item in staged}
            for path in old_paths:
                if path not in new_paths:
                    try:
                        path.unlink(missing_ok=True)
                    except OSError:
                        # The DB now points at the new archive; an inability
                        # to clean an unreachable old file must not turn a
                        # successful restore into a misleading 500 response.
                        pass
            return restored
        except Exception:
            if not committed:
                for final, temporary in staged:
                    final.unlink(missing_ok=True)
                    temporary.unlink(missing_ok=True)
                if snapshot_dir is not None:
                    shutil.rmtree(snapshot_dir, ignore_errors=True)
            raise

    @staticmethod
    def _membership_record(row: tuple | sqlite3.Row, *, token: str | None = None) -> dict:
        keys = ("principal", "role", "tokenHash", "enabled", "createdAt", "updatedAt", "disabledAt", "lastSeenAt")
        record = dict(zip(keys, row))
        # Token digests are never returned to a browser; callers only receive
        # the one-time clear token attached explicitly by create_membership.
        record.pop("tokenHash", None)
        record["enabled"] = bool(record.get("enabled"))
        if token is not None:
            record["token"] = token
        return record

    def membership_for_token(self, token: str) -> dict | None:
        """Resolve a bearer token by digest without returning token material."""
        if not token:
            return None
        digest = token_digest(token)
        with self.connection() as conn:
            row = conn.execute(
                "SELECT principal,role,token_hash,enabled,created_at,updated_at,disabled_at,last_seen_at FROM spatial_memberships WHERE token_hash=?",
                (digest,),
            ).fetchone()
            if row and row[3]:
                seen_at = now()
                conn.execute("UPDATE spatial_memberships SET last_seen_at=? WHERE token_hash=? AND enabled=1", (seen_at, digest))
                row = conn.execute(
                    "SELECT principal,role,token_hash,enabled,created_at,updated_at,disabled_at,last_seen_at FROM spatial_memberships WHERE token_hash=?",
                    (digest,),
                ).fetchone()
        return self._membership_record(row) if row else None

    def memberships(self) -> list[dict]:
        with self.connection() as conn:
            rows = conn.execute(
                "SELECT principal,role,token_hash,enabled,created_at,updated_at,disabled_at,last_seen_at FROM spatial_memberships ORDER BY principal ASC"
            ).fetchall()
        return [self._membership_record(row) for row in rows]

    def create_membership(self, *, principal: str, role: str, request_id: str) -> dict:
        """Create one high-entropy token and return it exactly once.

        The idempotency record stores only a request hash and principal.  A
        replay therefore cannot recover the original clear token.
        """
        principal = str(principal or "").strip()
        role = str(role or "").strip()
        request_id = str(request_id or "").strip()
        if not principal or len(principal) > 180:
            raise ValueError("成员标识不能为空且不能超过 180 个字符。")
        if any(ord(ch) < 32 for ch in principal):
            raise ValueError("成员标识不能包含控制字符。")
        if role not in WORKSPACE_ROLES:
            raise ValueError("成员角色配置无效。")
        if len(request_id) < 8 or len(request_id) > 128:
            raise ValueError("成员请求标识长度必须在 8 到 128 个字符之间。")
        request_hash = hashlib.sha256(canonical({"principal": principal, "role": role}).encode()).hexdigest()
        with self.connection() as conn:
            conn.execute("BEGIN IMMEDIATE")
            previous = conn.execute(
                "SELECT request_hash,principal FROM spatial_membership_operations WHERE request_id=?",
                (request_id,),
            ).fetchone()
            if previous:
                if previous[0] != request_hash:
                    raise HTTPException(409, "同一成员请求标识不能用于不同成员或角色。")
                row = conn.execute(
                    "SELECT principal,role,token_hash,enabled,created_at,updated_at,disabled_at,last_seen_at FROM spatial_memberships WHERE principal=?",
                    (previous[1],),
                ).fetchone()
                if row is None:
                    raise HTTPException(409, "成员请求记录与成员台账不一致。")
                return {**self._membership_record(row), "token": None, "tokenIssued": False}
            existing = conn.execute("SELECT 1 FROM spatial_memberships WHERE principal=?", (principal,)).fetchone()
            if existing:
                raise HTTPException(409, "成员标识已经存在；请使用新的成员标识。")
            clear_token = secrets.token_urlsafe(32)
            timestamp = now()
            conn.execute(
                "INSERT INTO spatial_memberships(principal,role,token_hash,enabled,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                (principal, role, token_digest(clear_token), 1, timestamp, timestamp),
            )
            conn.execute(
                "INSERT INTO spatial_membership_operations(request_id,request_hash,principal,created_at) VALUES(?,?,?,?)",
                (request_id, request_hash, principal, timestamp),
            )
            row = conn.execute(
                "SELECT principal,role,token_hash,enabled,created_at,updated_at,disabled_at,last_seen_at FROM spatial_memberships WHERE principal=?",
                (principal,),
            ).fetchone()
        return {**self._membership_record(row, token=clear_token), "tokenIssued": True}

    def disable_membership(self, principal: str) -> dict:
        principal = str(principal or "").strip()
        if not principal:
            raise ValueError("成员标识不能为空。")
        with self.connection() as conn:
            conn.execute("BEGIN IMMEDIATE")
            row = conn.execute(
                "SELECT principal,role,token_hash,enabled,created_at,updated_at,disabled_at,last_seen_at FROM spatial_memberships WHERE principal=?",
                (principal,),
            ).fetchone()
            if row is None:
                raise HTTPException(404, "成员不存在。")
            if row[3]:
                disabled_at = now()
                conn.execute("UPDATE spatial_memberships SET enabled=0,updated_at=?,disabled_at=COALESCE(disabled_at,?) WHERE principal=?", (disabled_at, disabled_at, principal))
            row = conn.execute(
                "SELECT principal,role,token_hash,enabled,created_at,updated_at,disabled_at,last_seen_at FROM spatial_memberships WHERE principal=?",
                (principal,),
            ).fetchone()
        return self._membership_record(row)

    def mutate(self, expected: int, request_id: str, action: str, payload: dict, actor: str) -> dict:
        fingerprint = hashlib.sha256(canonical(dict(type=action, payload=payload)).encode()).hexdigest()
        with self.connection() as conn:
            conn.execute("BEGIN IMMEDIATE")
            row = conn.execute("SELECT revision,payload FROM spatial_workspace WHERE id=1").fetchone()
            previous = conn.execute("SELECT request_hash FROM spatial_operations WHERE id=?", (request_id,)).fetchone()
            if previous:
                if previous[0] != fingerprint:
                    raise HTTPException(409, "同一操作标识不能用于不同请求。")
                return json.loads(row[1])
            if row[0] != expected:
                raise HTTPException(409, "工作区已被其他请求更新，请刷新后重试；未覆盖已有数据。")
            result = apply_operation(json.loads(row[1]), action, payload, actor)
            conn.execute("UPDATE spatial_workspace SET revision=?,payload=? WHERE id=1", (result["revision"], json.dumps(result, ensure_ascii=False)))
            conn.execute("INSERT INTO spatial_operations VALUES(?,?,?)", (request_id, fingerprint, result["revision"]))
            notification = mutation_notification(result, action, payload, actor, request_id)
            if notification:
                conn.execute(
                    """INSERT INTO spatial_outbox(
                       event_id,idempotency_key,event_type,aggregate_type,aggregate_id,
                       plan_id,audience,title,summary,actor,workspace_revision,
                       source_event_id,status,created_at,acknowledged_at,acknowledged_by
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        notification["eventId"], notification["idempotencyKey"], notification["eventType"],
                        notification["aggregateType"], notification["aggregateId"], notification.get("planId"),
                        notification["audience"], notification["title"], notification["summary"], notification["actor"],
                        notification["workspaceRevision"], notification.get("sourceEventId"), notification["status"],
                        notification["createdAt"], notification.get("acknowledgedAt"), notification.get("acknowledgedBy"),
                    ),
                )
            return result
    def save_analysis_run(self, payload: dict) -> None:
        trace = payload.get("trace") or {}
        run_id = str(trace.get("runId") or uid("analysis"))
        with self.connection() as conn:
            conn.execute(
                """INSERT INTO spatial_analysis_runs(run_id,plan_id,plan_revision,generated_at,degraded,payload)
                   VALUES(?,?,?,?,?,?) ON CONFLICT(run_id) DO UPDATE SET
                   plan_revision=excluded.plan_revision,generated_at=excluded.generated_at,
                   degraded=excluded.degraded,payload=excluded.payload""",
                (run_id, str(payload.get("planId", "")), int(payload.get("planRevision", 0)), str(trace.get("generatedAt") or now()), 1 if payload.get("degraded") else 0, json.dumps(payload, ensure_ascii=False)),
            )
    def analysis_runs(self, plan_id: str, limit: int = 20) -> list[dict]:
        with self.connection() as conn:
            rows = conn.execute("SELECT payload FROM spatial_analysis_runs WHERE plan_id=? ORDER BY generated_at DESC LIMIT ?", (plan_id, max(1, min(limit, 100)))).fetchall()
        return [json.loads(row[0]) for row in rows]
    def save_evaluation_run(self, payload: dict) -> dict:
        evaluation_id = str(payload.get("evaluationId") or uid("evaluation"))
        record = {**payload, "evaluationId": evaluation_id}
        with self.connection() as conn:
            conn.execute(
                """INSERT INTO spatial_evaluation_runs(evaluation_id,plan_id,plan_revision,evaluated_at,basis,status,payload)
                   VALUES(?,?,?,?,?,?,?) ON CONFLICT(evaluation_id) DO UPDATE SET
                   evaluated_at=excluded.evaluated_at,basis=excluded.basis,status=excluded.status,payload=excluded.payload""",
                (evaluation_id, str(record.get("planId", "")), int(record.get("planRevision", 0)), str(record.get("evaluatedAt") or now()), str(record.get("basis", "")), str(record.get("status", "completed")), json.dumps(record, ensure_ascii=False)),
            )
        return record
    def evaluation_runs(self, plan_id: str, limit: int = 20) -> list[dict]:
        with self.connection() as conn:
            rows = conn.execute("SELECT payload FROM spatial_evaluation_runs WHERE plan_id=? ORDER BY evaluated_at DESC LIMIT ?", (plan_id, max(1, min(limit, 100)))).fetchall()
        return [json.loads(row[0]) for row in rows]
    def begin_job(self, *, job_id: str, kind: str, plan_id: str | None, idempotency_key: str, context: dict | None = None) -> dict:
        """Create or resume a durable job record.

        ``context`` contains only bounded replay inputs (for example an expiry
        ``asOf`` timestamp).  It lets duplicate queued requests compare their
        inputs without pretending that a process-local runner is durable.
        """
        with self.connection() as conn:
            # Serialize the idempotency check and insert.  Without an
            # immediate transaction two browser clicks can both observe a
            # missing key and race into the UNIQUE constraint, even though
            # their eventual job state is supposed to be the same record.
            conn.execute("BEGIN IMMEDIATE")
            row = conn.execute("SELECT * FROM spatial_jobs WHERE idempotency_key=?", (idempotency_key,)).fetchone()
            if row:
                return self._job_row(row)
            created = now()
            encoded_context = json.dumps(context, ensure_ascii=False, separators=(",", ":")) if context else None
            conn.execute("INSERT INTO spatial_jobs(job_id,kind,plan_id,idempotency_key,status,created_at,request_payload) VALUES(?,?,?,?,?,?,?)",
                         (job_id, kind, plan_id, idempotency_key, "queued", created, encoded_context))
            return self._job_row(conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone())
    @staticmethod
    def _job_row(row: sqlite3.Row | tuple) -> dict:
        keys = ("jobId", "kind", "planId", "idempotencyKey", "status", "createdAt", "startedAt", "finishedAt", "result", "error", "context")
        item = dict(zip(keys, row))
        if item.get("result"):
            item["result"] = json.loads(item["result"])
        if item.get("context"):
            item["context"] = json.loads(item["context"])
        return item
    def set_job_running(self, job_id: str) -> dict:
        with self.connection() as conn:
            row = conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone()
            if not row:
                raise HTTPException(404, "后台任务不存在。")
            if row[4] != "queued":
                if row[4] in {"running"}:
                    return self._job_row(row)
                raise HTTPException(409, f"任务已经进入终态 {row[4]}，不能重新运行。")
            conn.execute("UPDATE spatial_jobs SET status='running',started_at=COALESCE(started_at,?),error=NULL WHERE job_id=? AND status='queued'", (now(), job_id))
            row = conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone()
        return self._job_row(row)
    def claim_next_job(self, kinds: tuple[str, ...] = ("analysis", "expiry_recheck")) -> dict | None:
        """Atomically claim the oldest queued job for an external Worker.

        SQLite's immediate transaction is enough for the single-node local
        deployment: two worker processes cannot claim the same row, and a
        cancelled/terminal row is left immutable.  A worker crash can leave a
        ``running`` row behind, so an expired lease is moved back to ``queued``
        in this same transaction.  The previous ``started_at`` value is the
        ownership token; callers must pass it to :meth:`finish_job` so an old
        worker cannot publish a result after a new worker reclaimed the job.
        """
        if not kinds:
            return None
        placeholders = ",".join("?" for _ in kinds)
        with self.connection() as conn:
            conn.execute("BEGIN IMMEDIATE")
            # Only a durable Worker opts into this method.  Requeueing is
            # deliberately lease-based rather than startup-based: a worker
            # that is still processing a slow model call keeps its ownership
            # until the configured timeout elapses.
            lease_cutoff = datetime.now(timezone.utc) - timedelta(seconds=worker_lease_seconds())
            cutoff_text = lease_cutoff.isoformat().replace("+00:00", "Z")
            recovery_message = (
                "Worker 租约已过期，任务已重新入队；旧 Worker 不再拥有写入权限。"
            )
            conn.execute(
                f"""UPDATE spatial_jobs
                    SET status='queued', started_at=NULL, finished_at=NULL,
                        error=?
                    WHERE status='running' AND kind IN ({placeholders})
                      AND started_at IS NOT NULL
                      AND julianday(started_at) < julianday(?)""",
                (recovery_message, *kinds, cutoff_text),
            )
            row = conn.execute(
                f"SELECT * FROM spatial_jobs WHERE status='queued' AND kind IN ({placeholders}) ORDER BY created_at ASC, job_id ASC LIMIT 1",
                tuple(kinds),
            ).fetchone()
            if not row:
                return None
            job_id = row[0]
            updated = conn.execute(
                "UPDATE spatial_jobs SET status='running', started_at=COALESCE(started_at,?) WHERE job_id=? AND status='queued'",
                (now(), job_id),
            )
            if updated.rowcount != 1:
                return None
            row = conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone()
        return self._job_row(row)

    def finish_job(
        self,
        job_id: str,
        *,
        status: str,
        result: dict | None = None,
        error: str | None = None,
        expected_started_at: str | None = None,
    ) -> dict:
        """Persist a terminal result, optionally guarded by a Worker lease.

        ``expected_started_at`` is intentionally optional for compatibility
        with process-local handlers.  A durable Worker passes the timestamp it
        received from :meth:`claim_next_job`; if another Worker has reclaimed
        the row, the compare-and-set fails and the stale process cannot mark
        the newer attempt failed or succeeded.
        """
        if status not in {"succeeded", "failed", "cancelled"}:
            raise ValueError("无效的任务终态。")
        with self.connection() as conn:
            row = conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone()
            if not row:
                raise HTTPException(404, "后台任务不存在。")
            current = row[4]
            if current in {"succeeded", "failed", "cancelled"}:
                if current == status:
                    return self._job_row(row)
                raise HTTPException(409, f"任务已经进入终态 {current}，不能改成 {status}。")
            if expected_started_at is not None and row[6] != expected_started_at:
                raise HTTPException(409, "Worker 租约已失效，任务已由其他 Worker 接管。")
            if status == "succeeded" and current != "running":
                raise HTTPException(409, "只有运行中的任务可以完成。")
            if expected_started_at is None:
                conn.execute("UPDATE spatial_jobs SET status=?,finished_at=?,result_payload=?,error=? WHERE job_id=? AND status IN ('queued','running')",
                             (status, now(), json.dumps(result, ensure_ascii=False) if result is not None else None, error, job_id))
            else:
                updated = conn.execute("UPDATE spatial_jobs SET status=?,finished_at=?,result_payload=?,error=? WHERE job_id=? AND status='running' AND started_at=?",
                                       (status, now(), json.dumps(result, ensure_ascii=False) if result is not None else None, error, job_id, expected_started_at))
                if updated.rowcount != 1:
                    raise HTTPException(409, "Worker 租约已失效，任务已由其他 Worker 接管。")
            row = conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone()
        return self._job_row(row)
    def cancel_job(self, job_id: str, *, reason: str = "由负责人取消") -> dict:
        """Cancel a queued/running job and keep terminal states immutable."""
        with self.connection() as conn:
            row = conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone()
            if not row:
                raise HTTPException(404, "后台任务不存在。")
            current = row[4]
            if current == "cancelled":
                return self._job_row(row)
            if current in {"succeeded", "failed"}:
                raise HTTPException(409, f"任务已经进入终态 {current}，不能取消。")
            conn.execute(
                "UPDATE spatial_jobs SET status='cancelled',finished_at=?,error=? WHERE job_id=? AND status IN ('queued','running')",
                (now(), reason[:500], job_id),
            )
            row = conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone()
        return self._job_row(row)
    def job(self, job_id: str) -> dict | None:
        with self.connection() as conn:
            row = conn.execute("SELECT * FROM spatial_jobs WHERE job_id=?", (job_id,)).fetchone()
        return self._job_row(row) if row else None
    def jobs(self, plan_id: str, limit: int = 20) -> list[dict]:
        with self.connection() as conn:
            rows = conn.execute("SELECT * FROM spatial_jobs WHERE plan_id=? ORDER BY created_at DESC LIMIT ?", (plan_id, max(1, min(limit, 100)))).fetchall()
        return [self._job_row(row) for row in rows]
    def expiry_recheck(self, *, request_id: str, as_of: str, actor: str, worker_name: str = "not-running") -> dict:
        """Run and persist a manual expiry scan using the durable job ledger.

        This is intentionally synchronous.  The returned ``worker`` marker is
        ``not-running`` and makes the boundary explicit until a real scheduler
        and outbox are introduced.  Idempotency is keyed by the caller's
        request ID, so retrying a browser request returns the original report.
        """
        try:
            dt(as_of)
        except (TypeError, ValueError) as error:
            raise ValueError("重检时间必须是带时区的 ISO 时间。") from error
        job_id = f"expiry-recheck:{request_id}"
        context = {"asOf": as_of}
        job = self.begin_job(job_id=job_id, kind="expiry_recheck", plan_id=None, idempotency_key=request_id, context=context)
        stored_context = job.get("context") or {}
        if stored_context.get("asOf") and dt(stored_context["asOf"]) != dt(as_of):
            raise HTTPException(409, "同一 requestId 不能用于不同的重检时间。")
        if job["status"] == "succeeded" and job.get("result"):
            stored_as_of = job["result"].get("asOf")
            if stored_as_of and dt(stored_as_of) != dt(as_of):
                raise HTTPException(409, "同一 requestId 不能用于不同的重检时间。")
            return job
        if job["status"] in {"failed", "cancelled"}:
            raise HTTPException(409, f"该重检任务已处于 {job['status']} 终态，请更换 requestId 重试。")
        self.set_job_running(job_id)
        try:
            report = expiry_recheck_report(self.read(), as_of, worker_name=worker_name)
            report.update({"requestId": request_id, "jobId": job_id, "generatedAt": now(), "actor": actor})
            return self.finish_job(job_id, status="succeeded", result=report)
        except Exception as error:
            current = self.job(job_id)
            if current and current["status"] in {"queued", "running"}:
                self.finish_job(job_id, status="failed", error=f"{type(error).__name__}: {str(error)[:500]}")
            raise
    def restore(self, *, expected_revision: int, request_id: str, candidate: dict, actor: str, migration_report: dict | None = None) -> dict:
        """Replace the workspace only after validation and a durable pre-restore snapshot.

        The imported revision is deliberately ignored.  Restore is an ordinary
        optimistic-concurrency mutation: a stale browser cannot overwrite a
        newer server state, and retrying the same request is idempotent.
        """
        candidate, migration = migrate_workspace_import(candidate)
        # The HTTP handler normalizes once so it can return the report to the
        # browser.  Carry that report into the durable event instead of
        # losing the provenance note on the second idempotent normalization.
        if migration_report is not None:
            migration = migration_report
        fingerprint = hashlib.sha256(canonical(dict(type="restore_workspace", workspace=candidate)).encode()).hexdigest()
        validate_workspace(candidate)
        with self.connection() as conn:
            conn.execute("BEGIN IMMEDIATE")
            row = conn.execute("SELECT revision,payload FROM spatial_workspace WHERE id=1").fetchone()
            previous = conn.execute("SELECT request_hash FROM spatial_operations WHERE id=?", (request_id,)).fetchone()
            if previous:
                if previous[0] != fingerprint:
                    raise HTTPException(409, "同一恢复操作标识不能用于不同存档。")
                return json.loads(row[1])
            if row[0] != expected_revision:
                raise HTTPException(409, "工作区已被其他请求更新，请刷新后重试；未覆盖已有数据。")
            before = json.loads(row[1])
            snapshot_id = uid("snapshot")
            saved_at = now()
            conn.execute(
                "INSERT INTO spatial_workspace_snapshots(snapshot_id,revision,saved_at,actor,reason,payload) VALUES(?,?,?,?,?,?)",
                (snapshot_id, int(row[0]), saved_at, actor, "restore_workspace_before_replace", json.dumps(before, ensure_ascii=False)),
            )
            restored = copy.deepcopy(candidate)
            restored["revision"] = int(row[0]) + 1
            restored.setdefault("events", [])
            restored["events"] = [*restored["events"], dict(
                id=uid("event"), at=saved_at, actor=actor, action="restore_workspace",
                objectId="workspace", summary=(
                    f"从服务端恢复文件替换工作区；恢复前快照 {snapshot_id}"
                    + (f"；存档迁移 v{migration['fromVersion']}→v{migration['toVersion']}" if migration["migrated"] else "")
                ),
            )]
            validate_workspace(restored)
            conn.execute("UPDATE spatial_workspace SET revision=?,payload=? WHERE id=1", (restored["revision"], json.dumps(restored, ensure_ascii=False)))
            conn.execute("INSERT INTO spatial_operations VALUES(?,?,?)", (request_id, fingerprint, restored["revision"]))
            return restored
    def snapshots(self, limit: int = 20) -> list[dict]:
        with self.connection() as conn:
            rows = conn.execute(
                "SELECT snapshot_id,revision,saved_at,actor,reason FROM spatial_workspace_snapshots ORDER BY saved_at DESC LIMIT ?",
                (max(1, min(limit, 100)),),
            ).fetchall()
        return [dict(snapshotId=row[0], revision=row[1], savedAt=row[2], actor=row[3], reason=row[4]) for row in rows]
    def snapshot(self, snapshot_id: str) -> dict | None:
        with self.connection() as conn:
            row = conn.execute("SELECT snapshot_id,revision,saved_at,actor,reason,payload FROM spatial_workspace_snapshots WHERE snapshot_id=?", (snapshot_id,)).fetchone()
        if not row:
            return None
        payload = json.loads(row[5])
        if isinstance(payload, dict) and isinstance(payload.get("workspace"), dict):
            artifacts = []
            for item in payload.get("artifacts", []):
                if isinstance(item, dict):
                    visible = dict(item)
                    visible.pop("snapshotStorageKey", None)
                    artifacts.append(visible)
            return dict(snapshotId=row[0], revision=row[1], savedAt=row[2], actor=row[3], reason=row[4], workspace=payload["workspace"], artifacts=artifacts)
        return dict(snapshotId=row[0], revision=row[1], savedAt=row[2], actor=row[3], reason=row[4], workspace=payload, artifacts=[])

    def snapshot_artifact_path(self, snapshot_id: str, artifact_id: str) -> tuple[dict, Path] | None:
        """Resolve a retained pre-restore artifact copy without exposing paths."""
        with self.connection() as conn:
            row = conn.execute("SELECT payload FROM spatial_workspace_snapshots WHERE snapshot_id=?", (snapshot_id,)).fetchone()
        if not row:
            return None
        payload = json.loads(row[0])
        if not isinstance(payload, dict) or not isinstance(payload.get("workspace"), dict):
            return None
        item = next((entry for entry in payload.get("artifacts", []) if isinstance(entry, dict) and entry.get("artifactId") == artifact_id), None)
        storage_key = item.get("snapshotStorageKey") if item else None
        if not item or not isinstance(storage_key, str):
            return None
        snapshot_root = artifact_root_for_db(self.path) / ".snapshots" / snapshot_id
        path = _safe_artifact_storage_path(snapshot_root, storage_key)
        if not path.is_file():
            raise HTTPException(410, "恢复前快照的原始文件已不可用。")
        return {key: value for key, value in item.items() if key != "snapshotStorageKey"}, path

    @staticmethod
    def _notification_row(row: sqlite3.Row | tuple) -> dict:
        keys = (
            "eventId", "idempotencyKey", "eventType", "aggregateType", "aggregateId", "planId",
            "audience", "title", "summary", "actor", "workspaceRevision", "sourceEventId",
            "status", "createdAt", "acknowledgedAt", "acknowledgedBy",
        )
        return dict(zip(keys, row))

    def notifications(self, *, status: str = "pending", audience: str | None = None, plan_id: str | None = None, limit: int = 50) -> list[dict]:
        """List durable local notifications without claiming external delivery.

        ``status=all`` is read-only history.  Audience matching is exact or
        prefix-based for task-owner entries (``任务负责人:...``), so a browser
        can request the current role inbox without exposing another audience's
        events by accident.
        """
        if status not in {"pending", "acknowledged", "all"}:
            raise ValueError("通知状态只能是 pending、acknowledged 或 all。")
        clauses: list[str] = []
        params: list[Any] = []
        if status != "all":
            clauses.append("status=?")
            params.append(status)
        if audience:
            clauses.append("(audience=? OR audience LIKE ?)")
            params.extend([audience, f"{audience}:%"])
        if plan_id:
            clauses.append("plan_id=?")
            params.append(plan_id)
        where = f" WHERE {' AND '.join(clauses)}" if clauses else ""
        params.append(max(1, min(limit, 100)))
        with self.connection() as conn:
            rows = conn.execute(
                f"SELECT event_id,idempotency_key,event_type,aggregate_type,aggregate_id,plan_id,audience,title,summary,actor,workspace_revision,source_event_id,status,created_at,acknowledged_at,acknowledged_by FROM spatial_outbox{where} ORDER BY created_at DESC,event_id DESC LIMIT ?",
                tuple(params),
            ).fetchall()
        return [self._notification_row(row) for row in rows]

    def acknowledge_notification(self, event_id: str, actor: str, *, allowed_audience: str | None = None) -> dict:
        """Acknowledge one local inbox item; repeated acknowledgement is safe."""
        with self.connection() as conn:
            conn.execute("BEGIN IMMEDIATE")
            row = conn.execute(
                "SELECT event_id,idempotency_key,event_type,aggregate_type,aggregate_id,plan_id,audience,title,summary,actor,workspace_revision,source_event_id,status,created_at,acknowledged_at,acknowledged_by FROM spatial_outbox WHERE event_id=?",
                (event_id,),
            ).fetchone()
            if not row:
                raise HTTPException(404, "通知不存在。")
            if allowed_audience and not (row[6] == allowed_audience or row[6].startswith(f"{allowed_audience}:")):
                raise HTTPException(403, "当前成员不能确认其他责任人的通知。")
            if row[12] == "pending":
                acknowledged_at = now()
                conn.execute(
                    "UPDATE spatial_outbox SET status='acknowledged',acknowledged_at=?,acknowledged_by=? WHERE event_id=? AND status='pending'",
                    (acknowledged_at, actor, event_id),
                )
            row = conn.execute(
                "SELECT event_id,idempotency_key,event_type,aggregate_type,aggregate_id,plan_id,audience,title,summary,actor,workspace_revision,source_event_id,status,created_at,acknowledged_at,acknowledged_by FROM spatial_outbox WHERE event_id=?",
                (event_id,),
            ).fetchone()
        return self._notification_row(row)


_store: WorkspaceStore | None = None

def store() -> WorkspaceStore:
    global _store
    if _store is None:
        _store = WorkspaceStore(os.getenv("AEGIS_SPATIAL_DB_PATH", str(ROOT / "aegis-spatial.db")))
    return _store


def _configured_identity() -> dict[str, str]:
    role = os.getenv("AEGIS_WORKSPACE_ROLE", "safety_owner").strip()
    if role not in WORKSPACE_ROLES:
        raise HTTPException(500, "工作区角色配置无效。")
    return {"principal": os.getenv("AEGIS_WORKSPACE_OWNER", "本机负责人"), "role": role, "auth": "local"}


def _bearer_token(request: Request) -> str | None:
    authorization = request.headers.get("authorization", "").strip()
    if not authorization:
        return None
    scheme, separator, value = authorization.partition(" ")
    if not separator or scheme.lower() != "bearer" or not value.strip():
        raise HTTPException(401, "工作区访问令牌格式无效。")
    return value.strip()


def _remember_identity(request: Request, identity: dict[str, str]) -> str:
    # Keep the public dependency return value as the principal for compatibility
    # with existing handlers, while retaining the server-resolved role beside
    # it.  The JSON body cannot overwrite either value.
    request.state.workspace_identity = identity
    return identity["principal"]


def local_owner(request: Request) -> str:
    """Authenticate the browser request and return its server-side principal.

    ``AEGIS_WORKSPACE_TOKEN`` remains a compatible static owner credential.
    Additional members are resolved from the hashed token table; raw bearer
    tokens never enter the database, logs, mutation payloads or response
    bodies except for the one-time create response.  With no configured
    credential, the existing localhost-only owner mode remains available.
    """
    configured_token = os.getenv("AEGIS_WORKSPACE_TOKEN", "")
    token = _bearer_token(request)
    if token:
        if configured_token and hmac.compare_digest(token, configured_token):
            return _remember_identity(request, _configured_identity() | {"auth": "env-token"})
        member = store().membership_for_token(token)
        if member:
            if not member.get("enabled"):
                raise HTTPException(403, "该工作区成员已被禁用。")
            identity = {"principal": str(member["principal"]), "role": str(member["role"]), "auth": "membership"}
            if identity["role"] not in WORKSPACE_ROLES:
                raise HTTPException(500, "成员台账中的角色配置无效。")
            return _remember_identity(request, identity)
        raise HTTPException(401, "工作区访问令牌无效。")
    if configured_token:
        raise HTTPException(401, "工作区访问令牌无效。")

    # Preserve local development behavior only when no credential is configured.
    peer = request.client.host if request.client else ""
    host = request.url.hostname
    origin = request.headers.get("origin")
    if peer not in {"127.0.0.1", "::1", "testclient"} or host not in {"localhost", "127.0.0.1", "::1", "testserver"}:
        raise HTTPException(403, "未配置访问令牌时仅允许本机访问。")
    if origin and urlparse(origin).hostname not in {"localhost", "127.0.0.1", "::1"}:
        raise HTTPException(403, "拒绝非本机网页来源。")
    return _remember_identity(request, _configured_identity())


def local_role(request: Request | None = None) -> str:
    """Return the role resolved by local_owner, never a body-provided role."""
    if request is not None:
        identity = getattr(request.state, "workspace_identity", None)
        if not identity:
            raise HTTPException(401, "请求身份尚未完成认证。")
        return str(identity["role"])
    # Direct Python callers have no Request context; preserve the configured
    # local-owner fallback for domain tests and worker code.
    return _configured_identity()["role"]


def local_principal(request: Request | None = None) -> str:
    if request is not None:
        identity = getattr(request.state, "workspace_identity", None)
        if not identity:
            raise HTTPException(401, "请求身份尚未完成认证。")
        return str(identity["principal"])
    return _configured_identity()["principal"]


def require_workspace_roles(*roles: str, request: Request | None = None) -> None:
    if local_role(request) not in set(roles):
        raise HTTPException(403, "当前工作区身份没有执行该成员管理操作的权限。")

class Mutation(StrictModel):
    expectedRevision: int = Field(ge=0)
    requestId: str = Field(min_length=8, max_length=128)
    type: str = Field(max_length=60)
    payload: dict

class RestoreRequest(StrictModel):
    expectedRevision: int = Field(ge=0)
    requestId: str = Field(min_length=8, max_length=128)
    confirm: Literal["RESTORE"]
    workspace: dict


class MemberCreateRequest(StrictModel):
    principal: str = Field(min_length=1, max_length=180)
    role: Literal["instructor", "researcher", "safety_owner", "ehs_manager", "administrator"]
    requestId: str = Field(default_factory=lambda: uid("member"), min_length=8, max_length=128)


class MemberDisableRequest(StrictModel):
    requestId: str = Field(default_factory=lambda: uid("member-disable"), min_length=8, max_length=128)


class ExpiryRecheckRequest(StrictModel):
    """A manually-triggered, durable expiry scan.

    ``asOf`` is injectable so an operator can reproduce a review at a known
    time.  The endpoint records the result as a durable job; it does not claim
    that a background worker or scheduler is running.
    """

    requestId: str = Field(default_factory=lambda: uid("expiry"), min_length=8, max_length=128)
    asOf: str | None = None

router = APIRouter(prefix="/api/spatial", tags=["web-workspace"])

# Keep this contract explicit instead of silently returning an empty graph for
# a misspelled type. Chemical and control nodes are source-anchored
# projections of saved plan text; they remain unknown until evidence is
# reviewed and are never treated as confirmed safety facts.
GRAPH_NODE_TYPES = frozenset({"plan", "step", "chemical", "control", "equipment", "document", "task", "check", "approval"})

@router.get("")
def get_workspace(actor: str = Depends(local_owner)):
    return envelope(store().read())

@router.get("/sources")
def get_source_registry(actor: str = Depends(local_owner)):
    """Return public source metadata used by equipment category citations."""
    sources = read_source_registry()
    return {"sources": sources, "scope": "公开类别与管理机制参考，不是机构资产台账"}


@router.get("/members")
def list_members(request: Request, actor: str = Depends(local_owner)):
    """List workspace members without exposing token digests or clear tokens."""
    require_workspace_roles("safety_owner", "administrator", request=request)
    return {
        "members": store().memberships(),
        "scope": "工作区成员台账；token 只在新增成功响应中显示一次。",
    }


@router.post("/members")
def create_member(body: MemberCreateRequest, request: Request, actor: str = Depends(local_owner)):
    """Create a member credential under an owner/admin identity.

    The clear bearer token is intentionally present only in this response;
    retries return the public member record with ``token: null``.
    """
    require_workspace_roles("safety_owner", "administrator", request=request)
    try:
        member = store().create_membership(principal=body.principal, role=body.role, request_id=body.requestId)
    except (ValueError, TypeError) as error:
        raise HTTPException(422, str(error)) from error
    return {
        "member": member,
        "tokenShownOnce": bool(member.get("tokenIssued")),
        "scope": "token 仅本次响应可见；请由负责人安全保存。",
    }


@router.post("/members/{principal}/disable")
def disable_member(principal: str, request: Request, body: MemberDisableRequest | None = None, actor: str = Depends(local_owner)):
    """Disable a member token; repeat calls are idempotent."""
    require_workspace_roles("safety_owner", "administrator", request=request)
    if principal == local_principal(request):
        raise HTTPException(409, "不能禁用当前登录成员；请由另一位管理员执行。")
    try:
        member = store().disable_membership(principal)
    except (ValueError, TypeError) as error:
        raise HTTPException(422, str(error)) from error
    return {"member": member, "requestId": body.requestId if body else None, "scope": "禁用仅撤销该 token 的后续访问，不删除历史审计记录。"}


@router.get("/plans/{plan_id}/approval")
def get_plan_approval(plan_id: str, request: Request, actor: str = Depends(local_owner)):
    """Return the server-derived approval chain for one current plan."""
    w = store().read()
    plan = next((item for item in w.get("plans", []) if item.get("id") == plan_id), None)
    if plan is None:
        raise HTTPException(404, "计划不存在。")
    checks = [item for item in check_workspace(w) if item.get("planId") == plan_id]
    state = approval_state(w, plan, checks)
    return {"planId": plan_id, "planRevision": plan.get("revision"), "role": local_role(request), "status": plan_status(w, plan, checks), **state}

@router.get("/graph")
def get_relationship_graph(
    planId: str | None = None,
    labId: str | None = None,
    types: str | None = None,
    nodeId: str | None = None,
    pathFrom: str | None = None,
    pathTo: str | None = None,
    dateFrom: str | None = None,
    dateTo: str | None = None,
    relations: str | None = None,
    hops: int = 0,
    actor: str = Depends(local_owner),
):
    """Return a server-derived relationship subgraph for the current snapshot.

    The browser graph is a projection, while this endpoint is the queryable
    source of truth: every edge is derived from a plan, step, reservation,
    document, task or deterministic check in the workspace payload.  ``types``
    filters node types (comma separated), ``nodeId`` plus ``hops`` requests a
    local neighbourhood, and date bounds retain reservations intersecting the
    requested interval.  Unknown or missing evidence is represented as a node
    state and is never promoted to a pass result.
    """
    if hops < 0 or hops > 2:
        raise HTTPException(422, "hops 必须在 0 到 2 之间。")
    if (pathFrom is None) != (pathTo is None):
        raise HTTPException(422, "pathFrom 与 pathTo 必须同时提供。")
    if pathFrom and nodeId:
        raise HTTPException(422, "路径查询不能同时使用 nodeId 局部邻域。")
    if pathFrom and hops:
        raise HTTPException(422, "路径查询不能同时使用 hops 局部邻域。")
    w = store().read()
    # Keep the query scoped to the active laboratory as well as the optional
    # plan.  The browser already applies this scope to its local projection,
    # but direct graph consumers must not accidentally receive relationships
    # from another laboratory in the same workspace.
    plans = [
        p for p in w.get("plans", [])
        if (not planId or p.get("id") == planId)
        and (not labId or p.get("labId") == labId)
    ]
    if planId and not plans:
        if labId and any(p.get("id") == planId for p in w.get("plans", [])):
            raise HTTPException(404, "计划不属于当前实验室。")
        raise HTTPException(404, "计划不存在。")
    allowed_types = {item.strip() for item in (types or "").split(",") if item.strip()}
    allowed_relations = {item.strip() for item in (relations or "").split(",") if item.strip()}
    unknown_types = sorted(allowed_types - GRAPH_NODE_TYPES)
    if unknown_types:
        raise HTTPException(422, f"types 包含不支持的节点类型：{', '.join(unknown_types)}。")
    try:
        start_bound = dt(dateFrom) if dateFrom else None
        end_bound = dt(dateTo) if dateTo else None
    except ValueError as error:
        raise HTTPException(422, "dateFrom/dateTo 必须是带时区的 ISO 时间。") from error
    if start_bound and end_bound and start_bound >= end_bound:
        raise HTTPException(422, "dateFrom 必须早于 dateTo。")

    nodes: dict[str, dict] = {}
    edges: list[dict] = []
    documents = w.get("documents", [])
    documents_by_id = {str(item.get("id")): item for item in documents}

    def graph_plan_id(*ids: str) -> str | None:
        for value in ids:
            parts = str(value or "").split(":")
            if len(parts) >= 2 and parts[0] in {"plan", "step", "chemical", "control", "task", "check", "approval"}:
                return parts[1]
        return None

    def graph_source_span(*, plan_id: str | None = None, line: int | None = None, document_id: str | None = None) -> dict[str, Any]:
        if document_id and document_id in documents_by_id:
            return _document_source_span(documents_by_id[document_id])
        if plan_id:
            source_plan = next((item for item in plans if item.get("id") == plan_id), None)
            if source_plan is not None:
                return _plan_source_span(source_plan, line)
        return {}

    def add_node(node_id: str, node_type: str, label: str, **extra: Any) -> None:
        if node_id not in nodes:
            payload = {"id": node_id, "type": node_type, "label": label, **extra}
            if not payload.get("sourceSpan"):
                payload["sourceSpan"] = graph_source_span(
                    plan_id=payload.get("planId"), line=payload.get("sourceLine"), document_id=payload.get("documentId")
                )
            nodes[node_id] = payload
    def canonical_edge_state(value: Any) -> str:
        if value in {"pass", "unknown", "blocked"}:
            return value
        return {"verified": "pass", "rejected": "blocked", "done": "pass"}.get(value, "unknown")

    def add_edge(edge_id: str, source: str, target: str, relation: str, *, state: str | None = None, provenance: str | None = None, **extra: Any) -> None:
        # Consumers can render a single state vocabulary for every relation.
        # Legacy status values are retained for evidence edges, while the
        # canonical state makes unknown/rejected evidence explicit.
        status = extra.get("status")
        canonical_state = canonical_edge_state(state if state is not None else status)
        edge = {"id": edge_id, "from": source, "to": target, "relation": relation, "state": canonical_state, **extra}
        if not edge.get("sourceSpan"):
            edge["sourceSpan"] = graph_source_span(
                plan_id=edge.get("planId") or graph_plan_id(source, target),
                line=edge.get("sourceLine"),
                document_id=edge.get("documentId") or (target.split(":", 1)[1] if target.startswith("document:") else None),
            )
        if provenance:
            edge["provenance"] = provenance
        edges.append(edge)
    def in_window(item: dict) -> bool:
        if not start_bound and not end_bound:
            return True
        try:
            item_start, item_end = dt(item["start"]), dt(item["end"])
        except (KeyError, ValueError):
            return False
        return (not end_bound or item_start < end_bound) and (not start_bound or item_end > start_bound)

    def slug(value: str) -> str:
        # Keep graph IDs stable without leaking arbitrary source text into a
        # URL or DOM identifier.  The human label remains the original term.
        return hashlib.sha1(value.encode("utf-8")).hexdigest()[:12]

    def source_line_for(plan: dict, term: str) -> int | None:
        for line_no, line in enumerate(plan.get("source", "").splitlines(), 1):
            if term in line:
                return line_no
        return None

    equipment_by_id = {e["id"]: e for e in w.get("equipment", [])}
    tasks = w.get("tasks", [])
    all_checks = check_workspace(w)
    for plan in plans:
        pid = plan["id"]
        window_reservations = [r for r in plan.get("reservations", []) if in_window(r)]
        plan_checks = [check for check in all_checks if check.get("planId") == pid]
        current_status = plan_status(w, plan, plan_checks)
        add_node(
            f"plan:{pid}", "plan", plan.get("name", pid), planId=pid, labId=plan.get("labId"),
            revision=plan.get("revision"), state=current_status, status=current_status,
            source=plan.get("sourceName") or plan.get("source"),
            reservationCount=len(window_reservations),
            timeFiltered=bool(start_bound or end_bound),
        )
        used_equipment: set[str] = set()
        for index, step in enumerate(plan.get("steps", []), 1):
            sid = f"step:{plan['id']}:{step.get('line', index)}"
            add_node(sid, "step", step.get("text", f"步骤 {index}"), planId=pid, sourceLine=step.get("line"), state="confirmed" if step.get("confirmed") else "candidate")
            add_edge(f"{pid}:step:{step.get('line', index)}", f"plan:{pid}", sid, "contains", state="pass" if step.get("confirmed") else "unknown", sourceLine=step.get("line"), provenance="SOP 原文")
            for equipment_id in step.get("equipmentIds", []):
                if equipment_id not in equipment_by_id:
                    continue
                used_equipment.add(equipment_id)
                edge_id = f"{pid}:step:{step.get('line', index)}:equipment:{equipment_id}"
                add_edge(edge_id, sid, f"equipment:{equipment_id}", "requires", state="pass", sourceLine=step.get("line"), provenance="SOP 原文")

        # Chemical and control nodes are derived by the same deterministic
        # parser used by the review service.  They are source-anchored
        # assertions, not AI candidates or confirmed SDS/control facts, so the
        # graph keeps them amber/unknown until a reviewer records evidence.
        try:
            parsed_sop = parse_sop(plan.get("source", ""))
        except (TypeError, ValueError):
            parsed_sop = {"substances": [], "steps": []}
        chemical_nodes: dict[str, str] = {}
        for substance in parsed_sop.get("substances", []):
            # ``parse_sop`` returns Pydantic ``Substance`` values, while a
            # future parser adapter may return mappings. Keep this projection
            # JSON-safe for both forms and preserve enum values (``flammable``)
            # instead of Python's ``HazardClass.FLAMMABLE`` representation.
            if isinstance(substance, dict):
                name = substance.get("name")
                cas_number = substance.get("cas_number")
                hazard_classes = substance.get("hazard_classes") or []
            else:
                name = getattr(substance, "name", None)
                cas_number = getattr(substance, "cas_number", None)
                hazard_classes = getattr(substance, "hazard_classes", None) or []
            if not name:
                continue
            chemical_id = f"chemical:{pid}:{slug(name)}"
            source_line = source_line_for(plan, str(name))
            hazard_values = [getattr(item, "value", str(item)) for item in hazard_classes]
            add_node(chemical_id, "chemical", str(name), planId=pid, name=str(name), cas=cas_number, hazardClasses=hazard_values, state="unknown", sourceLine=source_line, source=plan.get("sourceName"), provenance="SOP 原文关键词实体（待人工确认）", detail="来自已保存计划原文的确定性关键词识别；尚未关联经核验 SDS 或相容性结论。")
            chemical_nodes[str(name)] = chemical_id
        control_nodes: dict[str, str] = {}
        for index, parsed_step in enumerate(parsed_sop.get("steps", []), 1):
            source_line = getattr(parsed_step, "source_line", None) or getattr(parsed_step, "order", index)
            sid = f"step:{pid}:{source_line}"
            controls = list(getattr(parsed_step, "controls", None) or [])
            ventilation = getattr(parsed_step, "ventilation", None)
            waste_stream = getattr(parsed_step, "waste_stream", None)
            controls += ([f"通风：{ventilation}"] if ventilation else [])
            controls += ([f"废液：{waste_stream}"] if waste_stream else [])
            for control in dict.fromkeys(str(item).strip() for item in controls if str(item).strip()):
                control_id = f"control:{pid}:{slug(control)}"
                add_node(control_id, "control", control, planId=pid, sourceLine=source_line, state="unknown", provenance="SOP 原文确定性解析（待人工确认）", detail="原文包含该控制线索；需要负责人确认控制措施、适用范围与证据。")
                control_nodes[control] = control_id
                if sid in nodes:
                    add_edge(f"{pid}:step:{source_line}:control:{slug(control)}", sid, control_id, "requires-control", state="unknown", sourceLine=source_line, provenance="SOP 原文")
            instruction = str(getattr(parsed_step, "instruction", ""))
            for name, chemical_id in chemical_nodes.items():
                if name in instruction and sid in nodes:
                    add_edge(f"{pid}:step:{source_line}:chemical:{slug(name)}", sid, chemical_id, "mentions", state="unknown", sourceLine=source_line, provenance="SOP 原文")
        for reservation in plan.get("reservations", []):
            if not in_window(reservation):
                continue
            equipment_id = reservation.get("equipmentId")
            equipment = equipment_by_id.get(equipment_id)
            if not equipment:
                continue
            used_equipment.add(equipment_id)
            add_node(f"equipment:{equipment_id}", "equipment", equipment.get("name", equipment_id), equipmentId=equipment_id, labId=equipment.get("labId"), kind=equipment.get("kind"), state=equipment.get("status"), start=reservation.get("start"), end=reservation.get("end"))
            add_edge(f"{pid}:reservation:{reservation.get('id', equipment_id)}", f"plan:{pid}", f"equipment:{equipment_id}", "reserves", state="pass", start=reservation.get("start"), end=reservation.get("end"), provenance="计划预约")
        for equipment_id in used_equipment:
            equipment = equipment_by_id.get(equipment_id)
            if not equipment:
                continue
            add_node(f"equipment:{equipment_id}", "equipment", equipment.get("name", equipment_id), equipmentId=equipment_id, labId=equipment.get("labId"), kind=equipment.get("kind"), state=equipment.get("status"))
            for document in documents:
                if document.get("equipmentId") != equipment_id:
                    continue
                did = f"document:{document['id']}"
                add_node(did, "document", document.get("name", document["id"]), documentId=document["id"], equipmentId=equipment_id, state=document.get("status"), status=document.get("status"), validFrom=document.get("validFrom"), validUntil=document.get("validUntil"), source=document.get("source"), sourceType=document.get("sourceType"), sourceLocator=document.get("sourceLocator"), capturedAt=document.get("capturedAt"), sourceFileName=document.get("sourceFileName"), sourceFileSize=document.get("sourceFileSize"), sourceLocations=document.get("sourceLocations", []))
                add_edge(f"{equipment_id}:document:{document['id']}", f"equipment:{equipment_id}", did, "evidence", status=document.get("status"), specVersion=document.get("specVersion"), provenance=document.get("source") or document.get("sourceType"))
        for task in tasks:
            if task.get("planId") != pid:
                continue
            tid = f"task:{task['id']}"
            add_node(tid, "task", task.get("title", task["id"]), taskId=task["id"], planId=pid, equipmentId=task.get("equipmentId"), state=task.get("status"), dueAt=task.get("dueAt"))
            add_edge(f"{pid}:task:{task['id']}", f"plan:{pid}", tid, "action", state=task.get("status"), provenance="工作区任务")
            if task.get("equipmentId"):
                add_edge(f"{task['id']}:equipment", tid, f"equipment:{task['equipmentId']}", "targets", state="pass", provenance="任务关联字段")
        for check in all_checks:
            if check.get("planId") != pid:
                continue
            cid = f"check:{check['id']}"
            add_node(cid, "check", check.get("title", check["id"]), checkId=check["id"], planId=pid, equipmentId=check.get("equipmentId"), state=check.get("state"), code=check.get("code"), sourceLine=check.get("line"))
            add_edge(f"{pid}:check:{check['id']}", f"plan:{pid}", cid, "evaluates", state=check.get("state"), code=check.get("code"), sourceLine=check.get("line"), provenance="规则核验快照")
            if check.get("equipmentId"):
                add_edge(f"{check['id']}:equipment", cid, f"equipment:{check['equipmentId']}", "evaluates", state=check.get("state"), provenance="规则核验对象")
            # Evidence applicability is a first-class relationship.  A
            # verified document can still be expired, wrong-version, pending
            # or in conflict for this plan's exact reservation window; the
            # graph must expose that scoped result instead of only showing a
            # green document node from its upload-level status.
            for applicability in check.get("evidenceApplicability", []):
                document_node = f"document:{applicability.get('documentId')}"
                if document_node not in nodes:
                    continue
                applicability_state = applicability.get("state", "candidate")
                relation_state = "pass" if applicability_state == "supports" and applicability.get("effectiveForAllReservations") else "unknown"
                add_edge(
                    f"{check['id']}:evidence:{applicability.get('documentId')}", cid, document_node,
                    "supports-evidence", state=relation_state,
                    applicabilityState=applicability_state,
                    reason=applicability.get("reason"),
                    coveredReservationIds=applicability.get("coveredReservationIds", []),
                    missingReservationIds=applicability.get("missingReservationIds", []),
                    conflictingDocumentIds=applicability.get("conflictingDocumentIds", []),
                    provenance="证据适用性计算",
                )

        # Approval records are immutable human decisions.  A pending node is
        # explicit when no decision exists; neither node type implies an
        # approval result on its own.
        plan_decisions = [decision for decision in w.get("decisions", []) if decision.get("planId") == pid]
        if plan_decisions:
            for decision in plan_decisions:
                decision_id = f"approval:{decision.get('id')}"
                disposition = decision.get("disposition", "approved")
                same_basis = decision.get("basis") == basis(w, plan)
                state = "pass" if disposition == "approved" and same_basis and current_status == "approved" else "unknown"
                add_node(decision_id, "approval", "人工放行" if disposition == "approved" else "例外授权" if disposition == "exception" else "审批退回", approvalId=decision.get("id"), planId=pid, actor=decision.get("actor"), stage=decision.get("stage"), disposition=disposition, state=state, provenance="人工审批记录", detail=f"{decision.get('actor', '未记录')} · {decision.get('at', '')}{'' if same_basis else ' · 计划条件已变化，历史记录需重新核对'}")
                add_edge(f"{pid}:approval:{decision.get('id')}", f"plan:{pid}", decision_id, "approval", state=state, provenance="人工审批记录")
        else:
            pending_id = f"approval:{pid}:pending"
            add_node(pending_id, "approval", "等待人工放行", approvalId=None, planId=pid, state="unknown", provenance="审批链", detail="当前没有人工审批记录；AI 与规则结果不能替代负责人决策。")
            add_edge(f"{pid}:approval:pending", f"plan:{pid}", pending_id, "approval", state="unknown", provenance="审批链")

    # A path query is an explicit, stable view for review hand-off: the
    # server chooses the shortest undirected path over the same source-derived
    # edges that the graph already exposes.  Sorting neighbours by edge/id
    # makes equal-length paths deterministic across SQLite/Python runs.  The
    # result is a projection only; it never mutates the workspace or implies
    # that an unknown edge is safe.
    path_result: dict[str, Any] = {
        "requested": bool(pathFrom or pathTo),
        "from": pathFrom,
        "to": pathTo,
        "found": False,
        "nodeIds": [],
        "edgeIds": [],
        "hops": None,
    }
    if pathFrom and pathTo:
        if pathFrom not in nodes or pathTo not in nodes:
            missing = pathFrom if pathFrom not in nodes else pathTo
            raise HTTPException(404, f"关系图路径端点不存在：{missing}。")
        if allowed_types:
            invalid_types = [node_id for node_id in (pathFrom, pathTo) if nodes[node_id]["type"] not in allowed_types]
            if invalid_types:
                raise HTTPException(422, "路径端点类型被 types 过滤器排除。")
        adjacency: dict[str, list[tuple[str, dict[str, Any]]]] = {}
        for edge in edges:
            if allowed_relations and edge["relation"] not in allowed_relations:
                continue
            if allowed_types and (nodes[edge["from"]]["type"] not in allowed_types or nodes[edge["to"]]["type"] not in allowed_types):
                continue
            adjacency.setdefault(edge["from"], []).append((edge["to"], edge))
            adjacency.setdefault(edge["to"], []).append((edge["from"], edge))
        for links in adjacency.values():
            links.sort(key=lambda item: (item[1]["id"], item[0]))
        queue = [pathFrom]
        previous: dict[str, tuple[str, dict[str, Any]] | None] = {pathFrom: None}
        cursor = 0
        while cursor < len(queue) and pathTo not in previous:
            current = queue[cursor]
            cursor += 1
            for target, edge in adjacency.get(current, []):
                if target in previous:
                    continue
                previous[target] = (current, edge)
                queue.append(target)
                if target == pathTo:
                    break
        if pathTo in previous:
            path_nodes = [pathTo]
            path_edges: list[dict[str, Any]] = []
            current = pathTo
            while current != pathFrom:
                parent = previous[current]
                assert parent is not None
                parent_id, edge = parent
                path_edges.append(edge)
                path_nodes.append(parent_id)
                current = parent_id
            path_nodes.reverse()
            path_edges.reverse()
            path_result.update({
                "found": True,
                "nodeIds": path_nodes,
                "edgeIds": [edge["id"] for edge in path_edges],
                "hops": len(path_edges),
            })
            visible = set(path_nodes)
            edges = [edge for edge in path_edges if edge["from"] in visible and edge["to"] in visible]
            nodes = {key: value for key, value in nodes.items() if key in visible}
        else:
            # Keep both endpoints visible when no relation path exists. This
            # lets a reviewer see that the requested objects are disconnected
            # under the active date/type/relation scope rather than receiving
            # an ambiguous empty graph.
            nodes = {key: value for key, value in nodes.items() if key in {pathFrom, pathTo}}
            edges = []
    elif nodeId:
        adjacency: dict[str, set[str]] = {}
        for edge in edges:
            adjacency.setdefault(edge["from"], set()).add(edge["to"])
            adjacency.setdefault(edge["to"], set()).add(edge["from"])
        if nodeId not in nodes:
            raise HTTPException(404, "关系图对象不存在。")
        visible = {nodeId}
        frontier = {nodeId}
        for _ in range(hops):
            frontier = {linked for item in frontier for linked in adjacency.get(item, set())} - visible
            visible.update(frontier)
        nodes = {key: value for key, value in nodes.items() if key in visible}
    if allowed_types:
        nodes = {key: value for key, value in nodes.items() if value["type"] in allowed_types}
    visible_ids = set(nodes)
    edges = [edge for edge in edges if edge["from"] in visible_ids and edge["to"] in visible_ids]
    if allowed_relations:
        edges = [edge for edge in edges if edge["relation"] in allowed_relations]
    return {"revision": w.get("revision", 0), "planId": planId, "labId": labId, "nodes": list(nodes.values()), "edges": edges, "path": path_result, "filters": {"types": sorted(allowed_types), "relations": sorted(allowed_relations), "nodeId": nodeId, "pathFrom": pathFrom, "pathTo": pathTo, "dateFrom": dateFrom, "dateTo": dateTo, "hops": hops, "labId": labId}}

@router.post("/actions")
def mutate_workspace(body: Mutation, request: Request, actor: str = Depends(local_owner)):
    role = local_role(request)
    if body.type in {"approve", "approval_decision"} and role not in {"safety_owner", "ehs_manager", "administrator"}:
        raise HTTPException(403, "当前工作区身份没有记录正式复核的权限。")
    if body.type == "verify_document" and role not in {"safety_owner", "ehs_manager", "administrator"}:
        raise HTTPException(403, "当前工作区身份没有核验资料的权限。")
    if body.type == "transition_plan" and role not in {"instructor", "safety_owner", "ehs_manager", "administrator"}:
        raise HTTPException(403, "当前工作区身份没有变更计划生命周期的权限。")
    handoff_update = body.type == "update_task" and any(key in body.payload for key in ("newOwner", "expectedOwner", "handoffNote"))
    if (body.type == "assign_task" or handoff_update) and local_role(request) not in {"instructor", "safety_owner", "ehs_manager", "administrator"}:
        raise HTTPException(403, "当前工作区身份没有转交协作任务的权限。")
    if body.type == "review_branch_patch" and role not in {"instructor", "safety_owner", "ehs_manager", "administrator"}:
        raise HTTPException(403, "当前工作区身份没有复核候选字段的权限。")
    try:
        payload = dict(body.payload)
        if body.type in {"approve", "approval_decision"}:
            payload["_serverRole"] = role
        return envelope(store().mutate(body.expectedRevision, body.requestId, body.type, payload, actor))
    except (ValidationError, ValueError, KeyError, TypeError) as error:
        message = str(error) if not isinstance(error, ValidationError) else "输入结构不符合工作区契约。"
        raise HTTPException(422, message) from error

@router.get("/export")
def export_workspace(actor: str = Depends(local_owner)):
    w = store().read()
    return {"exportedAt": now(), "scope": "资源与证据审查记录，不是安全证书", "sourceRecord": "workspace snapshot", **envelope(w)}


def _workspace_artifact_export(store_instance: WorkspaceStore, workspace: dict) -> list[dict]:
    """Build a validated, public-safe artifact manifest for a workspace."""
    documents = {str(item["id"]): item for item in workspace.get("documents", [])}
    referenced = {str(item.get("sourceArtifactId")) for item in documents.values() if item.get("sourceArtifactId")}
    rows = store_instance.source_artifacts(document_ids=set(documents))
    by_id = {str(item.get("artifactId")): item for item in rows}
    if referenced != set(by_id).intersection(referenced):
        missing = sorted(referenced - set(by_id))
        if missing:
            raise HTTPException(409, f"工作区引用的原始来源文件不可导出：{', '.join(missing)}。")
    result: list[dict] = []
    for artifact_id in sorted(referenced):
        metadata = by_id[artifact_id]
        resolved = store_instance._source_artifact_storage(artifact_id, document_id=str(metadata["documentId"]))
        if not resolved:
            raise HTTPException(409, "来源文件记录与工作区资料不一致，无法导出。")
        public, path = resolved
        if public["documentId"] not in documents or documents[public["documentId"]].get("sourceArtifactId") != artifact_id:
            raise HTTPException(409, "来源文件记录未关联当前资料，无法导出。")
        filename = _artifact_filename(public["filename"])
        if filename != public["filename"] or Path(filename).suffix.lower() not in ARTIFACT_SUFFIXES:
            raise HTTPException(409, "来源文件名不符合存档安全要求，无法导出。")
        result.append({"metadata": public, "path": path, "member": f"artifacts/{artifact_id}/{filename}"})
    return result


@router.get("/export.zip")
def export_workspace_zip(actor: str = Depends(local_owner)):
    """Export workspace JSON plus retained source bytes as an authenticated ZIP."""
    from io import BytesIO

    workspace = store().read()
    artifacts = _workspace_artifact_export(store(), workspace)
    if len(artifacts) + 2 > ARCHIVE_MAX_FILES:
        raise HTTPException(413, "工作区 ZIP 存档条目过多。")
    manifest = {
        "format": ARCHIVE_FORMAT,
        "exportedAt": now(),
        "workspaceRevision": workspace.get("revision", 0),
        "scope": "资源与证据审查记录，不是安全证书",
        "included": ["workspace.json", "source-artifacts"],
        "excluded": ["analysis-runs", "evaluation-runs", "jobs", "outbox", "membership-ledger"],
        "artifacts": [{**item["metadata"], "member": item["member"]} for item in artifacts],
    }
    workspace_bytes = json.dumps(workspace, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    manifest_bytes = json.dumps(manifest, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    uncompressed_size = len(workspace_bytes) + len(manifest_bytes) + sum(int(item["metadata"]["size"]) for item in artifacts)
    if uncompressed_size > ARCHIVE_MAX_BYTES:
        raise HTTPException(413, "工作区 ZIP 展开后不能超过 32 MB。")
    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("workspace.json", workspace_bytes)
        archive.writestr("manifest.json", manifest_bytes)
        for item in artifacts:
            archive.write(item["path"], arcname=item["member"])
    payload = output.getvalue()
    if len(payload) > ARCHIVE_MAX_BYTES:
        raise HTTPException(413, "工作区 ZIP 存档不能超过 32 MB。")
    return Response(
        content=payload,
        media_type="application/zip",
        headers={"Content-Disposition": 'attachment; filename="aegislab-workspace.zip"', "X-Workspace-Revision": str(workspace.get("revision", 0))},
    )


async def _bounded_archive_body(request: Request) -> bytes:
    """Read an application/zip body without trusting Content-Length."""
    content_type = request.headers.get("content-type", "").split(";", 1)[0].strip().lower()
    if content_type and content_type not in {"application/zip", "application/x-zip-compressed"}:
        raise HTTPException(415, "ZIP 恢复接口只接受 application/zip 原始请求体。")
    declared = request.headers.get("content-length")
    try:
        if declared is not None and int(declared) > ARCHIVE_MAX_BYTES:
            raise HTTPException(413, "ZIP 存档不能超过 32 MB。")
    except ValueError:
        raise HTTPException(400, "Content-Length 无效。")
    body = bytearray()
    async for chunk in request.stream():
        body.extend(chunk)
        if len(body) > ARCHIVE_MAX_BYTES:
            raise HTTPException(413, "ZIP 存档不能超过 32 MB。")
    return bytes(body)


def _read_workspace_zip(raw: bytes) -> tuple[dict, list[dict]]:
    """Parse and validate a ZIP archive without extracting untrusted paths."""
    from io import BytesIO

    if len(raw) > ARCHIVE_MAX_BYTES:
        raise HTTPException(413, "ZIP 存档不能超过 32 MB。")
    try:
        archive = zipfile.ZipFile(BytesIO(raw), "r")
    except (zipfile.BadZipFile, OSError) as error:
        raise HTTPException(422, "ZIP 存档无法读取。") from error
    with archive:
        infos = archive.infolist()
        if len(infos) > ARCHIVE_MAX_FILES:
            raise HTTPException(413, "ZIP 存档条目过多。")
        names: list[str] = []
        total_uncompressed = 0
        for info in infos:
            try:
                name = _safe_archive_member(info.filename)
            except ValueError as error:
                raise HTTPException(422, str(error)) from error
            if info.is_dir() or name.endswith("/"):
                raise HTTPException(422, "ZIP 存档不能包含目录条目。")
            if info.flag_bits & 0x1:
                raise HTTPException(422, "不接受加密 ZIP 存档。")
            if name in names:
                raise HTTPException(422, "ZIP 存档包含重复条目。")
            names.append(name)
            total_uncompressed += int(info.file_size)
            if total_uncompressed > ARCHIVE_MAX_BYTES:
                raise HTTPException(413, "ZIP 展开后不能超过 32 MB。")
        if set(names) - {"workspace.json", "manifest.json"} and not any(name.startswith("artifacts/") for name in names):
            raise HTTPException(422, "ZIP 存档包含未知文件。")
        if "workspace.json" not in names or "manifest.json" not in names:
            raise HTTPException(422, "ZIP 存档必须包含 workspace.json 和 manifest.json。")
        try:
            workspace = json.loads(archive.read("workspace.json").decode("utf-8"))
            manifest = json.loads(archive.read("manifest.json").decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError, KeyError) as error:
            raise HTTPException(422, "ZIP 存档中的 JSON 清单无效。") from error
        if not isinstance(manifest, dict) or manifest.get("format") != ARCHIVE_FORMAT or not isinstance(manifest.get("artifacts"), list):
            raise HTTPException(422, "ZIP 存档格式或 artifact 清单无效。")
        if len(manifest["artifacts"]) > ARCHIVE_MAX_FILES:
            raise HTTPException(413, "ZIP 存档来源文件过多。")

        try:
            workspace, _ = migrate_workspace_import(workspace)
            validate_workspace(workspace)
        except (ValidationError, ValueError, KeyError, TypeError) as error:
            raise HTTPException(422, f"ZIP 工作区结构无效：{str(error)[:180]}") from error
        documents = {str(item["id"]): item for item in workspace.get("documents", [])}
        artifacts: list[dict] = []
        expected_names = {"workspace.json", "manifest.json"}
        for item in manifest["artifacts"]:
            if not isinstance(item, dict):
                raise HTTPException(422, "ZIP artifact 清单条目无效。")
            artifact_id = str(item.get("artifactId") or "")
            document_id = str(item.get("documentId") or "")
            filename = str(item.get("filename") or "")
            if not _ARTIFACT_ID_RE.fullmatch(artifact_id) or document_id not in documents:
                raise HTTPException(422, "ZIP artifact 标识或资料关联无效。")
            safe_name = _artifact_filename(filename)
            if safe_name != filename or Path(filename).suffix.lower() not in ARTIFACT_SUFFIXES:
                raise HTTPException(422, "ZIP artifact 文件名不安全或类型不受支持。")
            raw_size = item.get("size")
            if isinstance(raw_size, bool) or not isinstance(raw_size, int) or raw_size < 0 or raw_size > ARTIFACT_MAX_BYTES:
                raise HTTPException(422, "ZIP artifact size 必须是 0 到 8 MB 的整数。")
            raw_version = item.get("version", 1)
            if isinstance(raw_version, bool) or not isinstance(raw_version, int) or raw_version < 1:
                raise HTTPException(422, "ZIP artifact version 无效。")
            media_type = item.get("mediaType", "application/octet-stream")
            if not isinstance(media_type, str) or not media_type.strip() or len(media_type) > 180:
                raise HTTPException(422, "ZIP artifact mediaType 无效。")
            status = item.get("status", "stored")
            if not isinstance(status, str) or status not in {"stored"}:
                raise HTTPException(422, "ZIP artifact status 无效。")
            if documents[document_id].get("sourceArtifactId") != artifact_id:
                raise HTTPException(422, "ZIP artifact 与资料 sourceArtifactId 不一致。")
            source_version = documents[document_id].get("sourceArtifactVersion")
            if source_version is not None and source_version != raw_version:
                raise HTTPException(422, "资料 sourceArtifactVersion 与 ZIP artifact 不一致。")
            member = str(item.get("member") or f"artifacts/{artifact_id}/{filename}")
            try:
                _safe_archive_member(member)
            except ValueError as error:
                raise HTTPException(422, str(error)) from error
            if member != f"artifacts/{artifact_id}/{filename}" or member in expected_names or member not in names:
                raise HTTPException(422, "ZIP artifact 成员路径不符合清单。")
            if member in expected_names:
                raise HTTPException(422, "ZIP artifact 成员重复。")
            expected_names.add(member)
            info = archive.getinfo(member)
            if info.file_size > ARTIFACT_MAX_BYTES:
                raise HTTPException(413, "单份来源文件不能超过 8 MB。")
            data = archive.read(member)
            if len(data) != raw_size:
                raise HTTPException(422, "ZIP artifact 文件大小与清单不一致。")
            artifacts.append({"metadata": item, "data": data})
        if set(names) != expected_names:
            raise HTTPException(422, "ZIP 存档包含未列入 manifest 的文件。")
        referenced = {str(item.get("sourceArtifactId")) for item in documents.values() if item.get("sourceArtifactId")}
        if referenced != {str(item["metadata"]["artifactId"]) for item in artifacts}:
            raise HTTPException(422, "ZIP 工作区引用与 artifact 清单不一致。")
        return workspace, artifacts


@router.post("/restore.zip")
async def restore_workspace_zip(request: Request, actor: str = Depends(local_owner)):
    """Restore workspace and source bytes from an authenticated raw ZIP body."""
    if local_role(request) not in {"safety_owner", "ehs_manager", "administrator"}:
        raise HTTPException(403, "当前工作区身份没有恢复服务端存档的权限。")
    expected_raw = request.query_params.get("expectedRevision")
    expected_revision: int | None = None
    if expected_raw is not None:
        try:
            expected_revision = int(expected_raw)
        except (TypeError, ValueError) as error:
            raise HTTPException(422, "expectedRevision 必须是非负整数。") from error
        if expected_revision < 0:
            raise HTTPException(422, "expectedRevision 必须是非负整数。")
    request_id = request.query_params.get("requestId") or ""
    if len(request_id) > 180:
        raise HTTPException(422, "requestId 过长。")
    raw = await _bounded_archive_body(request)
    workspace, artifacts = _read_workspace_zip(raw)
    current = store().read()
    restored = store().restore_with_artifacts(
        expected_revision=int(current.get("revision", 0)) if expected_revision is None else expected_revision,
        request_id=request_id or f"zip-{hashlib.sha256(raw).hexdigest()[:32]}",
        candidate=workspace,
        artifacts=artifacts,
        actor=actor,
    )
    return {**envelope(restored), "archive": {"format": ARCHIVE_FORMAT, "artifacts": len(artifacts)}}

@router.post("/restore")
def restore_workspace(body: RestoreRequest, request: Request, actor: str = Depends(local_owner)):
    """Restore a validated browser export with a durable rollback point.

    This endpoint is intentionally role-gated and never trusts the imported
    revision.  The current server payload is snapshotted before replacement;
    a stale browser receives 409 instead of overwriting newer work.
    """
    if local_role(request) not in {"safety_owner", "ehs_manager", "administrator"}:
        raise HTTPException(403, "当前工作区身份没有恢复服务端存档的权限。")
    try:
        candidate, migration = migrate_workspace_import(body.workspace)
        restored = store().restore(expected_revision=body.expectedRevision, request_id=body.requestId, candidate=candidate, actor=actor, migration_report=migration)
        return {**envelope(restored), "migration": migration}
    except (ValidationError, ValueError, KeyError, TypeError) as error:
        message = str(error) if not isinstance(error, ValidationError) else "恢复文件结构不符合工作区契约。"
        raise HTTPException(422, message) from error

@router.get("/snapshots")
def list_workspace_snapshots(limit: int = 20, actor: str = Depends(local_owner)):
    return {"snapshots": store().snapshots(limit)}

@router.get("/snapshots/{snapshot_id}")
def get_workspace_snapshot(snapshot_id: str, actor: str = Depends(local_owner)):
    snapshot = store().snapshot(snapshot_id)
    if not snapshot:
        raise HTTPException(404, "恢复前快照不存在。")
    return snapshot


@router.get("/snapshots/{snapshot_id}/artifacts/{artifact_id}/content")
def get_snapshot_artifact_content(snapshot_id: str, artifact_id: str, actor: str = Depends(local_owner)):
    resolved = store().snapshot_artifact_path(snapshot_id, artifact_id)
    if not resolved:
        raise HTTPException(404, "恢复前快照中的来源文件不存在。")
    metadata, path = resolved
    return FileResponse(path, media_type=metadata.get("mediaType") or "application/octet-stream", filename=metadata.get("filename") or "source")


def _notification_audience(request: Request, requested: str | None) -> str | None:
    """Constrain inbox reads to the authenticated member's responsibility."""
    role = local_role(request)
    if role in {"safety_owner", "ehs_manager", "administrator"}:
        return requested
    principal = local_principal(request)
    own_audience = f"任务负责人:{principal}"
    if requested and requested not in {principal, own_audience}:
        raise HTTPException(403, "当前成员不能读取其他责任人的通知。")
    return own_audience


@router.get("/notifications")
def list_notifications(
    request: Request,
    status: Literal["pending", "acknowledged", "all"] = "pending",
    audience: str | None = None,
    planId: str | None = None,
    limit: int = 50,
    actor: str = Depends(local_owner),
):
    """Read the durable local inbox generated by workspace mutations.

    This is deliberately scoped to the browser workspace.  It exposes pending
    hand-off records for the UI and does not claim email, chat, or push
    delivery.  Acknowledgement is an auditable local state transition.
    """
    effective_audience = _notification_audience(request, audience)
    try:
        notifications = store().notifications(status=status, audience=effective_audience, plan_id=planId, limit=limit)
    except ValueError as error:
        raise HTTPException(422, str(error)) from error
    return {
        "status": status,
        "audience": effective_audience,
        "planId": planId,
        "notifications": notifications,
        "scope": "本地 outbox 收件箱；未发送外部通知，需要负责人打开对象后复核。",
    }


@router.post("/notifications/{event_id}/ack")
def acknowledge_notification(event_id: str, request: Request, actor: str = Depends(local_owner)):
    """Acknowledge a notification without changing the underlying business object."""
    allowed_audience = _notification_audience(request, None)
    return {
        "notification": store().acknowledge_notification(event_id, actor, allowed_audience=allowed_audience),
        "scope": "仅确认本机收件箱已读，不代表审批、证据核验或实验放行。",
    }


def _line_locations(text: str) -> list[dict[str, Any]]:
    """Return bounded line ranges for a plain-text source.

    A location is a navigation contract, not a claim about PDF page layout.
    Long files are grouped into deterministic ranges so a large upload cannot
    turn provenance metadata into an unbounded workspace payload.
    """
    lines = text.splitlines() or [""]
    if len(lines) <= 2000:
        return [{"kind": "line", "lineStart": index, "lineEnd": index, "label": f"L{index}"} for index in range(1, len(lines) + 1)]
    chunk_size = 100
    return [
        {"kind": "range", "lineStart": start, "lineEnd": min(start + chunk_size - 1, len(lines)), "label": f"L{start}–L{min(start + chunk_size - 1, len(lines))}"}
        for start in range(1, len(lines) + 1, chunk_size)
    ]


@router.post('/import-text')
async def import_text(request: Request, filename: str, documentId: str | None = None, actor: str = Depends(local_owner)):
    """Bounded document extraction with explicit source coordinates.

    Unless ``documentId`` is supplied, the service keeps only extracted text
    and provenance metadata.  A document upload may explicitly retain the
    original bytes in the controlled artifact volume; this never executes
    macros or claims that extraction itself verifies source applicability.
    """
    from io import BytesIO
    import zipfile
    raw = bytearray()
    async for chunk in request.stream():
        raw.extend(chunk)
        if len(raw) > 8 * 1024 * 1024:
            raise HTTPException(413, '文件不能超过 8 MB。')
    suffix = Path(filename).suffix.lower()
    source_file_name = Path(filename).name
    source_media_type = request.headers.get("content-type") or "application/octet-stream"
    if source_media_type == "application/octet-stream":
        source_media_type = {
            ".txt": "text/plain", ".md": "text/markdown", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".pdf": "application/pdf",
        }.get(Path(filename).suffix.lower(), source_media_type)
    source_locations: list[dict[str, Any]] = []
    locator = "extracted-text-lines"
    try:
        if suffix in {'.txt', '.md'}:
            text = raw.decode('utf-8-sig')
            source_locations = _line_locations(text)
        elif suffix == '.docx':
            from docx import Document as WordDocument
            with zipfile.ZipFile(BytesIO(raw)) as archive:
                if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValueError('文档展开后过大。')
            doc = WordDocument(BytesIO(raw))
            # Preserve paragraph/table order. No unsupported page/line attribution.
            from docx.oxml.ns import qn
            from docx.text.paragraph import Paragraph
            from docx.table import Table
            parts = []
            for element in doc.element.body:
                if element.tag == qn('w:p'):
                    parts.append(Paragraph(element, doc).text)
                elif element.tag == qn('w:tbl'):
                    parts.extend(' | '.join(c.text for c in row.cells) for row in Table(element, doc).rows)
            text = '\n'.join(parts)
            source_locations = [{"kind": "paragraph", "paragraph": index, "label": f"段落 {index}"} for index in range(1, len(parts) + 1)]
            locator = "docx-paragraphs"
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(raw), strict=False)
            if reader.is_encrypted:
                raise ValueError('请先移除文档密码。')
            if len(reader.pages) > 100:
                raise ValueError('本次导入最多支持 100 页。')
            page_texts = [page.extract_text() or '' for page in reader.pages]
            text = '\n\n'.join(page_texts)
            source_locations = [{"kind": "page", "page": index, "label": f"PDF 第 {index} 页"} for index in range(1, len(page_texts) + 1)]
            locator = "pdf-pages-text-layer"
        else:
            raise ValueError('仅支持 TXT、Markdown、DOCX 和带文本层的 PDF。')
        if not text.strip():
            raise ValueError('没有提取到可用文本；扫描件请先转成文本后导入，不自动编造内容。')
        if len(text) > 100000:
            raise ValueError('提取文本超过 100,000 字符，请按实验计划拆分。')
        artifact = None
        if documentId:
            artifact = store().save_source_artifact(
                document_id=documentId,
                filename=source_file_name,
                media_type=source_media_type,
                data=bytes(raw),
                actor=actor,
            )
        return {
            'filename': source_file_name,
            'text': text,
            'sourceType': 'upload',
            'sourceLocator': locator,
            'locator': locator,
            'capturedAt': now(),
            'contentLength': len(text),
            'sourceFileName': source_file_name,
            'sourceFileSize': len(raw),
            'sourceMediaType': source_media_type,
            'sourceLocations': source_locations,
            'sourceArtifactId': artifact.get('artifactId') if artifact else None,
            'sourceArtifactVersion': artifact.get('version') if artifact else None,
            'note': '原文定位对应提取后的行、段落或 PDF 文本层；不是未经 OCR 验证的版面坐标。上传成功不会自动核验或证明适用性。',
        }
    except HTTPException:
        raise
    except ImportError as error:
        raise HTTPException(503, '服务端缺少文档解析依赖，请安装 backend/requirements.txt。') from error
    except Exception as error:
        raise HTTPException(422, f'文件无法提取：{str(error)[:180]}') from error


@router.get('/documents/{document_id}/source')
def source_artifact_metadata(document_id: str, actor: str = Depends(local_owner)):
    """Return metadata for the original file associated with one document."""
    document = next((item for item in store().read().get('documents', []) if item.get('id') == document_id), None)
    if not document:
        raise HTTPException(404, '资料不存在。')
    artifact_id = document.get('sourceArtifactId')
    if not artifact_id:
        raise HTTPException(404, '该资料没有保留原始文件。')
    metadata = store().source_artifact(str(artifact_id), document_id=document_id)
    if not metadata:
        raise HTTPException(404, '该资料的原始文件记录不存在。')
    return metadata


@router.get('/documents/{document_id}/source/span')
def source_span_metadata(
    document_id: str,
    line: int | None = None,
    page: int | None = None,
    sourceArtifactVersion: int | None = None,
    actor: str = Depends(local_owner),
):
    """Resolve one recorded document coordinate with version protection.

    A caller may ask for a line or PDF page, but the coordinate must already
    exist in ``sourceLocations``.  Supplying a stale artifact version is
    rejected instead of returning a span into another uploaded file.
    """
    if line is not None and line < 1:
        raise HTTPException(422, "line 必须为正整数。")
    if page is not None and page < 1:
        raise HTTPException(422, "page 必须为正整数。")
    if line is not None and page is not None:
        raise HTTPException(422, "line 与 page 不能同时提供。")
    document = next((item for item in store().read().get("documents", []) if item.get("id") == document_id), None)
    if not document:
        raise HTTPException(404, "资料不存在。")
    current_version = document.get("sourceArtifactVersion")
    if sourceArtifactVersion is not None and sourceArtifactVersion != current_version:
        raise HTTPException(409, "资料来源版本已变化，拒绝跨版本回链。")
    span = _document_source_span(document, line=line, page=page)
    if not span:
        raise HTTPException(422, "资料没有与请求坐标匹配的来源定位。")
    return {"documentId": document_id, "sourceSpan": span}


@router.get('/documents/{document_id}/source/content')
def source_artifact_content(document_id: str, actor: str = Depends(local_owner)):
    """Stream the retained original only through the authenticated workspace."""
    document = next((item for item in store().read().get('documents', []) if item.get('id') == document_id), None)
    if not document:
        raise HTTPException(404, '资料不存在。')
    artifact_id = document.get('sourceArtifactId')
    if not artifact_id:
        raise HTTPException(404, '该资料没有保留原始文件。')
    resolved = store().source_artifact_path(str(artifact_id), document_id=document_id)
    if not resolved:
        raise HTTPException(404, '该资料的原始文件记录不存在。')
    metadata, path = resolved
    return FileResponse(path, media_type=metadata['mediaType'], filename=metadata['filename'], headers={'X-Source-Artifact-Version': str(metadata['version'])})


@router.get('/plans/{plan_id}/analysis/runs')
def list_analysis_runs(plan_id: str, limit: int = 20, actor: str = Depends(local_owner)):
    w = store().read()
    if not any(plan["id"] == plan_id for plan in w["plans"]):
        raise HTTPException(404, '计划不存在。')
    return {"planId": plan_id, "runs": store().analysis_runs(plan_id, limit)}

@router.get('/plans/{plan_id}/evaluations')
def list_evaluation_runs(plan_id: str, limit: int = 20, actor: str = Depends(local_owner)):
    w = store().read()
    if not any(plan["id"] == plan_id for plan in w["plans"]):
        raise HTTPException(404, '计划不存在。')
    return {"planId": plan_id, "evaluations": store().evaluation_runs(plan_id, limit)}

@router.get('/plans/{plan_id}/jobs')
def list_plan_jobs(plan_id: str, limit: int = 20, actor: str = Depends(local_owner)):
    if not any(plan["id"] == plan_id for plan in store().read()["plans"]):
        raise HTTPException(404, '计划不存在。')
    return {"planId": plan_id, "jobs": store().jobs(plan_id, limit)}

@router.get('/jobs/{job_id}')
def get_job(job_id: str, actor: str = Depends(local_owner)):
    job = store().job(job_id)
    if not job:
        raise HTTPException(404, '后台任务不存在。')
    return job


def _run_expiry_background(request_id: str, as_of: str, actor: str) -> None:
    """Execute one queued expiry job in the process-local FastAPI pool."""
    try:
        store().expiry_recheck(request_id=request_id, as_of=as_of, actor=actor)
    except Exception:
        # ``expiry_recheck`` persists failures; background callbacks cannot
        # propagate an exception after the HTTP response has been sent.
        return


@router.post('/expiry/recheck')
def recheck_expiry(body: ExpiryRecheckRequest, background_tasks: BackgroundTasks = None, background: bool = False, actor: str = Depends(local_owner)):
    """Persist a manual expiry/overdue scan for review.

    The endpoint is deliberately explicit about its boundary: it evaluates
    effective evidence windows and task deadlines at ``asOf`` but does not
    pretend that a background Worker, notification channel, or automatic task
    creator is present.  A client may omit ``asOf`` to use the current UTC
    instant; tests and incident reviews should pass a fixed, timezone-aware
    timestamp for reproducibility.
    """
    # FastAPI injects a string actor at request time.  Keeping a concrete
    # fallback makes direct Python calls (used by smoke tools and notebooks)
    # safe as well; a ``Depends`` marker must never leak into persisted JSON.
    if not isinstance(actor, str):
        actor = "本机负责人"
    as_of = body.asOf or now()
    if background:
        try:
            dt(as_of)
        except (TypeError, ValueError) as error:
            raise HTTPException(422, "重检时间必须是带时区的 ISO 时间。") from error
        job_id = f"expiry-recheck:{body.requestId}"
        existing = store().job(job_id)
        if existing:
            context = existing.get("context") or {}
            if context.get("asOf") and dt(context["asOf"]) != dt(as_of):
                raise HTTPException(409, "同一 requestId 不能用于不同的重检时间。")
            if existing["status"] in {"failed", "cancelled"}:
                raise HTTPException(409, f"该重检任务已处于 {existing['status']} 终态，请更换 requestId 重试。")
            return existing
        job = store().begin_job(job_id=job_id, kind="expiry_recheck", plan_id=None,
                                idempotency_key=body.requestId, context={"asOf": as_of})
        if job_execution_mode() == "worker":
            return {**job, "scope": "独立 durable worker 执行；任务会在 Worker 领取后运行，可跨 API 进程恢复。"}
        if background_tasks is None:
            raise HTTPException(500, "后台任务调度器不可用。")
        background_tasks.add_task(_run_expiry_background, body.requestId, as_of, actor)
        return job
    try:
        return store().expiry_recheck(request_id=body.requestId, as_of=as_of, actor=actor)
    except HTTPException:
        raise
    except (ValueError, TypeError) as error:
        raise HTTPException(422, str(error)) from error
    except Exception as error:
        raise HTTPException(500, f"到期重检失败：{type(error).__name__}: {str(error)[:180]}") from error

@router.post('/jobs/{job_id}/cancel')
def cancel_job(job_id: str, request: Request, actor: str = Depends(local_owner)):
    """Request cancellation for a durable local job.

    Analysis currently runs in the request process, so cancellation can stop
    queued work or prevent a later retry, but it cannot interrupt Python code
    that is already executing.  The state transition is persisted and a
    synchronous handler must still check it before publishing a result.
    """
    if local_role(request) not in {"safety_owner", "ehs_manager", "administrator"}:
        raise HTTPException(403, "当前工作区身份没有取消分析任务的权限。")
    return store().cancel_job(job_id, reason=f"{actor} 取消任务")


def _run_analysis_sync(
    plan_id: str,
    retry: bool = False,
    actor: str = "本机负责人",
    job_id_override: str | None = None,
    job_started_at: str | None = None,
):
    """Reuse legacy parsing/rules and optional LLM explanation without approval."""
    from .parser import parse_sop
    from .llm import LLMAdapter
    from .models import Review, Evidence, LabStep
    from .rules import ReviewRuleEngine
    w = store().read()
    p = next((p for p in w['plans'] if p['id'] == plan_id), None)
    if p is None:
        raise HTTPException(404, '计划不存在。')
    # A plan revision can have several immutable attempts after a provider
    # timeout or an explicit rerun. Opening Copilot without `retry=true` should
    # show the newest completed attempt for the current revision rather than
    # resurrecting the first (possibly degraded) idempotency record.
    if not retry and job_id_override is None:
        latest = next((run for run in store().analysis_runs(p['id'], 100)
                       if int(run.get('planRevision', -1)) == int(p['revision'])), None)
        if latest:
            run_id = str((latest.get('trace') or {}).get('runId') or '')
            latest_job_id = run_id.removeprefix('copilot:')
            latest_job = store().job(latest_job_id) if latest_job_id else None
            if latest_job and latest_job.get('status') == 'succeeded':
                return {**latest, 'job': latest_job}
    job_id = job_id_override or f"analysis:{p['id']}:v{p['revision']}"
    existing = store().job(job_id)
    if job_id_override is None and retry and existing and existing["status"] in {"failed", "cancelled", "succeeded"}:
        # A retry/rerun is a new durable attempt. Terminal history stays
        # immutable, while the new idempotency key allows the same plan
        # revision to run again after a transient model failure or a provider
        # configuration change. The current plan is never rewritten.
        suffix = "rerun" if existing["status"] == "succeeded" else "retry"
        job_id = f"{job_id}:{suffix}:{str(uuid4())[:8]}"
    job = store().begin_job(job_id=job_id, kind="analysis", plan_id=p['id'], idempotency_key=job_id)
    if job["status"] == "succeeded" and job.get("result"):
        return {**job["result"], "job": job}
    if job["status"] in {"cancelled", "failed"}:
        raise HTTPException(409, f"该分析任务已处于 {job['status']} 终态，请修改计划修订后使用 retry=true 重新发起。")
    store().set_job_running(job_id)
    try:
        parsed = parse_sop(p['source'])
        # Spatial plans already contain reviewer-visible, line-preserving steps.
        # Many institutional SOPs are prose rather than numbered instructions, so
        # the generic parser can legitimately return only its fallback first line.
        # In that case use the plan's confirmed structure for the copilot and rule
        # review instead of pretending the remaining steps do not exist.
        if len(parsed.get('steps', [])) < len(p.get('steps', [])):
            structured_steps = []
            for index, source_step in enumerate(p['steps'], 1):
                text = source_step['text']
                # Re-run the same deterministic parser on the exact source line;
                # do not invent PPE or controls merely because a word suggests a
                # hazard.  A missing control must remain an explicit unknown.
                line_parse = parse_sop(text)
                candidate = line_parse.get('steps', [None])[0]
                if candidate is None:
                    candidate = LabStep(order=index, instruction=text)
                equipment_names = [e['name'] for e in w['equipment'] if e['id'] in source_step.get('equipmentIds', [])]
                structured_steps.append(candidate.model_copy(update={
                    'order': index,
                    'instruction': text,
                    'equipment': sorted(set(candidate.equipment + equipment_names)),
                    'evidence_refs': [f"plan-line-{source_step['line']}"],
                    'document_revision_id': f"planrev-{p['id']}-v{p['revision']}",
                    'source_line': source_step['line'],
                    'source_text_hash': hashlib.sha256(source_step['text'].encode('utf-8')).hexdigest(),
                }))
            parsed['steps'] = structured_steps
            parsed['document_revision_id'] = f"planrev-{p['id']}-v{p['revision']}"
            parsed['ppe'] = sorted({ppe for step in structured_steps for ppe in step.requires_ppe})
            parsed['equipment'] = sorted({item for step in structured_steps for item in step.equipment})
            parsed['ventilation'] = '通风柜' if any(step.ventilation for step in structured_steps) else parsed.get('ventilation')
            parsed['waste_streams'] = sorted({step.waste_stream for step in structured_steps if step.waste_stream})
        adapter = LLMAdapter()
        model_fields = adapter.extract_sop(p['source'])
        parsed, model_used = _merge_model_parse(parsed, model_fields)
        review = Review(id=p['id'], title=p['name'], objective='实验计划辅助文本预审', requester=p['owner'], room=p['labId'],
            sop=p['source'], created_at=now(), updated_at=now(), steps=parsed['steps'], substances=parsed['substances'],
            ppe=parsed['ppe'], equipment=parsed['equipment'], ventilation=parsed['ventilation'], waste_streams=parsed['waste_streams'],
            evidence=[Evidence(**item) for item in parsed['evidence']])
        engine = ReviewRuleEngine()
        evaluation = engine.evaluate(review)
        result = adapter.summarize(review.model_dump(mode='json'), evaluation.model_dump(mode='json'))
        copilot = _copilot_packet(w, p, parsed, evaluation)
        # The trace is intentionally explicit in the response so the UI can show
        # what came from the model, what came from deterministic rules, and what
        # still requires a human decision.  It is metadata, not an approval.
        plan_checks = [item for item in evaluation.results if getattr(item, "rule_id", None)]
        check_counts = {
            "pass": sum(1 for item in plan_checks if item.passed),
            "needs_review": sum(1 for item in plan_checks if not item.passed),
        }
        decision_state = "blocked" if copilot["ruleSummary"]["blocked"] else "needs_review" if copilot["ruleSummary"]["unknown"] else "ready_for_human_review"
        decision_label = {
            "blocked": "存在阻断，安全门保持关闭",
            "needs_review": "存在未知项，需补证后人工复核",
            "ready_for_human_review": "规则范围内未见未满足项，仍需人工复核",
        }[decision_state]
        # This is a provenance ledger for the run, not a claim that the model or
        # public sources certify the experiment.  Keeping the input/output counts
        # next to each stage makes it visible which work was deterministic and
        # which part (if any) came from the configured model.
        trace = {
            # Include the durable job attempt so an explicit rerun of the same
            # plan revision remains a separate immutable history record.
            'runId': f"copilot:{job_id}",
            'generatedAt': now(),
            'stages': [
                {'id': 'understand', 'label': '结构化理解', 'status': 'model' if model_used else 'deterministic', 'detail': f"识别 {len(parsed['steps'])} 个步骤、{len(parsed['substances'])} 个化学品实体", 'inputs': ['计划原文'], 'outputs': ['步骤实体', '化学品实体', '设备语义']},
                {'id': 'check', 'label': '规则核验', 'status': 'rule', 'detail': f"{check_counts['pass']} 项满足，{check_counts['needs_review']} 项待处理", 'inputs': ['结构化步骤', '资源台账', '证据目录'], 'outputs': ['阻断项', '未知项', '规则来源']},
                {'id': 'scenario', 'label': '场景推演', 'status': 'deterministic', 'detail': f"生成 {len(copilot['scenarios'])} 个隔离候选方案", 'inputs': ['阻断项', '未知项', '预约安排'], 'outputs': ['替代资源候选', '先补证候选']},
                {'id': 'review', 'label': '人工复核', 'status': 'required', 'detail': '候选结果需要负责人确认后才可形成计划修订', 'inputs': ['候选结果', '来源与证据'], 'outputs': ['人工意见', '版本化决定']},
            ],
            'model': {'configured': adapter.status()['configured'], 'used': model_used, 'provider': result.provider, 'name': adapter.model or None, 'degradedReason': result.reason},
            'sources': [
                {'kind': 'sop', 'label': '计划原文', 'reference': p['sourceName'], 'detail': f"{len(parsed['steps'])} 个结构化步骤 · 保留原文行号", 'status': 'used'},
                {'kind': 'rules', 'label': '确定性规则包', 'reference': engine.version, 'detail': f"{len(plan_checks)} 条规则结果 · {check_counts['needs_review']} 项待处理", 'status': 'used'},
                {'kind': 'resources', 'label': '资源与证据目录', 'reference': f"{len(w['equipment'])} 台设备 · {len(w['documents'])} 份资料", 'detail': '只依据当前工作区登记，不替代现场核验', 'status': 'used'},
                {'kind': 'model', 'label': '模型解释层', 'reference': adapter.model or '未配置', 'detail': result.reason or '仅生成候选解释；不可自动批准或改写计划', 'status': 'used' if model_used else 'degraded'},
            ],
            'decision': {'state': decision_state, 'label': decision_label, 'approvalRequired': True, 'reason': 'AI 候选和规则结果不能替代机构审批或现场确认。'},
            'source_policy': '所有候选都绑定 SOP 行号、资源检查或证据记录；无法建立来源时保留为待确认。',
        }
        evaluation_record = {
            'evaluationId': uid('evaluation'), 'planId': p['id'], 'planRevision': p['revision'],
            'evaluatedAt': now(), 'basis': basis(w, p), 'status': 'completed',
            'ruleVersion': engine.version, 'results': [r.model_dump(mode='json') for r in evaluation.results],
            'checkCounts': copilot['ruleSummary'], 'decisionState': decision_state,
            'sourceRefs': [p['sourceName'], RULE_PACK],
        }
        response = {'planId': p['id'], 'planRevision': p['revision'], 'ruleVersion': engine.version,
            'parser': 'model-enriched-deterministic-sop-v2' if model_used else 'deterministic-sop-v1', 'provider': result.provider, 'model': adapter.model or None,
            'evaluationId': evaluation_record['evaluationId'], 'degraded': result.provider == 'offline' or result.degraded or not result.available,
            'degradedReason': result.reason,
            'summary': result.text, 'results': [r.model_dump(mode='json') for r in evaluation.results],
            'copilot': copilot,
            'trace': trace,
            'scope': 'AI 负责结构化理解与解释；确定性规则和人工复核负责约束结论。不会自动批准、修改计划或生成事故处置指令。'}
        # A cancellation request can arrive while a synchronous analysis is
        # assembling its result.  Never publish a successful result after the
        # durable job has been cancelled.
        current_job = store().job(job_id)
        if current_job and current_job["status"] == "cancelled":
            raise HTTPException(409, "分析任务已取消，未发布分析结果。")
        store().save_evaluation_run(evaluation_record)
        store().save_analysis_run(response)
        job = store().finish_job(
            job_id, status="succeeded", result=response,
            expected_started_at=job_started_at,
        )
        return {**response, "job": job}
    except HTTPException:
        current_job = store().job(job_id)
        if current_job and current_job["status"] in {"queued", "running"} and (
            job_started_at is None or current_job.get("startedAt") == job_started_at
        ):
            store().finish_job(
                job_id, status="failed", error="分析请求被拒绝或中断。",
                expected_started_at=job_started_at,
            )
        raise
    except Exception as error:
        current_job = store().job(job_id)
        if current_job and current_job["status"] in {"queued", "running"} and (
            job_started_at is None or current_job.get("startedAt") == job_started_at
        ):
            store().finish_job(
                job_id, status="failed",
                error=f"{type(error).__name__}: {str(error)[:500]}",
                expected_started_at=job_started_at,
            )
        raise HTTPException(500, "分析任务失败；请查看任务记录后修改计划或重新发起。") from error


def _prepare_background_analysis(plan_id: str, retry: bool) -> tuple[dict, dict | None, str | None, bool]:
    """Create an idempotent analysis job without doing model work in-request.

    The boolean return value indicates whether this request created the job;
    only the creator may enqueue the process-local runner.  A completed run is
    returned as ``job=None`` so opening Copilot remains read-only/idempotent.
    """
    workspace = store().read()
    plan = next((item for item in workspace["plans"] if item["id"] == plan_id), None)
    if plan is None:
        raise HTTPException(404, "计划不存在。")
    if not retry:
        latest = next((run for run in store().analysis_runs(plan_id, 100)
                       if int(run.get("planRevision", -1)) == int(plan["revision"])), None)
        if latest:
            run_id = str((latest.get("trace") or {}).get("runId") or "")
            latest_job = store().job(run_id.removeprefix("copilot:")) if run_id else None
            if latest_job and latest_job.get("status") == "succeeded":
                return plan, {**latest, "job": latest_job}, None, False
    base_id = f"analysis:{plan_id}:v{plan['revision']}"
    job_id = base_id
    existing = store().job(job_id)
    if retry and existing and existing["status"] in {"failed", "cancelled", "succeeded"}:
        job_id = f"{base_id}:{'rerun' if existing['status'] == 'succeeded' else 'retry'}:{str(uuid4())[:8]}"
        existing = None
    if existing and existing["status"] in {"failed", "cancelled"}:
        raise HTTPException(409, f"该分析任务已处于 {existing['status']} 终态，请修改计划修订后使用 retry=true 重新发起。")
    job = store().begin_job(job_id=job_id, kind="analysis", plan_id=plan_id, idempotency_key=job_id)
    return plan, job, job_id, existing is None


def _run_analysis_background(plan_id: str, job_id: str, actor: str) -> None:
    """Run one durable analysis in FastAPI's process-local background pool.

    This is deliberately not described as a multi-process Worker.  A restart
    marks queued/running rows failed; the browser can inspect the error and
    explicitly retry with a new immutable attempt.
    """
    try:
        # The queued row already owns an immutable attempt ID.  Force the
        # retry path so a completed attempt for the same plan revision cannot
        # short-circuit this worker and leave the newly queued row unfinished.
        _run_analysis_sync(plan_id, retry=True, actor=actor, job_id_override=job_id)
    except Exception:
        # The synchronous runner has already recorded the terminal error.  A
        # background callback must not surface an exception after the response.
        return


@router.post('/plans/{plan_id}/analysis')
def analyze_plan(
    plan_id: str,
    retry: bool = False,
    background: bool = False,
    background_tasks: BackgroundTasks = None,
    actor: str = Depends(local_owner),
):
    """Analyze a plan synchronously or via an explicit process-local job.

    ``background=true`` returns a durable queued job immediately.  It is
    intended for the browser UI to poll; it does not claim a persistent worker
    or survive a process restart.  The default remains synchronous for API
    compatibility and deterministic command-line checks.
    """
    if not background:
        return _run_analysis_sync(plan_id, retry=retry, actor=actor)
    plan, prepared, job_id, created = _prepare_background_analysis(plan_id, retry)
    if job_id is None:
        return prepared
    job = prepared
    if job is None:
        raise HTTPException(500, "分析任务创建失败。")
    if created:
        if job_execution_mode() == "worker":
            return {
                "planId": plan_id,
                "planRevision": plan["revision"],
                "background": True,
                "job": job,
                "scope": "独立 durable worker 执行；任务会在 Worker 领取后运行，可跨 API 进程恢复。",
            }
        if background_tasks is None:
            # Direct Python calls are useful in smoke tests and must fail with
            # a meaningful HTTP error rather than an AttributeError.  The
            # actual FastAPI route always injects BackgroundTasks here.
            raise HTTPException(500, "后台任务调度器不可用。")
        background_tasks.add_task(_run_analysis_background, plan_id, job_id, actor)
    return {
        "planId": plan_id,
        "planRevision": plan["revision"],
        "background": True,
        "job": job,
        "scope": "进程内后台执行；服务重启不会伪称任务继续，失败后需显式 retry。",
    }
